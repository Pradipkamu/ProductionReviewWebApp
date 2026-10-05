from __future__ import annotations

from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..auth import get_current_user
from ..db import get_db
from ..models import (
    CastingDefectDaily, CastingDefectPhenomenon, Customer, CustomerRejectionDaily,
    CustomerRejectionPhenomenon, Product, User, Vendor,
)
from ..services.special_quality_import import create_phenomenon, norm

router = APIRouter(tags=["casting and customer quality"])


class PhenomenonCreate(BaseModel):
    name: str = Field(min_length=2, max_length=220)
    phenomenon_group: str | None = Field(default=None, max_length=120)
    criticality: str = Field(default="NORMAL", pattern="^(LOW|NORMAL|HIGH|CRITICAL)$")


class CastingCreate(BaseModel):
    defect_date: date
    shift: str = Field(pattern="^(A|B|C|General)$")
    product_id: int
    vendor_id: int | None = None
    heat_batch_no: str | None = Field(default=None, max_length=120)
    phenomenon_id: int
    inspected_qty: Decimal = Field(gt=0)
    defect_qty: Decimal = Field(ge=0)
    rework_qty: Decimal = Field(default=Decimal("0"), ge=0)
    scrap_qty: Decimal = Field(default=Decimal("0"), ge=0)
    remark: str | None = Field(default=None, max_length=2000)


class CustomerRejectionCreate(BaseModel):
    rejection_date: date
    customer_id: int
    product_id: int
    reference_no: str | None = Field(default=None, max_length=120)
    batch_no: str | None = Field(default=None, max_length=120)
    phenomenon_id: int
    dispatch_qty: Decimal = Field(gt=0)
    reject_qty: Decimal = Field(ge=0)
    remark: str | None = Field(default=None, max_length=2000)


def _phenomena(db: Session, model):
    return [{"id": x.id, "code": x.code, "name": x.name, "phenomenon_group": x.phenomenon_group,
             "criticality": x.criticality, "active": x.is_active}
            for x in db.scalars(select(model).order_by(model.name)).all()]


@router.get("/casting-quality/phenomena")
def casting_phenomena(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    return _phenomena(db, CastingDefectPhenomenon)


@router.post("/casting-quality/phenomena")
def add_casting_phenomenon(payload: PhenomenonCreate, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    row = create_phenomenon(db, CastingDefectPhenomenon, payload.name, group=payload.phenomenon_group, criticality=payload.criticality)
    db.commit(); db.refresh(row)
    return {"id": row.id, "code": row.code, "name": row.name}


@router.get("/customer-quality/phenomena")
def customer_phenomena(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    return _phenomena(db, CustomerRejectionPhenomenon)


@router.post("/customer-quality/phenomena")
def add_customer_phenomenon(payload: PhenomenonCreate, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    row = create_phenomenon(db, CustomerRejectionPhenomenon, payload.name, group=payload.phenomenon_group, criticality=payload.criticality)
    db.commit(); db.refresh(row)
    return {"id": row.id, "code": row.code, "name": row.name}


def _ppm(reject: Decimal, denominator: Decimal) -> Decimal:
    return (reject * Decimal("1000000") / denominator).quantize(Decimal("0.001"))


@router.post("/casting-quality/daily")
def add_casting(payload: CastingCreate, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    product = db.get(Product, payload.product_id)
    vendor = db.get(Vendor, payload.vendor_id) if payload.vendor_id else None
    phenomenon = db.get(CastingDefectPhenomenon, payload.phenomenon_id)
    if not product: raise HTTPException(422, "Product was not found")
    if payload.vendor_id and not vendor: raise HTTPException(422, "Vendor was not found")
    if not phenomenon or not phenomenon.is_active: raise HTTPException(422, "Casting phenomenon was not found or is inactive")
    if payload.defect_qty > payload.inspected_qty: raise HTTPException(422, "Defect Qty cannot exceed Inspected Qty")
    if payload.rework_qty + payload.scrap_qty > payload.defect_qty:
        raise HTTPException(422, "Rework Qty plus Scrap Qty cannot exceed Defect Qty")
    batch = payload.heat_batch_no.strip() if payload.heat_batch_no else None
    scope_key = "|".join([payload.defect_date.isoformat(), norm(payload.shift), str(product.id), str(vendor.id if vendor else 0), norm(batch)])
    prior_denominator = db.scalar(select(CastingDefectDaily.inspected_qty).where(CastingDefectDaily.record_key.startswith(scope_key + "|")).limit(1))
    if prior_denominator is not None and prior_denominator != payload.inspected_qty:
        raise HTTPException(422, "Inspected Qty must match other phenomena for the same date/shift/product/vendor/batch")
    key = scope_key + "|" + str(phenomenon.id)
    if db.scalar(select(CastingDefectDaily.id).where(CastingDefectDaily.record_key == key)):
        raise HTTPException(409, "This date/shift/product/vendor/batch/phenomenon record already exists")
    row = CastingDefectDaily(record_key=key, defect_date=payload.defect_date, shift=payload.shift, product_id=product.id,
        plant=product.plant, vendor_id=vendor.id if vendor else None, heat_batch_no=batch, phenomenon_id=phenomenon.id,
        inspected_qty=payload.inspected_qty, defect_qty=payload.defect_qty, rework_qty=payload.rework_qty,
        scrap_qty=payload.scrap_qty, ppm=_ppm(payload.defect_qty, payload.inspected_qty),
        remark=payload.remark.strip() if payload.remark else None, source="MANUAL", entered_by_id=user.id)
    db.add(row); db.commit(); db.refresh(row)
    return {"id": row.id, "ppm": float(row.ppm)}


@router.post("/customer-quality/daily")
def add_customer_rejection(payload: CustomerRejectionCreate, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    customer = db.get(Customer, payload.customer_id)
    product = db.get(Product, payload.product_id)
    phenomenon = db.get(CustomerRejectionPhenomenon, payload.phenomenon_id)
    if not customer: raise HTTPException(422, "Customer was not found")
    if not product: raise HTTPException(422, "Product was not found")
    if product.customer_id and product.customer_id != customer.id:
        raise HTTPException(422, "Product is not assigned to the selected Customer")
    if not phenomenon or not phenomenon.is_active: raise HTTPException(422, "Customer rejection phenomenon was not found or is inactive")
    if payload.reject_qty > payload.dispatch_qty: raise HTTPException(422, "Reject Qty cannot exceed Dispatch Qty")
    reference = payload.reference_no.strip() if payload.reference_no else None
    batch = payload.batch_no.strip() if payload.batch_no else None
    scope_key = "|".join([payload.rejection_date.isoformat(), str(customer.id), str(product.id), norm(reference), norm(batch)])
    prior_denominator = db.scalar(select(CustomerRejectionDaily.dispatch_qty).where(CustomerRejectionDaily.record_key.startswith(scope_key + "|")).limit(1))
    if prior_denominator is not None and prior_denominator != payload.dispatch_qty:
        raise HTTPException(422, "Dispatch Qty must match other phenomena for the same date/customer/product/reference/batch")
    key = scope_key + "|" + str(phenomenon.id)
    if db.scalar(select(CustomerRejectionDaily.id).where(CustomerRejectionDaily.record_key == key)):
        raise HTTPException(409, "This date/customer/product/reference/batch/phenomenon record already exists")
    row = CustomerRejectionDaily(record_key=key, rejection_date=payload.rejection_date, customer_id=customer.id,
        product_id=product.id, reference_no=reference, batch_no=batch, phenomenon_id=phenomenon.id,
        dispatch_qty=payload.dispatch_qty, reject_qty=payload.reject_qty,
        ppm=_ppm(payload.reject_qty, payload.dispatch_qty), remark=payload.remark.strip() if payload.remark else None,
        source="MANUAL", entered_by_id=user.id)
    db.add(row); db.commit(); db.refresh(row)
    return {"id": row.id, "ppm": float(row.ppm)}


@router.get("/casting-quality/report")
def casting_report(from_date: date, to_date: date, product_id: int | None = None, vendor_id: int | None = None,
                   phenomenon_id: int | None = None, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    q = (select(CastingDefectDaily, Product, Vendor, CastingDefectPhenomenon)
         .join(Product, Product.id == CastingDefectDaily.product_id)
         .outerjoin(Vendor, Vendor.id == CastingDefectDaily.vendor_id)
         .join(CastingDefectPhenomenon, CastingDefectPhenomenon.id == CastingDefectDaily.phenomenon_id)
         .where(CastingDefectDaily.defect_date.between(from_date, to_date)))
    if product_id: q = q.where(CastingDefectDaily.product_id == product_id)
    if vendor_id: q = q.where(CastingDefectDaily.vendor_id == vendor_id)
    if phenomenon_id: q = q.where(CastingDefectDaily.phenomenon_id == phenomenon_id)
    rows = db.execute(q.order_by(CastingDefectDaily.defect_date.desc(), CastingDefectDaily.id.desc())).all()
    data = [{"id": r.id, "date": r.defect_date, "shift": r.shift, "product": p.name, "product_id": p.id,
             "vendor": v.name if v else None, "vendor_id": v.id if v else None, "heat_batch_no": r.heat_batch_no,
             "phenomenon": ph.name, "phenomenon_id": ph.id, "inspected_qty": float(r.inspected_qty),
             "defect_qty": float(r.defect_qty), "rework_qty": float(r.rework_qty), "scrap_qty": float(r.scrap_qty),
             "ppm": float(r.ppm), "remark": r.remark, "source": r.source} for r, p, v, ph in rows]
    return _report_payload(data, "inspected_qty", "defect_qty", extras=("rework_qty", "scrap_qty"),
                           denominator_scope=("date", "shift", "product_id", "vendor_id", "heat_batch_no"))


@router.get("/customer-quality/report")
def customer_report(from_date: date, to_date: date, customer_id: int | None = None, product_id: int | None = None,
                    phenomenon_id: int | None = None, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    q = (select(CustomerRejectionDaily, Customer, Product, CustomerRejectionPhenomenon)
         .join(Customer, Customer.id == CustomerRejectionDaily.customer_id)
         .join(Product, Product.id == CustomerRejectionDaily.product_id)
         .join(CustomerRejectionPhenomenon, CustomerRejectionPhenomenon.id == CustomerRejectionDaily.phenomenon_id)
         .where(CustomerRejectionDaily.rejection_date.between(from_date, to_date)))
    if customer_id: q = q.where(CustomerRejectionDaily.customer_id == customer_id)
    if product_id: q = q.where(CustomerRejectionDaily.product_id == product_id)
    if phenomenon_id: q = q.where(CustomerRejectionDaily.phenomenon_id == phenomenon_id)
    rows = db.execute(q.order_by(CustomerRejectionDaily.rejection_date.desc(), CustomerRejectionDaily.id.desc())).all()
    data = [{"id": r.id, "date": r.rejection_date, "customer": c.name, "customer_id": c.id,
             "product": p.name, "product_id": p.id, "reference_no": r.reference_no, "batch_no": r.batch_no,
             "phenomenon": ph.name, "phenomenon_id": ph.id, "dispatch_qty": float(r.dispatch_qty),
             "reject_qty": float(r.reject_qty), "ppm": float(r.ppm), "remark": r.remark, "source": r.source}
            for r, c, p, ph in rows]
    return _report_payload(data, "dispatch_qty", "reject_qty",
                           denominator_scope=("date", "customer_id", "product_id", "reference_no", "batch_no"))


def _report_payload(rows: list[dict], denominator_key: str, reject_key: str, extras: tuple[str, ...] = (),
                    denominator_scope: tuple[str, ...] = ()):
    denominators: dict[tuple, float] = {}
    for row in rows:
        scope = tuple(row.get(key) for key in denominator_scope) if denominator_scope else (row["id"],)
        denominators[scope] = max(denominators.get(scope, 0), row[denominator_key])
    denominator = sum(denominators.values())
    rejected = sum(x[reject_key] for x in rows)
    summary = {denominator_key: denominator, reject_key: rejected, "ppm": rejected * 1000000 / denominator if denominator else 0,
               "records": len(rows)}
    for key in extras: summary[key] = sum(x[key] for x in rows)
    trend_rejected, trend_denominators = defaultdict(float), defaultdict(dict)
    phenomena, products, customers = defaultdict(float), defaultdict(float), defaultdict(float)
    for row in rows:
        label = str(row["date"])
        scope = tuple(row.get(key) for key in denominator_scope) if denominator_scope else (row["id"],)
        trend_denominators[label][scope] = max(trend_denominators[label].get(scope, 0), row[denominator_key])
        trend_rejected[label] += row[reject_key]
        phenomena[row["phenomenon"]] += row[reject_key]
        products[row["product"]] += row[reject_key]
        if row.get("customer"): customers[row["customer"]] += row[reject_key]
    trend_rows = []
    for label in sorted(trend_rejected):
        daily_denominator = sum(trend_denominators[label].values())
        daily_rejected = trend_rejected[label]
        trend_rows.append({"label": label, denominator_key: daily_denominator, reject_key: daily_rejected,
                           "ppm": daily_rejected * 1000000 / daily_denominator if daily_denominator else 0})
    pareto = lambda values: [{"name": name, reject_key: value} for name, value in sorted(values.items(), key=lambda x: x[1], reverse=True)]
    return {"summary": summary, "trend": trend_rows, "phenomenon_pareto": pareto(phenomena),
            "product_pareto": pareto(products), "customer_pareto": pareto(customers), "rows": rows}


def _template(db: Session, kind: str):
    products = db.execute(select(Product, Customer).outerjoin(Customer, Customer.id == Product.customer_id).order_by(Product.name)).all()
    vendors = db.scalars(select(Vendor).where(Vendor.is_active.is_(True)).order_by(Vendor.name)).all()
    if kind == "casting":
        title, sheet, phenomena = "Casting Defect Upload", "Casting_Defect_Data", db.scalars(select(CastingDefectPhenomenon).where(CastingDefectPhenomenon.is_active.is_(True)).order_by(CastingDefectPhenomenon.name)).all()
        columns = ["Date", "Shift", "Product", "Vendor", "Heat/Batch No", "Inspected Qty", "Phenomenon", "Defect Qty", "Rework Qty", "Scrap Qty", "Remark"]
        required = {1, 2, 3, 6, 7, 8}
    else:
        title, sheet, phenomena = "Customer Rejection Upload", "Customer_Rejection_Data", db.scalars(select(CustomerRejectionPhenomenon).where(CustomerRejectionPhenomenon.is_active.is_(True)).order_by(CustomerRejectionPhenomenon.name)).all()
        columns = ["Date", "Customer", "Product", "Reference No", "Batch No", "Dispatch Qty", "Phenomenon", "Reject Qty", "Remark"]
        required = {1, 2, 3, 6, 7, 8}
    wb = Workbook(); instructions = wb.active; instructions.title = "Instructions"; ws = wb.create_sheet(sheet); masters = wb.create_sheet("Masters")
    instructions.append([title + " v0.5.6"]); instructions.append(["Use", f"Enter one {kind} rejection/defect combination per row."])
    instructions.append(["Required", ", ".join(columns[i-1] for i in sorted(required))])
    instructions.append(["PPM", "Calculated by the application from Reject/Defect Qty divided by Dispatch/Inspected Qty × 1,000,000."])
    instructions.append(["Control", "Select current master values from dropdowns, preview the workbook, then confirm only when there are no errors."])
    instructions.column_dimensions["A"].width = 18; instructions.column_dimensions["B"].width = 105
    instructions["A1"].font = Font(size=15, bold=True, color="1F4E78")
    ws.append(columns)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="C65911" if cell.column in required else "1F4E78")
        cell.font = Font(color="FFFFFF", bold=True); cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:{chr(64+len(columns))}501"
    for index, width in enumerate([14,22,30,24,20,18,34,16,16,16,40][:len(columns)], 1): ws.column_dimensions[chr(64+index)].width = width
    masters.append(["Product", "Customer", "Vendor", "Phenomenon", "Shift"])
    count = max(len(products), len(vendors), len(phenomena), 4, 1)
    for i in range(count):
        product, customer = products[i] if i < len(products) else (None, None)
        masters.append([product.name if product else None, customer.name if customer else None,
                        vendors[i].name if i < len(vendors) else None, phenomena[i].name if i < len(phenomena) else None,
                        ["A", "B", "C", "General"][i] if i < 4 else None])
    ranges = {"ProductList": ("A", len(products)), "CustomerList": ("B", len({c.id for _, c in products if c})),
              "VendorList": ("C", len(vendors)), "PhenomenonList": ("D", len(phenomena)), "ShiftList": ("E", 4)}
    # Rebuild a compact, unique customer list so the dropdown has no blanks/duplicates.
    unique_customers = sorted({c.name for _, c in products if c})
    for row in range(2, count + 2): masters.cell(row, 2).value = unique_customers[row-2] if row-2 < len(unique_customers) else None
    for name, (column, size) in ranges.items(): wb.defined_names.add(DefinedName(name, attr_text=f"'Masters'!${column}$2:${column}${max(2,size+1)}"))
    def add_list(column: str, range_name: str, allow_blank: bool):
        dv = DataValidation(type="list", formula1=f"={range_name}", allow_blank=allow_blank); dv.errorStyle = "stop"; dv.showErrorMessage = True
        ws.add_data_validation(dv); dv.add(f"{column}2:{column}501")
    if kind == "casting": add_list("B", "ShiftList", False); add_list("C", "ProductList", False); add_list("D", "VendorList", True); add_list("G", "PhenomenonList", False)
    else: add_list("B", "CustomerList", False); add_list("C", "ProductList", False); add_list("G", "PhenomenonList", False)
    date_dv = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2100,12,31)", allow_blank=False); ws.add_data_validation(date_dv); date_dv.add("A2:A501")
    denominator_dv = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=False); denominator_dv.errorStyle = "stop"; denominator_dv.showErrorMessage = True
    ws.add_data_validation(denominator_dv); denominator_dv.add("F2:F501")
    qty_dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=False); qty_dv.errorStyle = "stop"; qty_dv.showErrorMessage = True
    ws.add_data_validation(qty_dv); qty_dv.add("H2:H501")
    if kind == "casting":
        optional_qty_dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True); optional_qty_dv.errorStyle = "stop"; optional_qty_dv.showErrorMessage = True
        ws.add_data_validation(optional_qty_dv); optional_qty_dv.add("I2:J501")
    for row in range(2, 502):
        for column in range(1, len(columns)+1):
            ws.cell(row, column).protection = Protection(locked=False)
            if column in required: ws.cell(row, column).fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(row, 1).number_format = "dd-mmm-yyyy"
    ws.protection.sheet = True; ws.protection.password = "PRWUpload"; ws.protection.autoFilter = False
    masters.sheet_state = "hidden"
    bio = BytesIO(); wb.save(bio); bio.seek(0)
    filename = "Casting_Defect_Upload_v0.5.6.xlsx" if kind == "casting" else "Customer_Rejection_Upload_v0.5.6.xlsx"
    return StreamingResponse(bio, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get("/casting-quality/template")
def casting_template(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    return _template(db, "casting")


@router.get("/customer-quality/template")
def customer_template(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    return _template(db, "customer")
