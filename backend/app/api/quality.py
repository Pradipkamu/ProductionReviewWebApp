from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
import json

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook as XLWorkbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from pydantic import BaseModel
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..auth import get_current_user
from ..db import get_db
from ..enums import ActionStatus, Priority
from ..models import (
    Action, ActionContext, ActionHistory, ActionWhyWhy, Customer, DailyMIS, Machine, Operation, Product,
    QualityActionLink, QualityPhenomenon, QualityRejectionDaily, QualityRejectionImportBatch,
    QualityRejectionMonthlyHistory, ReviewActionLink, ReviewSession, RouteOperation, User,
)
from ..services.filtering import csv_ints, csv_strings
from ..services.quality_import import (
    _daily_record_key, _route_for_date, denominator_for_rejection,
    upsert_phenomenon,
)

router = APIRouter(prefix="/quality", tags=["quality"])
class PhenomenonCreate(BaseModel):
    name: str
    phenomenon_group: str | None = None
    default_responsible_team: str | None = "Operation"
    criticality: str = "NORMAL"


class QualityActionCreate(BaseModel):
    owner_id: int | None = None
    due_at: datetime | None = None
    priority: Priority = Priority.HIGH
    action_description: str = "Investigate using standard Why-Why and implement corrective action"


class QualityRejectionManualCreate(BaseModel):
    rejection_date: date
    shift: str
    product_id: int
    detection_route_operation_id: int
    responsible_route_operation_id: int | None = None
    machine_id: int | None = None
    phenomenon_id: int
    reject_qty: Decimal
    rework_qty: Decimal = Decimal("0")
    scrap_qty: Decimal = Decimal("0")
    remark: str | None = None
    action_required: bool = False


def _num(value) -> float:
    return float(value or 0)


def _bool(value) -> bool:
    return bool(value)


def _label_key(value) -> str:
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def _by_id(db: Session, model, ids: set[int]) -> dict:
    return {item.id: item for item in db.scalars(select(model).where(model.id.in_(ids))).all()} if ids else {}


def _daily_context(db: Session, rows: list[QualityRejectionDaily]) -> dict:
    products = _by_id(db, Product, {x.product_id for x in rows})
    phenomena = _by_id(db, QualityPhenomenon, {x.phenomenon_id for x in rows})
    machines = _by_id(db, Machine, {x.machine_id for x in rows if x.machine_id})
    routes = _by_id(db, RouteOperation, {
        route_id for x in rows for route_id in (x.detection_route_operation_id, x.responsible_route_operation_id) if route_id
    })
    operations = _by_id(db, Operation, {x.operation_id for x in routes.values()})
    links = db.scalars(select(QualityActionLink).where(
        QualityActionLink.daily_rejection_id.in_([x.id for x in rows])
    ).order_by(QualityActionLink.id.desc())).all() if rows else []
    latest_links = {}
    for link in links:
        latest_links.setdefault(link.daily_rejection_id, link.action_id)
    actions = _by_id(db, Action, set(latest_links.values()))
    return dict(products=products, phenomena=phenomena, machines=machines,
                routes=routes, operations=operations, latest_links=latest_links, actions=actions)


def _serialize_daily(row: QualityRejectionDaily, context: dict) -> dict:
    p = context['products'].get(row.product_id)
    ph = context['phenomena'].get(row.phenomenon_id)
    machine = context['machines'].get(row.machine_id)
    det = context['routes'].get(row.detection_route_operation_id)
    resp = context['routes'].get(row.responsible_route_operation_id)
    det_op = context['operations'].get(det.operation_id) if det else None
    resp_op = context['operations'].get(resp.operation_id) if resp else None
    action = context['actions'].get(context['latest_links'].get(row.id))
    return {
        "id": row.id,
        "date": row.rejection_date,
        "shift": row.shift,
        "product_id": row.product_id,
        "product": p.name if p else None,
        "customer_id": p.customer_id if p else None,
        "plant": row.plant or (p.plant if p else None),
        "product_group": p.product_group if p else None,
        "detection_route_operation_id": row.detection_route_operation_id,
        "detection_process": det_op.name if det_op else None,
        "responsible_route_operation_id": row.responsible_route_operation_id,
        "responsible_process": resp_op.name if resp_op else None,
        "responsible_team": row.responsible_team,
        "machine_id": row.machine_id,
        "machine": machine.name if machine else None,
        "phenomenon_id": row.phenomenon_id,
        "phenomenon": ph.name if ph else None,
        "phenomenon_group": ph.phenomenon_group if ph else None,
        "reject_qty": _num(row.reject_qty),
        "rework_qty": _num(row.rework_qty),
        "scrap_qty": _num(row.scrap_qty),
        "denominator_source": row.denominator_source,
        "denominator_qty": _num(row.denominator_qty) if row.denominator_qty is not None else None,
        "ppm": _num(row.ppm) if row.ppm is not None else None,
        "remark": row.remark,
        "action_required": row.action_required,
        "action": {"id": action.id, "action_no": action.action_no, "status": action.status.value} if action else None,
    }


def _daily_rows(db: Session, *, from_date: date, to_date: date, plant: str | None = None,
                product_id: str | int | None = None, phenomenon_id: str | int | None = None,
                operation_id: str | int | None = None, machine_id: str | int | None = None, shift: str | None = None,
                customer_id: str | int | None = None, product_group: str | None = None) -> list[QualityRejectionDaily]:
    q = select(QualityRejectionDaily).join(Product, Product.id == QualityRejectionDaily.product_id).where(
        QualityRejectionDaily.rejection_date >= from_date,
        QualityRejectionDaily.rejection_date <= to_date,
    )
    plants = csv_strings(plant)
    product_ids = csv_ints(product_id)
    phenomenon_ids = csv_ints(phenomenon_id)
    operation_ids = csv_ints(operation_id)
    machine_ids = csv_ints(machine_id)
    shifts = csv_strings(shift)
    customer_ids = csv_ints(customer_id)
    groups = csv_strings(product_group)
    if plants:
        q = q.where(QualityRejectionDaily.plant.in_(plants))
    if product_ids:
        q = q.where(QualityRejectionDaily.product_id.in_(product_ids))
    if phenomenon_ids:
        q = q.where(QualityRejectionDaily.phenomenon_id.in_(phenomenon_ids))
    if operation_ids:
        q = q.join(RouteOperation, RouteOperation.id == QualityRejectionDaily.detection_route_operation_id).where(RouteOperation.operation_id.in_(operation_ids))
    if machine_ids:
        q = q.where(QualityRejectionDaily.machine_id.in_(machine_ids))
    if shifts:
        q = q.where(QualityRejectionDaily.shift.in_(shifts))
    if customer_ids:
        q = q.where(Product.customer_id.in_(customer_ids))
    if groups:
        q = q.where(Product.product_group.in_(groups))
    return list(db.scalars(q.order_by(QualityRejectionDaily.rejection_date.desc(), QualityRejectionDaily.id.desc())).all())


def _aggregate_daily(rows: list[QualityRejectionDaily]) -> dict:
    reject = sum((_num(x.reject_qty) for x in rows), 0.0)
    rework = sum((_num(x.rework_qty) for x in rows), 0.0)
    scrap = sum((_num(x.scrap_qty) for x in rows), 0.0)
    denoms: dict[tuple, float] = {}
    pending = 0
    for x in rows:
        key = (x.rejection_date, x.shift, x.product_id, x.detection_route_operation_id, x.machine_id, x.denominator_source)
        if x.denominator_qty is not None and _num(x.denominator_qty) > 0:
            denoms[key] = max(denoms.get(key, 0.0), _num(x.denominator_qty))
        elif _num(x.reject_qty) > 0:
            pending += 1
    denominator = sum(denoms.values())
    ppm = reject / denominator * 1_000_000 if denominator > 0 and pending == 0 else None
    return {"reject_qty": reject, "rework_qty": rework, "scrap_qty": scrap, "denominator_qty": denominator,
            "ppm": ppm, "ppm_pending_rows": pending}


def _month_start(value: date) -> date:
    return value.replace(day=1)


def _is_dispatch_denominator(value: str | None) -> bool:
    text = str(value or "").strip().upper().replace(" ", "_")
    return "DISP" in text or "DISPATCH" in text


def _history_dispatch_resolution(
    db: Session,
    rows: list[QualityRejectionMonthlyHistory],
    *,
    plant_filter: str | None = None,
) -> dict[int, tuple[float | None, str | None]]:
    """Resolve historical Dispatch Done denominator without breaking filter scope.

    An unfiltered product/month uses matching Historical Daily MIS as the current
    authoritative dispatch total. A plant-filtered report uses explicit plant/product
    Dispatch_Qty because Daily MIS has no independent historical plant split. When
    one Product + Month exists under multiple historical plants and only a subset is
    selected, a product-total MIS denominator must not be reused for that numerator.
    """
    if not rows:
        return {}

    selected_plants = set(csv_strings(plant_filter))
    # Dispatch Done is a denominator, not part of the rejection numerator.
    # Resolve it for defect-detail rows too so each historical phenomenon can show
    # PPM from the same product/month Historical Daily MIS dispatch total. The
    # include_in_aggregate flag continues to control rejection aggregation only.
    eligible_keys: set[tuple[date, int]] = set()
    for x in rows:
        if _is_dispatch_denominator(x.denominator_source):
            eligible_keys.add((_month_start(x.month), x.product_id))

    monthly: dict[tuple[date, int], float] = defaultdict(float)
    all_plants_by_key: dict[tuple[date, int], set[str]] = defaultdict(set)
    if eligible_keys:
        months = [m for m, _ in eligible_keys]
        product_ids = sorted({pid for _, pid in eligible_keys})
        start = min(months)
        last = max(months)
        next_month = (last.replace(day=28) + timedelta(days=4)).replace(day=1)

        mis_q = select(DailyMIS).where(
            DailyMIS.product_id.in_(product_ids),
            DailyMIS.mis_date >= start,
            DailyMIS.mis_date < next_month,
        )
        for mis in db.scalars(mis_q).all():
            monthly[(_month_start(mis.mis_date), mis.product_id)] += _num(mis.actual_qty)

        if selected_plants:
            plant_q = select(
                QualityRejectionMonthlyHistory.month,
                QualityRejectionMonthlyHistory.product_id,
                QualityRejectionMonthlyHistory.plant,
            ).where(
                QualityRejectionMonthlyHistory.product_id.in_(product_ids),
                QualityRejectionMonthlyHistory.month >= start,
                QualityRejectionMonthlyHistory.month <= last,
            )
            for month, product_id, plant in db.execute(plant_q).all():
                if plant not in (None, ""):
                    all_plants_by_key[(_month_start(month), product_id)].add(str(plant).strip())

    resolved: dict[int, tuple[float | None, str | None]] = {}
    for x in rows:
        key = (_month_start(x.month), x.product_id)
        eligible = _is_dispatch_denominator(x.denominator_source)
        stored_qty = _num(x.dispatch_qty) if x.dispatch_qty is not None else 0.0
        mis_qty = monthly.get(key, 0.0)

        # A plant filter is scope-safe only when it covers every historical plant for
        # this Product + Month. Otherwise Daily MIS is too broad because it has no
        # separate plant dimension.
        known_plants = all_plants_by_key.get(key, set())
        plant_scope_restricted = bool(selected_plants and known_plants and not known_plants.issubset(selected_plants))

        # Explicit plant-scoped Dispatch_Qty must win over Daily MIS. Historical
        # imports from older app versions may have persisted the MIS fallback into
        # dispatch_qty; if that value exactly equals product-total MIS while the
        # plant scope is restricted, treat it as unsafe rather than as explicit.
        stored_looks_like_mis = bool(stored_qty > 0 and mis_qty > 0 and abs(stored_qty - mis_qty) < 0.0005)
        unsafe_stored = plant_scope_restricted and stored_looks_like_mis
        if selected_plants and stored_qty > 0 and not unsafe_stored:
            resolved[x.id] = (stored_qty, "UPLOADED")
        elif eligible and mis_qty > 0 and not plant_scope_restricted:
            resolved[x.id] = (mis_qty, "MIS_HISTORY")
        elif stored_qty > 0 and not unsafe_stored:
            resolved[x.id] = (stored_qty, "UPLOADED")
        elif plant_scope_restricted:
            resolved[x.id] = (None, "PLANT_SCOPE_UNAVAILABLE")
        else:
            resolved[x.id] = (None, None)
    return resolved


def _historical_rows(db: Session, *, from_date: date, to_date: date, plant: str | None = None,
                     product_id: str | int | None = None, phenomenon_id: str | int | None = None,
                     customer_id: str | int | None = None, product_group: str | None = None,
                     include_detail: bool = False) -> list[QualityRejectionMonthlyHistory]:
    first = _month_start(from_date)
    last = _month_start(to_date)
    q = select(QualityRejectionMonthlyHistory).join(Product, Product.id == QualityRejectionMonthlyHistory.product_id).where(
        QualityRejectionMonthlyHistory.month >= first,
        QualityRejectionMonthlyHistory.month <= last,
    )
    plants = csv_strings(plant)
    product_ids = csv_ints(product_id)
    phenomenon_ids = csv_ints(phenomenon_id)
    customer_ids = csv_ints(customer_id)
    groups = csv_strings(product_group)
    if plants:
        q = q.where(QualityRejectionMonthlyHistory.plant.in_(plants))
    elif not include_detail:
        # Aggregate calculations use only approved aggregate rows. Detail/history
        # views may still retrieve non-aggregate phenomenon rows without allowing
        # them to double-count management totals.
        q = q.where(QualityRejectionMonthlyHistory.include_in_aggregate.is_(True))
    if product_ids:
        q = q.where(QualityRejectionMonthlyHistory.product_id.in_(product_ids))
    if phenomenon_ids:
        q = q.where(QualityRejectionMonthlyHistory.phenomenon_id.in_(phenomenon_ids))
    if customer_ids:
        q = q.where(Product.customer_id.in_(customer_ids))
    if groups:
        q = q.where(Product.product_group.in_(groups))
    return list(db.scalars(q.order_by(QualityRejectionMonthlyHistory.month.desc(), QualityRejectionMonthlyHistory.id.desc())).all())


def _exclude_history_covered_by_daily(history: list[QualityRejectionMonthlyHistory],
                                      daily: list[QualityRejectionDaily]) -> list[QualityRejectionMonthlyHistory]:
    # Historical monthly aggregates are authoritative for their product/month.
    # A partial daily upload must not make the completed historical month disappear.
    # Callers that combine both sources should instead exclude overlapping DAILY rows.
    return history


def _exclude_daily_covered_by_history(daily: list[QualityRejectionDaily],
                                      history: list[QualityRejectionMonthlyHistory]) -> list[QualityRejectionDaily]:
    covered = {(_month_start(x.month), x.product_id) for x in history}
    return [x for x in daily if (_month_start(x.rejection_date), x.product_id) not in covered]


def _aggregate_history(rows: list[QualityRejectionMonthlyHistory], dispatch_resolution: dict[int, tuple[float | None, str | None]] | None = None) -> dict:
    reject = sum((_num(x.reject_qty) for x in rows), 0.0)
    denoms: dict[tuple, float] = {}
    pending_keys: set[tuple] = set()
    for x in rows:
        key = (_month_start(x.month), x.product_id, x.source, x.source_sheet, x.record_scope, x.plant, x.denominator_source)
        qty = dispatch_resolution.get(x.id, (None, None))[0] if dispatch_resolution is not None else (_num(x.dispatch_qty) if x.dispatch_qty is not None else None)
        if qty is not None and qty > 0:
            denoms[key] = max(denoms.get(key, 0.0), qty)
        elif _num(x.reject_qty) > 0:
            pending_keys.add(key)
    denominator = sum(denoms.values())
    pending = len(pending_keys)
    ppm = reject / denominator * 1_000_000 if denominator > 0 and pending == 0 else None
    return {
        "reject_qty": reject,
        "rework_qty": 0.0,
        "scrap_qty": 0.0,
        "denominator_qty": denominator,
        "ppm": ppm,
        "ppm_pending_rows": pending,
    }


def _aggregate_combined(daily: list[QualityRejectionDaily], history: list[QualityRejectionMonthlyHistory], dispatch_resolution: dict[int, tuple[float | None, str | None]] | None = None) -> dict:
    da = _aggregate_daily(daily)
    ha = _aggregate_history(history, dispatch_resolution)
    reject = da["reject_qty"] + ha["reject_qty"]
    denominator = da["denominator_qty"] + ha["denominator_qty"]
    pending = da["ppm_pending_rows"] + ha["ppm_pending_rows"]
    return {
        "reject_qty": reject,
        "rework_qty": da["rework_qty"] + ha["rework_qty"],
        "scrap_qty": da["scrap_qty"] + ha["scrap_qty"],
        "denominator_qty": denominator,
        "ppm": reject / denominator * 1_000_000 if denominator > 0 and pending == 0 else None,
        "ppm_pending_rows": pending,
        "daily_records": len(daily),
        "historical_records": len(history),
    }


def _serialize_history(row: QualityRejectionMonthlyHistory, products: dict, phenomena: dict,
                       dispatch_resolution: dict[int, tuple[float | None, str | None]] | None = None) -> dict:
    p = products.get(row.product_id)
    ph = phenomena.get(row.phenomenon_id)
    resolved_qty, resolved_source = (dispatch_resolution.get(row.id, (None, None)) if dispatch_resolution is not None else ((_num(row.dispatch_qty), "UPLOADED") if row.dispatch_qty is not None else (None, None)))
    resolved_ppm = (_num(row.reject_qty) / resolved_qty * 1_000_000) if resolved_qty is not None and resolved_qty > 0 else None
    return {
        "id": row.id,
        "month": row.month,
        "source": row.source,
        "source_sheet": row.source_sheet,
        "record_scope": row.record_scope,
        "include_in_aggregate": row.include_in_aggregate,
        "product_id": row.product_id,
        "product": p.name if p else None,
        "plant": row.plant or (p.plant if p else None),
        "product_group": p.product_group if p else None,
        "phenomenon_id": row.phenomenon_id,
        "phenomenon": ph.name if ph else None,
        "detection_operation": row.detection_operation,
        "responsible_team": row.responsible_team,
        "responsible_operation": row.responsible_operation,
        "machine": row.machine,
        "reject_qty": _num(row.reject_qty),
        "denominator_source": row.denominator_source,
        "dispatch_qty": resolved_qty,
        "dispatch_qty_source": resolved_source,
        "ppm": resolved_ppm,
        "source_reported_ppm": _num(row.source_reported_ppm) if row.source_reported_ppm is not None else None,
        "data_quality_note": row.data_quality_note,
    }


@router.get("/phenomena")
def phenomena(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    rows = db.scalars(select(QualityPhenomenon).where(QualityPhenomenon.is_active.is_(True)).order_by(QualityPhenomenon.name)).all()
    return [{"id": x.id, "code": x.code, "name": x.name, "group": x.phenomenon_group,
             "default_responsible_team": x.default_responsible_team, "criticality": x.criticality} for x in rows]


@router.post("/phenomena")
def create_phenomenon(payload: PhenomenonCreate, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    try:
        row = upsert_phenomenon(db, payload.name, group=payload.phenomenon_group,
                                default_team=payload.default_responsible_team, criticality=payload.criticality)
        db.commit(); db.refresh(row)
        return {"id": row.id, "code": row.code, "name": row.name, "group": row.phenomenon_group, "criticality": row.criticality}
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


@router.get("/phenomena-template")
def download_phenomenon_template(_: User = Depends(get_current_user)):
    wb = XLWorkbook()
    ws = wb.active
    ws.title = "Phenomenon_Master"
    instructions = wb.create_sheet("Instructions")
    headers = ["Phenomenon Name", "Phenomenon Group", "Responsible Team", "Criticality", "Active"]
    ws.append(headers)
    samples = [
        ["Porosity", "Casting Rejection", "Casting", "NORMAL", "Yes"],
        ["Blow Hole", "Casting Rejection", "Casting", "NORMAL", "Yes"],
    ]
    for row in samples:
        ws.append(row)
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:E501"
    for column, width in zip(("A","B","C","D","E"), (38,24,24,16,12)):
        ws.column_dimensions[column].width = width
    criticality = DataValidation(type="list", formula1='"LOW,NORMAL,HIGH,CRITICAL"', allow_blank=False)
    active = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
    ws.add_data_validation(criticality); criticality.add("D2:D501")
    ws.add_data_validation(active); active.add("E2:E501")
    instructions.append(["Phenomenon Master Upload"])
    instructions.append(["Required", "Phenomenon Name. Names are matched case-insensitively after normalization."])
    instructions.append(["Group", "Use groups such as Casting Rejection, Machining Rejection, Assembly Rejection or Supplier Defect."])
    instructions.append(["Responsible Team", "Default team responsible for the phenomenon, for example Casting or Operation."])
    instructions.append(["Criticality", "Select LOW, NORMAL, HIGH or CRITICAL."])
    instructions.append(["Active", "Yes keeps the phenomenon available; No deactivates an existing phenomenon."])
    instructions.append(["Import", "Upload on Quality / Rejection, review New/Updated/Unchanged/Rejected, then confirm."])
    instructions.append(["Duplicates", "Duplicate phenomenon names in one workbook are rejected. Existing names are safely updated, not duplicated."])
    instructions["A1"].font = Font(size=15, bold=True, color="1F4E78")
    for row in range(2, 9):
        instructions.cell(row, 1).font = Font(bold=True, color="1F4E78")
        instructions.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    instructions.column_dimensions["A"].width = 20
    instructions.column_dimensions["B"].width = 100
    out = BytesIO()
    wb.save(out); out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="Phenomenon_Master_Upload.xlsx"'})


@router.get("/template")
def download_daily_template(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    products = db.execute(select(Product, Customer).outerjoin(Customer, Customer.id == Product.customer_id).where(Product.is_active.is_(True)).order_by(Product.sort_order, Product.name)).all()
    operations = db.scalars(select(Operation).where(Operation.is_active.is_(True)).order_by(Operation.name)).all()
    machines = db.scalars(select(Machine).where(Machine.is_active.is_(True)).order_by(Machine.name)).all()
    phen = db.scalars(select(QualityPhenomenon).where(QualityPhenomenon.is_active.is_(True)).order_by(QualityPhenomenon.name)).all()
    if not products:
        raise HTTPException(409, "No active products are available. Add Product Master data before downloading the Daily Rejection template.")
    if not phen:
        raise HTTPException(409, "No active rejection phenomena are available. Add the approved Phenomenon Master first, then download the template again.")

    wb = XLWorkbook()
    ws = wb.active
    ws.title = "Daily_Rejection_Data"
    instructions = wb.create_sheet("Instructions")
    master = wb.create_sheet("Masters")
    headers = ["Date", "Shift", "Product", "Plant", "Customer", "Type", "Detection Process", "Responsible Process", "Machine",
               "Phenomenon", "Reject Qty", "Rework Qty", "Scrap Qty", "Remark", "Raise Action"]
    ws.append(headers)
    required_columns = {1, 2, 3, 7, 10, 11}
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="C65911" if cell.column in required_columns else "1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:O501"
    widths = [14,10,30,14,18,14,24,24,18,38,12,12,12,40,14]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64+i)].width = width

    master_headers = ["Product", "Plant", "Customer", "Type", "Process", "Machine", "Phenomenon", "Shift", "YesNo"]
    master.append(master_headers)
    max_rows = max(len(products), len(operations)+1, len(machines), len(phen), 4, 2)
    process_names = ["Disp_Done"] + [x.name for x in operations if _label_key(x.name) not in {"disp done", "dispatch done"}]
    for i in range(max_rows):
        p, c = products[i] if i < len(products) else (None, None)
        master.append([
            p.name if p else None,
            p.plant if p else None,
            c.name if c else None,
            p.product_group if p else None,
            process_names[i] if i < len(process_names) else None,
            machines[i].name if i < len(machines) else None,
            phen[i].name if i < len(phen) else None,
            ["A","B","C","General"][i] if i < 4 else None,
            ["Yes","No"][i] if i < 2 else None,
        ])

    product_last = max(2, len(products)+1)
    process_last = max(2, len(process_names)+1)
    machine_last = max(2, len(machines)+1)
    phen_last = max(2, len(phen)+1)
    shift_last = 5
    yesno_last = 3
    named_ranges = {
        "ProductList": f"'Masters'!$A$2:$A${product_last}",
        "ProcessList": f"'Masters'!$E$2:$E${process_last}",
        "MachineList": f"'Masters'!$F$2:$F${machine_last}",
        "PhenomenonList": f"'Masters'!$G$2:$G${phen_last}",
        "ShiftList": f"'Masters'!$H$2:$H${shift_last}",
        "YesNoList": f"'Masters'!$I$2:$I${yesno_last}",
    }
    for name, reference in named_ranges.items():
        wb.defined_names.add(DefinedName(name, attr_text=reference))

    def add_list_validation(cell_range: str, list_name: str, label: str, required: bool) -> None:
        dv = DataValidation(type="list", formula1=f"={list_name}", allow_blank=not required)
        dv.errorStyle = "stop"
        dv.errorTitle = f"Invalid {label}"
        dv.error = f"Select {label} from the dropdown. Typed values outside the approved master are not accepted."
        dv.promptTitle = label
        dv.prompt = f"Select an approved {label} from the dropdown."
        dv.showErrorMessage = True
        dv.showInputMessage = True
        ws.add_data_validation(dv)
        dv.add(cell_range)

    add_list_validation("B2:B501", "ShiftList", "Shift", True)
    add_list_validation("C2:C501", "ProductList", "Product", True)
    add_list_validation("G2:G501", "ProcessList", "Detection Process", True)
    add_list_validation("H2:H501", "ProcessList", "Responsible Process", False)
    add_list_validation("I2:I501", "MachineList", "Machine", False)
    add_list_validation("J2:J501", "PhenomenonList", "Phenomenon", True)
    add_list_validation("O2:O501", "YesNoList", "Raise Action", False)

    date_dv = DataValidation(type="date", operator="between", formula1="DATE(2020,1,1)", formula2="DATE(2100,12,31)", allow_blank=False)
    date_dv.errorStyle = "stop"; date_dv.errorTitle = "Date required"; date_dv.error = "Enter a valid rejection date."
    date_dv.showErrorMessage = True; ws.add_data_validation(date_dv); date_dv.add("A2:A501")
    reject_dv = DataValidation(type="decimal", operator="greaterThan", formula1="0", allow_blank=False)
    reject_dv.errorStyle = "stop"; reject_dv.errorTitle = "Reject Qty required"; reject_dv.error = "Reject Qty must be greater than zero."
    reject_dv.showErrorMessage = True; ws.add_data_validation(reject_dv); reject_dv.add("K2:K501")
    nonnegative_dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    nonnegative_dv.errorStyle = "stop"; nonnegative_dv.errorTitle = "Invalid quantity"; nonnegative_dv.error = "Quantity cannot be negative."
    nonnegative_dv.showErrorMessage = True; ws.add_data_validation(nonnegative_dv); nonnegative_dv.add("L2:M501")

    missing_fill = PatternFill("solid", fgColor="FCE4D6")
    for column in ("A", "B", "C", "G", "J", "K"):
        ws.conditional_formatting.add(
            f"{column}2:{column}501",
            FormulaRule(formula=[f'AND(COUNTA($A2:$O2)>3,{column}2="")'], fill=missing_fill),
        )
    for r in range(2, 502):
        ws.cell(r, 4).value = f'=IFERROR(VLOOKUP(C{r},Masters!$A$2:$D${product_last},2,FALSE),"")'
        ws.cell(r, 5).value = f'=IFERROR(VLOOKUP(C{r},Masters!$A$2:$D${product_last},3,FALSE),"")'
        ws.cell(r, 6).value = f'=IFERROR(VLOOKUP(C{r},Masters!$A$2:$D${product_last},4,FALSE),"")'
        for c in (4,5,6):
            ws.cell(r,c).fill = PatternFill("solid", fgColor="E7E6E6")
            ws.cell(r,c).protection = Protection(locked=True)
        for c in (1,2,3,7,8,9,10,11,12,13,14,15):
            ws.cell(r,c).protection = Protection(locked=False)
        for c in required_columns:
            if c not in (4,5,6):
                ws.cell(r,c).fill = PatternFill("solid", fgColor="FFF2CC")
        ws.cell(r,1).number_format = "dd-mmm-yyyy"
    instructions.append(["Daily Rejection Upload v0.4.6"])
    instructions.append(["Use", "Enter one rejection combination per row in Daily_Rejection_Data."])
    instructions.append(["Required", "Date, Shift, Product, Detection Process, Phenomenon and Reject Qty."])
    instructions.append(["Dropdowns", "Always select fixed master data from the dropdown. Do not type alternate spellings."])
    instructions.append(["Auto-filled", "Plant, Customer and Type come from Product Master and are protected."])
    instructions.append(["Quantities", "Reject Qty must be greater than zero. Rework and Scrap may be blank or zero."])
    instructions.append(["Before upload", "Use Excel Import Preview. Correct all rejected rows before confirmation."])
    instructions.append(["Master change", "Download a fresh template after Product, Process, Machine or Phenomenon masters change."])
    instructions["A1"].font = Font(size=15, bold=True, color="1F4E78")
    instructions["A2"].font = Font(bold=True, color="C65911")
    for r in range(2, 9):
        instructions.cell(r, 1).font = Font(bold=True, color="1F4E78")
        instructions.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    instructions.column_dimensions["A"].width = 18
    instructions.column_dimensions["B"].width = 95
    instructions.freeze_panes = "A2"
    instructions.sheet_view.showGridLines = False
    ws.protection.sheet = True
    ws.protection.password = "PRWUpload"
    ws.protection.autoFilter = False
    ws.protection.sort = False
    master["K1"] = "Template Version"
    master["K2"] = "v0.4.6"
    master.sheet_state = "hidden"
    bio = BytesIO(); wb.save(bio); bio.seek(0)
    return StreamingResponse(bio, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="Daily_Rejection_Upload_v0.4.6.xlsx"'})


@router.post("/import-daily")
def import_daily(_: User = Depends(get_current_user)):
    raise HTTPException(409, "Use /api/import/preview/quality-daily then confirm")


@router.post("/import-history")
def import_history(_: User = Depends(get_current_user)):
    raise HTTPException(409, "Use /api/import/preview/quality-history then confirm")


@router.get("/imports")
def quality_imports(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    rows = db.scalars(select(QualityRejectionImportBatch).order_by(QualityRejectionImportBatch.created_at.desc()).limit(100)).all()
    return [{"id": x.id, "file_name": x.file_name, "sha256": x.file_sha256, "type": x.import_type,
             "status": x.status, "imported_at": x.created_at, "stats": json.loads(x.stats_json) if x.stats_json else {}} for x in rows]


@router.post("/daily")
def create_daily_manual(payload: QualityRejectionManualCreate, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    shift_map = {"a": "A", "b": "B", "c": "C", "general": "General"}
    shift = shift_map.get(str(payload.shift or "").strip().lower())
    if not shift:
        raise HTTPException(400, "Shift must be A, B, C or General")
    if payload.reject_qty <= 0:
        raise HTTPException(400, "Reject Qty must be greater than zero")
    if payload.rework_qty < 0 or payload.scrap_qty < 0:
        raise HTTPException(400, "Rework Qty and Scrap Qty cannot be negative")

    product = db.get(Product, payload.product_id)
    if not product or not product.is_active:
        raise HTTPException(400, "Product is not available in the active Product Master")
    phenomenon = db.get(QualityPhenomenon, payload.phenomenon_id)
    if not phenomenon or not phenomenon.is_active:
        raise HTTPException(400, "Phenomenon is not available in the active Quality Phenomenon Master")

    route = _route_for_date(db, product.id, payload.rejection_date)
    if not route:
        raise HTTPException(400, f"No route is available for {product.name} on {payload.rejection_date}")

    detection_ro = db.get(RouteOperation, payload.detection_route_operation_id)
    if not detection_ro or detection_ro.route_version_id != route.id:
        raise HTTPException(400, "Detection Process is not in the active route for the selected Product and Date")

    responsible_ro = None
    if payload.responsible_route_operation_id is not None:
        responsible_ro = db.get(RouteOperation, payload.responsible_route_operation_id)
        if not responsible_ro or responsible_ro.route_version_id != route.id:
            raise HTTPException(400, "Responsible Process is not in the active route for the selected Product and Date")

    machine = None
    if payload.machine_id is not None:
        machine = db.get(Machine, payload.machine_id)
        if not machine or not machine.is_active:
            raise HTTPException(400, "Machine is not available in the active Machine Master")

    record_key = _daily_record_key(
        payload.rejection_date, shift, product.id, detection_ro.id,
        responsible_ro.id if responsible_ro else None, machine.id if machine else None, phenomenon.id,
    )
    existing = db.scalar(select(QualityRejectionDaily).where(QualityRejectionDaily.record_key == record_key))
    if existing:
        raise HTTPException(
            409,
            "A rejection already exists for the same Date, Shift, Product, Detection Process, Responsible Process, Machine and Phenomenon. Use the existing record/import workflow for corrections.",
        )

    denominator_source, denominator_qty = denominator_for_rejection(
        db, on_date=payload.rejection_date, shift=shift, product=product,
        detection_ro=detection_ro, machine=machine,
    )
    ppm = None
    if denominator_qty is not None and denominator_qty > 0:
        ppm = (payload.reject_qty / denominator_qty * Decimal("1000000")).quantize(Decimal("0.001"))

    row = QualityRejectionDaily(
        record_key=record_key,
        rejection_date=payload.rejection_date,
        shift=shift,
        product_id=product.id,
        plant=product.plant,
        detection_route_operation_id=detection_ro.id,
        responsible_route_operation_id=responsible_ro.id if responsible_ro else None,
        responsible_team=phenomenon.default_responsible_team or "Operation",
        machine_id=machine.id if machine else None,
        phenomenon_id=phenomenon.id,
        reject_qty=payload.reject_qty,
        rework_qty=payload.rework_qty,
        scrap_qty=payload.scrap_qty,
        denominator_source=denominator_source,
        denominator_qty=denominator_qty,
        ppm=ppm,
        remark=payload.remark.strip() if payload.remark and payload.remark.strip() else None,
        source="MANUAL",
        action_required=payload.action_required,
        entered_by_id=user.id,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    context = _daily_context(db, [row])
    result = _serialize_daily(row, context)
    result["created"] = True
    result["ppm_pending"] = ppm is None
    return result


@router.get("/daily")
def list_daily(from_date: date, to_date: date, plant: str | None = None, product_id: str | None = None,
               phenomenon_id: str | None = None, operation_id: str | None = None, machine_id: str | None = None,
               shift: str | None = None, customer_id: str | None = None, product_group: str | None = None,
               db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    rows = _daily_rows(db, from_date=from_date, to_date=to_date, plant=plant, product_id=product_id,
                       phenomenon_id=phenomenon_id, operation_id=operation_id, machine_id=machine_id, shift=shift,
                       customer_id=customer_id, product_group=product_group)
    context = _daily_context(db, rows)
    return [_serialize_daily(x, context) for x in rows]


@router.get("/dashboard")
def dashboard(from_date: date, to_date: date, plant: str | None = None, product_id: str | None = None,
              phenomenon_id: str | None = None, operation_id: str | None = None, machine_id: str | None = None,
              shift: str | None = None, customer_id: str | None = None, product_group: str | None = None,
              db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    daily = _daily_rows(db, from_date=from_date, to_date=to_date, plant=plant, product_id=product_id,
                       phenomenon_id=phenomenon_id, operation_id=operation_id, machine_id=machine_id, shift=shift,
                       customer_id=customer_id, product_group=product_group)

    # Monthly history has no dependable shift / route-operation / machine granularity.
    # Keep it visible for normal product/plant/phenomenon reporting, but do not mix it into
    # filters that it cannot faithfully satisfy.
    history_supported = not csv_ints(operation_id) and not csv_ints(machine_id) and not csv_strings(shift)
    history: list[QualityRejectionMonthlyHistory] = []
    if history_supported:
        history = _historical_rows(db, from_date=from_date, to_date=to_date, plant=plant, product_id=product_id,
                                   phenomenon_id=phenomenon_id, customer_id=customer_id, product_group=product_group)
        history = _exclude_history_covered_by_daily(history, daily)
        daily = _exclude_daily_covered_by_history(daily, history)

    dispatch_resolution = _history_dispatch_resolution(db, history, plant_filter=plant) if history else {}
    return _dashboard_data(db, daily, history, dispatch_resolution, history_supported)


def _dashboard_data(db, daily, history, dispatch_resolution, history_supported):
    agg = _aggregate_combined(daily, history, dispatch_resolution)
    daily_links = list(db.scalars(select(QualityActionLink).where(
        QualityActionLink.daily_rejection_id.in_([x.id for x in daily] or [-1])
    )).all())
    history_links = list(db.scalars(select(QualityActionLink).where(
        QualityActionLink.monthly_history_id.in_([x.id for x in history] or [-1])
    )).all())
    actions_by_id = _by_id(db, Action, {x.action_id for x in [*daily_links, *history_links]})
    actions = [actions_by_id.get(x.action_id) for x in [*daily_links, *history_links]]
    agg.update({
        "records": len(daily) + len(history),
        "open_actions": sum(1 for a in actions if a and a.status != ActionStatus.CLOSED),
        "overdue_actions": sum(1 for a in actions if a and a.status != ActionStatus.CLOSED and a.due_at and a.due_at < datetime.utcnow()),
        "history_included": history_supported,
        "history_note": None if history_supported else "Historical monthly data is excluded while Shift, Process or Machine filters are applied.",
    })
    return agg


@router.get("/pareto")
def pareto(from_date: date, to_date: date, group_by: str = "phenomenon", plant: str | None = None,
           product_id: str | None = None, phenomenon_id: str | None = None, operation_id: str | None = None,
           machine_id: str | None = None, shift: str | None = None, customer_id: str | None = None,
           product_group: str | None = None, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    daily = _daily_rows(db, from_date=from_date, to_date=to_date, plant=plant, product_id=product_id,
                       phenomenon_id=phenomenon_id, operation_id=operation_id, machine_id=machine_id, shift=shift,
                       customer_id=customer_id, product_group=product_group)
    history_supported = not csv_ints(operation_id) and not csv_ints(machine_id) and not csv_strings(shift)
    history: list[QualityRejectionMonthlyHistory] = []
    if history_supported:
        history = _historical_rows(db, from_date=from_date, to_date=to_date, plant=plant, product_id=product_id,
                                   phenomenon_id=phenomenon_id, customer_id=customer_id, product_group=product_group)
        history = _exclude_history_covered_by_daily(history, daily)
        daily = _exclude_daily_covered_by_history(daily, history)

    dispatch_resolution = _history_dispatch_resolution(db, history, plant_filter=plant) if history else {}

    return _pareto_data(db, daily, history, dispatch_resolution, group_by)


def _pareto_data(db, daily, history, dispatch_resolution, group_by, daily_context=None, products=None, phenomena=None):
    valid = {"product", "plant", "phenomenon", "operation", "machine"}
    if group_by not in valid:
        raise HTTPException(400, "group_by must be product, plant, phenomenon, operation or machine")

    daily_context = daily_context or _daily_context(db, daily)
    products = products or _by_id(db, Product, {x.product_id for x in history})
    phenomena = phenomena or _by_id(db, QualityPhenomenon, {x.phenomenon_id for x in history})
    daily_groups: dict[str, list[QualityRejectionDaily]] = defaultdict(list)
    history_groups: dict[str, list[QualityRejectionMonthlyHistory]] = defaultdict(list)

    for x in daily:
        p = daily_context['products'].get(x.product_id)
        ph = daily_context['phenomena'].get(x.phenomenon_id)
        m = daily_context['machines'].get(x.machine_id)
        ro = daily_context['routes'].get(x.detection_route_operation_id)
        op = daily_context['operations'].get(ro.operation_id) if ro else None
        key = {
            "product": p.name if p else "Unknown",
            "plant": x.plant or (p.plant if p else None) or "Unassigned",
            "phenomenon": ph.name if ph else "Unknown",
            "operation": op.name if op else "Unassigned",
            "machine": m.name if m else "Unassigned",
        }[group_by]
        daily_groups[key].append(x)

    for x in history:
        p = products.get(x.product_id)
        ph = phenomena.get(x.phenomenon_id)
        key = {
            "product": p.name if p else "Unknown",
            "plant": x.plant or (p.plant if p else None) or "Unassigned",
            "phenomenon": ph.name if ph else "Unknown",
            "operation": x.detection_operation or "Unassigned",
            "machine": x.machine or "Unassigned",
        }[group_by]
        history_groups[key].append(x)

    out=[]
    for name in sorted(set(daily_groups) | set(history_groups)):
        da = _aggregate_daily(daily_groups.get(name, []))
        ha = _aggregate_history(history_groups.get(name, []), dispatch_resolution)
        reject = da["reject_qty"] + ha["reject_qty"]
        denominator = da["denominator_qty"] + ha["denominator_qty"]
        pending = da["ppm_pending_rows"] + ha["ppm_pending_rows"]
        out.append({
            "name": name,
            "value": reject,
            "reject_qty": reject,
            "ppm": reject / denominator * 1_000_000 if denominator > 0 and pending == 0 else None,
            "denominator_qty": denominator,
            "ppm_pending_rows": pending,
            "historical_records": len(history_groups.get(name, [])),
        })
    out.sort(key=lambda x: x["reject_qty"], reverse=True)
    return out


@router.get("/trend")
def trend(from_date: date, to_date: date, plant: str | None = None, product_id: str | None = None,
          phenomenon_id: str | None = None, operation_id: str | None = None, machine_id: str | None = None,
          shift: str | None = None, customer_id: str | None = None, product_group: str | None = None,
          db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    rows = _daily_rows(db, from_date=from_date, to_date=to_date, plant=plant, product_id=product_id,
                       phenomenon_id=phenomenon_id, operation_id=operation_id, machine_id=machine_id, shift=shift,
                       customer_id=customer_id, product_group=product_group)
    return _trend_data(rows, from_date, to_date)


def _trend_data(rows, from_date, to_date):
    grouped: dict[date, list[QualityRejectionDaily]] = defaultdict(list)
    d=from_date
    while d<=to_date:
        grouped[d]=[]; d+=timedelta(days=1)
    for x in rows: grouped[x.rejection_date].append(x)
    out=[]
    for d, grp in sorted(grouped.items()):
        a=_aggregate_daily(grp)
        out.append({"date": d, "label": d.strftime("%d-%b"), **a})
    return out


@router.get('/report-pack')
def report_pack(from_date: date, to_date: date, plant: str | None = None, product_id: str | None = None,
                phenomenon_id: str | None = None, operation_id: str | None = None,
                machine_id: str | None = None, shift: str | None = None,
                customer_id: str | None = None, product_group: str | None = None,
                db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    daily = _daily_rows(db, from_date=from_date, to_date=to_date, plant=plant, product_id=product_id,
                        phenomenon_id=phenomenon_id, operation_id=operation_id, machine_id=machine_id,
                        shift=shift, customer_id=customer_id, product_group=product_group)
    history = _historical_rows(db, from_date=from_date, to_date=to_date, plant=plant,
                               product_id=product_id, phenomenon_id=phenomenon_id,
                               customer_id=customer_id, product_group=product_group, include_detail=True)
    supported = not csv_ints(operation_id) and not csv_ints(machine_id) and not csv_strings(shift)
    aggregate_candidates = [x for x in history if x.include_in_aggregate]
    aggregate_history = _exclude_history_covered_by_daily(aggregate_candidates, daily) if supported else []
    aggregate_daily = _exclude_daily_covered_by_history(daily, aggregate_history) if supported else daily
    resolved = _history_dispatch_resolution(db, history, plant_filter=plant) if history else {}
    context = _daily_context(db, daily)
    products = _by_id(db, Product, {x.product_id for x in history})
    phenomena = _by_id(db, QualityPhenomenon, {x.phenomenon_id for x in history})
    month_first = from_date.replace(day=1)
    month_last = to_date.replace(day=1)
    return {
        'summary': _dashboard_data(db, aggregate_daily, aggregate_history, resolved, supported),
        'rows': [_serialize_daily(x, context) for x in daily],
        'trend': _trend_data(daily, from_date, to_date),
        'part_pareto': _pareto_data(db, aggregate_daily, aggregate_history, resolved, 'product', context, products, phenomena),
        'phenomenon_pareto': _pareto_data(db, aggregate_daily, aggregate_history, resolved, 'phenomenon', context, products, phenomena),
        'monthly_trend': monthly_trend(from_month=month_first, to_month=month_last, plant=plant,
                                      product_id=product_id, phenomenon_id=phenomenon_id,
                                      customer_id=customer_id, product_group=product_group, db=db, _=user),
        'history': [_serialize_history(x, products, phenomena, resolved) for x in history],
    }


@router.get("/history-range")
def history_range(db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    first, last, count = db.execute(select(func.min(QualityRejectionMonthlyHistory.month),
                                           func.max(QualityRejectionMonthlyHistory.month),
                                           func.count(QualityRejectionMonthlyHistory.id))).one()
    return {"min_month": first, "max_month": last, "records": count}


@router.get("/history")
def list_history(from_date: date, to_date: date, plant: str | None = None, product_id: str | None = None,
                 phenomenon_id: str | None = None, customer_id: str | None = None, product_group: str | None = None,
                 db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    rows = _historical_rows(db, from_date=from_date, to_date=to_date, plant=plant, product_id=product_id,
                            phenomenon_id=phenomenon_id, customer_id=customer_id, product_group=product_group,
                            include_detail=True)
    dispatch_resolution = _history_dispatch_resolution(db, rows, plant_filter=plant) if rows else {}
    products = _by_id(db, Product, {x.product_id for x in rows})
    phenomena = _by_id(db, QualityPhenomenon, {x.phenomenon_id for x in rows})
    return [_serialize_history(x, products, phenomena, dispatch_resolution) for x in rows]


@router.get("/monthly-trend")
def monthly_trend(from_month: date, to_month: date, plant: str | None = None, product_id: str | None = None,
                  phenomenon_id: str | None = None, customer_id: str | None = None, product_group: str | None = None,
                  db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    first=from_month.replace(day=1); last=to_month.replace(day=1)
    daily_to=(last.replace(day=28)+timedelta(days=4)).replace(day=1)-timedelta(days=1)
    daily = _daily_rows(db, from_date=first, to_date=daily_to, plant=plant, product_id=product_id,
                        phenomenon_id=phenomenon_id, customer_id=customer_id, product_group=product_group)
    daily_by_month: dict[date,list[QualityRejectionDaily]]=defaultdict(list)
    for x in daily: daily_by_month[x.rejection_date.replace(day=1)].append(x)

    q=select(QualityRejectionMonthlyHistory).join(Product, Product.id==QualityRejectionMonthlyHistory.product_id).where(
        QualityRejectionMonthlyHistory.month>=first, QualityRejectionMonthlyHistory.month<=last)
    plants=csv_strings(plant); product_ids=csv_ints(product_id); phenomenon_ids=csv_ints(phenomenon_id); customer_ids=csv_ints(customer_id); groups=csv_strings(product_group)
    if plants:
        q=q.where(QualityRejectionMonthlyHistory.plant.in_(plants))
    else:
        q=q.where(QualityRejectionMonthlyHistory.include_in_aggregate.is_(True))
    if product_ids: q=q.where(QualityRejectionMonthlyHistory.product_id.in_(product_ids))
    if phenomenon_ids: q=q.where(QualityRejectionMonthlyHistory.phenomenon_id.in_(phenomenon_ids))
    if customer_ids: q=q.where(Product.customer_id.in_(customer_ids))
    if groups: q=q.where(Product.product_group.in_(groups))
    hist=list(db.scalars(q).all())
    dispatch_resolution = _history_dispatch_resolution(db, hist, plant_filter=plant) if hist else {}
    hist_by_month: dict[date,list[QualityRejectionMonthlyHistory]]=defaultdict(list)
    for x in hist: hist_by_month[x.month.replace(day=1)].append(x)

    out=[]; m=first
    while m<=last:
        rows=hist_by_month.get(m,[])
        if rows:
            reject=sum(_num(x.reject_qty) for x in rows)
            denoms={}
            pending=0
            for x in rows:
                # Same dispatch denominator repeats for each phenomenon; count it once per product/source/scope/month.
                key=(x.product_id,x.source,x.source_sheet,x.record_scope,x.plant)
                qty = dispatch_resolution.get(x.id, (None, None))[0]
                if qty is not None and qty>0: denoms[key]=max(denoms.get(key,0),qty)
                elif _num(x.reject_qty)>0: pending+=1
            denominator=sum(denoms.values()); ppm=reject/denominator*1_000_000 if denominator>0 and pending==0 else None
            a={"reject_qty":reject,"rework_qty":0,"scrap_qty":0,"denominator_qty":denominator,"ppm":ppm,"ppm_pending_rows":pending}; source="HISTORICAL"
        elif daily_by_month.get(m):
            a=_aggregate_daily(daily_by_month[m]); source="DAILY_ROLLUP"
        else:
            a={"reject_qty":0.0,"rework_qty":0.0,"scrap_qty":0.0,"denominator_qty":0.0,"ppm":None,"ppm_pending_rows":0}; source="NONE"
        out.append({"month":m,"label":m.strftime("%b-%Y"),"source":source,**a})
        m=(m.replace(day=28)+timedelta(days=4)).replace(day=1)
    return out


@router.post("/rejections/{rejection_id}/raise-action")
def raise_action(rejection_id: int, payload: QualityActionCreate, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    rej=db.get(QualityRejectionDaily,rejection_id)
    if not rej: raise HTTPException(404,"Rejection record not found")
    existing=db.scalar(select(QualityActionLink).where(QualityActionLink.daily_rejection_id==rejection_id).order_by(QualityActionLink.id.desc()).limit(1))
    if existing:
        a=db.get(Action,existing.action_id)
        return {"status":"already_linked","action_id":a.id,"action_no":a.action_no}
    p=db.get(Product,rej.product_id); ph=db.get(QualityPhenomenon,rej.phenomenon_id)
    count=db.query(Action).count()+1
    action_no=f"ACT-{rej.rejection_date.year}-{count:05d}"
    ppm_text=f"{_num(rej.ppm):,.0f} PPM" if rej.ppm is not None else "PPM pending"
    problem=f"{ph.name if ph else 'Rejection'} - {_num(rej.reject_qty):g} Nos. ({ppm_text})"
    a=Action(action_no=action_no, reference_date=rej.rejection_date, problem_category="Quality / Rejection",
             problem_description=problem, action_description=payload.action_description, owner_id=payload.owner_id,
             due_at=payload.due_at, priority=payload.priority, kpi_type="PPM",
             kpi_value_when_raised=rej.ppm, gap_when_raised=rej.reject_qty)
    db.add(a);db.flush()
    db.add(ActionContext(action_id=a.id,context_date=rej.rejection_date,product_id=rej.product_id,
                         route_operation_id=rej.detection_route_operation_id,machine_id=rej.machine_id))
    db.add(ActionWhyWhy(action_id=a.id,containment_action=payload.action_description,updated_by_id=user.id))
    db.add(ActionHistory(action_id=a.id,changed_by_id=user.id,new_status=a.status,comment="Action raised from Quality / Rejection; standard Why-Why plan opened"))
    db.add(QualityActionLink(action_id=a.id,daily_rejection_id=rej.id))
    active=db.scalar(select(ReviewSession).where(ReviewSession.review_date==rej.rejection_date,ReviewSession.ended_at.is_(None)).order_by(ReviewSession.id.desc()).limit(1))
    if active: db.add(ReviewActionLink(review_session_id=active.id,action_id=a.id))
    db.commit();db.refresh(a)
    return {"status":"created","action_id":a.id,"action_no":a.action_no}
