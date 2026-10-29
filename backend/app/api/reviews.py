from ..time_utils import ist_today, ist_date, UTCTimestamp, format_ist
from datetime import date, datetime, timedelta
from io import BytesIO
import re

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from pydantic import BaseModel, Field
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..auth import get_current_user
from ..db import get_db
from ..enums import ActionStatus, Priority
from ..models import Action, ActionContext, ActionHistory, ActionWhyWhy, Product, ReviewActionLink, ReviewMeetingRequest, ReviewParticipantPresence, ReviewPoint, ReviewSession, User, UserSession
from ..schemas import ReviewClose, ReviewStart

router = APIRouter(prefix="/reviews", tags=["reviews"])


class ReviewPointCreate(BaseModel):
    category: str | None = Field(default=None, max_length=120)
    discussion_point: str = Field(min_length=2, max_length=5000)
    action_required: bool = False
    action_description: str | None = Field(default=None, max_length=5000)
    owner_id: int | None = None
    due_at: UTCTimestamp | None = None
    priority: Priority = Priority.MEDIUM
    product_id: int | None = None


class ReviewUpdate(BaseModel):
    participants: str | None = Field(default=None, max_length=5000)
    comments: str | None = Field(default=None, max_length=5000)


def serialize_review(row: ReviewSession, db: Session | None = None) -> dict:
    result = {"id": row.id, "review_date": row.review_date, "started_at": row.started_at,
              "ended_at": row.ended_at, "participants": row.participants,
              "total_sales_gap": float(row.total_sales_gap) if row.total_sales_gap is not None else None,
              "comments": row.comments, "owner_id": row.owner_id,
              "owner_name": (db.get(User, row.owner_id).full_name if db is not None and row.owner_id and db.get(User, row.owner_id) else None),
              "auto_close_reason": row.auto_close_reason}
    if db is not None:
        action_ids = select(ReviewPoint.action_id).where(
            ReviewPoint.review_session_id == row.id,
            ReviewPoint.action_id.is_not(None),
        ).union(
            select(ReviewActionLink.action_id).where(ReviewActionLink.review_session_id == row.id)
        )
        discussion_point_count = db.scalar(
            select(func.count(ReviewPoint.id)).where(ReviewPoint.review_session_id == row.id)
        ) or 0
        # Include actions linked to the MOM from Dashboard, Quality, OEE or
        # the Actions page. Actions already represented by a typed ReviewPoint
        # are excluded here so an actionable discussion point is counted once.
        linked_action_point_count = db.scalar(
            select(func.count(func.distinct(ReviewActionLink.action_id))).where(
                ReviewActionLink.review_session_id == row.id,
                ReviewActionLink.action_id.not_in(
                    select(ReviewPoint.action_id).where(
                        ReviewPoint.review_session_id == row.id,
                        ReviewPoint.action_id.is_not(None),
                    )
                ),
            )
        ) or 0
        result["total_point_count"] = discussion_point_count + linked_action_point_count
        result["action_count"] = db.scalar(select(func.count(Action.id)).where(Action.id.in_(action_ids))) or 0
        result["closed_action_count"] = db.scalar(select(func.count(Action.id)).where(
            Action.id.in_(action_ids), Action.status == ActionStatus.CLOSED
        )) or 0
        result["overdue_action_count"] = db.scalar(select(func.count(Action.id)).where(
            Action.id.in_(action_ids), Action.status != ActionStatus.CLOSED,
            Action.due_at.is_not(None), Action.due_at < datetime.utcnow()
        )) or 0
        overdue_dates = db.scalars(select(Action.due_at).where(
            Action.id.in_(action_ids), Action.status != ActionStatus.CLOSED,
            Action.due_at.is_not(None), Action.due_at < datetime.utcnow()
        )).all()
        result["max_overdue_days"] = max(
            (max(0, (ist_today() - ist_date(due_at)).days) for due_at in overdue_dates),
            default=0,
        )
    return result


def action_data(db: Session, a: Action | None, followup: dict | None = None) -> dict | None:
    if not a:
        return None
    owner = db.get(User, a.owner_id) if a.owner_id else None
    product_names = db.scalars(
        select(Product.name).join(ActionContext, ActionContext.product_id == Product.id)
        .where(ActionContext.action_id == a.id).distinct().order_by(Product.name)
    ).all()
    return {"id": a.id, "action_no": a.action_no, "description": a.action_description,
            "owner_id": a.owner_id, "owner": owner.full_name if owner else None,
            "products": product_names,
            "due_at": a.due_at, "priority": a.priority.value, "status": a.status.value,
            "closure_remark": a.closure_remark, "closed_at": a.closed_at,
            "age_days": max(0, (ist_today() - a.reference_date).days),
            "overdue_days": max(0, (ist_today() - ist_date(a.due_at)).days) if a.due_at and a.status != ActionStatus.CLOSED else 0,
            "last_followup_at": (followup or {}).get("at"),
            "last_followup_comment": (followup or {}).get("comment"),
            "requires_whywhy": a.requires_whywhy}


def point_data(db: Session, p: ReviewPoint) -> dict:
    product = db.get(Product, p.product_id) if p.product_id else None
    creator = db.get(User, p.created_by_id) if p.created_by_id else None
    return {"id": p.id, "sequence_no": p.sequence_no, "category": p.category,
            "discussion_point": p.discussion_point, "action_required": p.action_required,
            "created_by": creator.full_name if creator else None,
            "product_id": p.product_id, "product_code": product.code if product else None,
            "product": product.name if product else None,
            "action": action_data(db, db.get(Action, p.action_id) if p.action_id else None)}


@router.get("/active")
def get_active_review(review_date: date, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    row = db.scalar(select(ReviewSession).where(ReviewSession.review_date == review_date, ReviewSession.ended_at.is_(None),
                    ReviewSession.owner_id == user.id)
                    .order_by(ReviewSession.started_at.desc(), ReviewSession.id.desc()).limit(1))
    return serialize_review(row, db) if row else None


@router.get("")
def list_reviews(participant: str | None = None, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    stmt = select(ReviewSession)
    if participant:
        stmt = stmt.where(ReviewSession.participants.ilike(f"%{participant.strip()}%"))
    rows = db.scalars(stmt.order_by(ReviewSession.review_date.desc(), ReviewSession.id.desc()).limit(200)).all()
    return [serialize_review(x, db) for x in rows]


@router.get("/{review_id}")
def get_review(review_id: int, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    row = db.get(ReviewSession, review_id)
    if not row:
        raise HTTPException(404, "Review session not found")
    return serialize_review(row, db)


@router.patch("/{review_id}")
def update_review(review_id: int, payload: ReviewUpdate, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    row = db.get(ReviewSession, review_id)
    if not row:
        raise HTTPException(404, "Review session not found")
    if row.ended_at is not None:
        raise HTTPException(409, "Closed meetings cannot be edited; reopen the meeting first")
    for field, value in payload.model_dump(exclude_unset=True).items():
        if isinstance(value, str):
            value = value.strip() or None
        setattr(row, field, value)
    db.commit()
    db.refresh(row)
    return serialize_review(row, db)


@router.patch("/{review_id}/reopen")
def reopen_review(review_id: int, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    row = db.get(ReviewSession, review_id)
    if not row:
        raise HTTPException(404, "Review session not found")
    # Multiple independent MOMs may be active for the same date. Reopening
    # one historical meeting must not be blocked by another meeting's state.
    row.ended_at = None
    row.auto_close_reason = None
    db.commit()
    db.refresh(row)
    return serialize_review(row, db)


@router.post("")
def start_review(payload: ReviewStart, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    # Each explicit Start action creates an independent MOM, even when another
    # meeting for the same date is already active. Collaboration happens inside
    # a selected meeting; it must not silently merge separate meetings.
    row = ReviewSession(review_date=payload.review_date, started_at=datetime.utcnow(), owner_id=user.id,
                        participants=payload.participants, comments=payload.comments)
    db.add(row)
    db.commit()
    db.refresh(row)
    result = serialize_review(row, db)
    result["existing"] = False
    return result


@router.patch("/{review_id}/close")
def close_review(review_id: int, payload: ReviewClose, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    row = db.get(ReviewSession, review_id)
    if not row: raise HTTPException(404, "Review session not found")

    point_count = db.scalar(
        select(func.count(ReviewPoint.id)).where(ReviewPoint.review_session_id == review_id)
    ) or 0
    linked_action_count = db.scalar(
        select(func.count(ReviewActionLink.id)).where(ReviewActionLink.review_session_id == review_id)
    ) or 0
    if point_count == 0 and linked_action_count == 0:
        db.delete(row)
        db.commit()
        return {"id": review_id, "discarded": True, "message": "Meeting closed with no discussion points and was not saved."}

    if row.ended_at is None: row.ended_at = datetime.utcnow()
    row.auto_close_reason = None
    if payload.comments is not None: row.comments = payload.comments
    db.commit(); db.refresh(row)
    result = serialize_review(row, db)
    result["discarded"] = False
    return result


@router.post("/{review_id}/points")
def add_point(review_id: int, payload: ReviewPointCreate, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    # Lock the parent MOM row for the duration of this transaction. Concurrent
    # contributors to the same meeting are serialized while allocating the next
    # sequence number, preventing duplicate point numbers without locking other
    # meetings or blocking independent MOMs.
    review = db.scalar(select(ReviewSession).where(ReviewSession.id == review_id).with_for_update())
    if not review: raise HTTPException(404, "Review session not found")
    if review.ended_at is not None: raise HTTPException(409, "Closed meetings cannot be edited")
    action = None
    product = None
    if payload.product_id is not None:
        product = db.get(Product, payload.product_id)
        if not product or not product.is_active:
            raise HTTPException(422, "Selected product is not available in the active Product Master")

    if payload.action_required:
        if not (payload.action_description or "").strip():
            raise HTTPException(422, "Action description is required when Action Required is Yes")
        if payload.owner_id is None:
            raise HTTPException(422, "Select a responsible owner before assigning this action")
        owner = db.get(User, payload.owner_id)
        if not owner or not owner.is_active:
            raise HTTPException(422, "Selected action owner is not active")
        if payload.due_at is None:
            raise HTTPException(422, "Enter a target date and time before assigning this action")
        count = db.query(Action).count() + 1
        action = Action(action_no=f"ACT-{review.review_date.year}-{count:05d}", reference_date=review.review_date,
                        problem_category=payload.category or "Daily Review", problem_description=payload.discussion_point.strip(),
                        action_description=payload.action_description.strip(), owner_id=payload.owner_id,
                        due_at=payload.due_at, priority=payload.priority, action_type="DAILY_REVIEW", requires_whywhy=False)
        db.add(action); db.flush()
        db.add(ActionHistory(action_id=action.id, changed_by_id=user.id, new_status=action.status,
                             comment=f"Created from Daily Review #{review.id}; simple follow-up action"))
        db.add(ReviewActionLink(review_session_id=review.id, action_id=action.id))
        if product:
            db.add(ActionContext(action_id=action.id, context_date=review.review_date, product_id=product.id))
    seq = (db.scalar(select(func.max(ReviewPoint.sequence_no)).where(ReviewPoint.review_session_id == review.id)) or 0) + 1
    point = ReviewPoint(review_session_id=review.id, sequence_no=seq, category=payload.category,
                        discussion_point=payload.discussion_point.strip(), action_required=payload.action_required,
                        product_id=product.id if product else None,
                        action_id=action.id if action else None, created_by_id=user.id)
    db.add(point); db.commit(); db.refresh(point)
    return point_data(db, point)


@router.post("/{review_id}/presence")
def update_meeting_presence(review_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    meeting = db.get(ReviewSession, review_id)
    if not meeting or meeting.ended_at is not None:
        raise HTTPException(409, "Meeting is not active")
    now = datetime.utcnow()
    row = db.scalar(select(ReviewParticipantPresence).where(ReviewParticipantPresence.review_session_id == review_id,
        ReviewParticipantPresence.user_id == user.id).with_for_update())
    if row:
        row.last_seen_at = now
    else:
        db.add(ReviewParticipantPresence(review_session_id=review_id, user_id=user.id, last_seen_at=now))
    db.commit()
    return {"ok": True, "last_seen_at": now}


@router.get("/collaboration/active")
def active_collaboration(review_date: date, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    rows = db.scalars(select(ReviewSession).where(ReviewSession.review_date == review_date, ReviewSession.ended_at.is_(None)).order_by(ReviewSession.started_at.desc(), ReviewSession.id.desc())).all()
    now = datetime.utcnow()
    result = []
    for row in rows:
        owner = db.get(User, row.owner_id) if row.owner_id else None
        recent = False
        if owner:
            recent = bool(db.scalar(select(UserSession.id).where(UserSession.user_id == owner.id, UserSession.revoked_at.is_(None), UserSession.expires_at > now, UserSession.last_seen_at > now - timedelta(minutes=30)).limit(1)))
        presence_rows = db.scalars(select(ReviewParticipantPresence).where(ReviewParticipantPresence.review_session_id == row.id).order_by(ReviewParticipantPresence.last_seen_at.desc())).all()
        participant_presence = []
        for presence in presence_rows:
            participant = db.get(User, presence.user_id)
            if participant:
                participant_presence.append({"user_id": participant.id, "name": participant.full_name,
                    "last_seen_at": presence.last_seen_at, "online": presence.last_seen_at > now - timedelta(minutes=2)})
        result.append({**serialize_review(row, db), "owner_name": owner.full_name if owner else None, "owner_online": recent,
                       "participants_presence": participant_presence,
                       "is_owner": bool(row.owner_id == user.id),
                       "pending_request": bool(db.scalar(select(ReviewMeetingRequest.id).where(ReviewMeetingRequest.review_session_id == row.id, ReviewMeetingRequest.requester_id == user.id, ReviewMeetingRequest.status == "PENDING").limit(1)))})
    return result


@router.post("/{review_id}/request-owner")
def request_meeting_owner(review_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    row = db.get(ReviewSession, review_id)
    if not row or row.ended_at is not None:
        raise HTTPException(404, "Active meeting not found")
    if not row.owner_id:
        raise HTTPException(409, "This meeting has no recorded owner; ask an administrator to review it.")
    if row.owner_id == user.id:
        raise HTTPException(409, "You are the owner of this meeting.")
    existing = db.scalar(select(ReviewMeetingRequest).where(ReviewMeetingRequest.review_session_id == row.id,
        ReviewMeetingRequest.requester_id == user.id, ReviewMeetingRequest.status == "PENDING"))
    if not existing:
        db.add(ReviewMeetingRequest(review_session_id=row.id, requester_id=user.id, owner_id=row.owner_id,
                                    message=f"{user.full_name} is requesting you to join the meeting."))
        db.commit()
    return {"ok": True, "message": "Meeting owner has been notified in the application."}


@router.get("/collaboration/owner-requests")
def list_owner_requests(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    rows = db.scalars(select(ReviewMeetingRequest).where(ReviewMeetingRequest.owner_id == user.id,
        ReviewMeetingRequest.status == "PENDING").order_by(ReviewMeetingRequest.created_at.desc()).limit(50)).all()
    result = []
    for item in rows:
        requester = db.get(User, item.requester_id)
        meeting = db.get(ReviewSession, item.review_session_id)
        if meeting and meeting.ended_at is None:
            result.append({"id": item.id, "review_session_id": meeting.id, "review_date": meeting.review_date,
                           "requester_name": requester.full_name if requester else "User", "created_at": item.created_at,
                           "mom_no": f"MOM-{meeting.review_date.isoformat()}-{meeting.id:03d}"})
    return result


@router.get("/collaboration/my-requests")
def list_my_meeting_requests(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    rows = db.scalars(select(ReviewMeetingRequest).where(
        ReviewMeetingRequest.requester_id == user.id,
        ReviewMeetingRequest.status.in_(["PENDING", "ACKNOWLEDGED"])
    ).order_by(ReviewMeetingRequest.created_at.desc()).limit(50)).all()
    result = []
    for item in rows:
        meeting = db.get(ReviewSession, item.review_session_id)
        if meeting:
            result.append({
                "id": item.id,
                "review_session_id": meeting.id,
                "review_date": meeting.review_date,
                "status": item.status,
                "responded_at": item.responded_at,
                "mom_no": f"MOM-{meeting.review_date.isoformat()}-{meeting.id:03d}",
            })
    return result


@router.patch("/collaboration/owner-requests/{request_id}/acknowledge")
def acknowledge_owner_request(request_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    item = db.get(ReviewMeetingRequest, request_id)
    if not item or item.owner_id != user.id:
        raise HTTPException(404, "Meeting request not found")
    item.status = "ACKNOWLEDGED"
    item.responded_at = datetime.utcnow()
    db.commit()
    return {"ok": True}


@router.post("/actions/{action_id}/escalate-whywhy")
def escalate_whywhy(action_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    action = db.get(Action, action_id)
    if not action: raise HTTPException(404, "Action not found")
    if not action.requires_whywhy:
        action.requires_whywhy = True
        action.action_type = "WHY_WHY"
        if not db.scalar(select(ActionWhyWhy).where(ActionWhyWhy.action_id == action.id)):
            db.add(ActionWhyWhy(action_id=action.id, containment_action=action.action_description, updated_by_id=user.id))
        db.add(ActionHistory(action_id=action.id, changed_by_id=user.id, old_status=action.status,
                             new_status=action.status, comment="Escalated from Daily Review to Why-Why"))
        db.commit()
    return {"id": action.id, "requires_whywhy": True}


@router.get("/{review_id}/mom")
def meeting_mom(review_id: int, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    review = db.get(ReviewSession, review_id)
    if not review: raise HTTPException(404, "Review session not found")
    points = db.scalars(select(ReviewPoint).where(ReviewPoint.review_session_id == review.id).order_by(ReviewPoint.sequence_no)).all()
    point_action_ids = {p.action_id for p in points if p.action_id}
    linked_action_ids = set(db.scalars(
        select(ReviewActionLink.action_id).where(ReviewActionLink.review_session_id == review.id)
    ).all())
    current_action_ids = linked_action_ids | point_action_ids

    # Actions raised from Dashboard, OEE, Quality or the Actions page may be
    # linked to the active meeting without having a typed ReviewPoint. Keep
    # them visible as a separate register and avoid duplicating point actions.
    other_linked_ids = linked_action_ids - point_action_ids
    other_linked_rows = db.scalars(
        select(Action).where(Action.id.in_(other_linked_ids or {-1}))
        .order_by(Action.due_at, Action.priority.desc(), Action.reference_date)
    ).all()
    other_linked_actions = [action_data(db, action) for action in other_linked_rows]

    linked_pending = db.scalars(
        select(Action).join(ReviewActionLink, ReviewActionLink.action_id == Action.id)
        .join(ReviewSession, ReviewSession.id == ReviewActionLink.review_session_id)
        .where(ReviewSession.review_date < review.review_date, Action.status != ActionStatus.CLOSED)
        .distinct().order_by(Action.due_at, Action.reference_date)
    ).all()
    legacy_pending = db.scalars(
        select(Action).where(
            Action.action_type == "DAILY_REVIEW",
            Action.status != ActionStatus.CLOSED,
            Action.reference_date < review.review_date,
        ).order_by(Action.due_at, Action.reference_date)
    ).all()
    pending_rows = {a.id: a for a in linked_pending}
    for action in legacy_pending:
        pending_rows.setdefault(action.id, action)
    pending_rows = {aid: a for aid, a in pending_rows.items() if aid not in current_action_ids}

    pending_ids = list(pending_rows)
    followups: dict[int, dict] = {}
    if pending_ids:
        history = db.scalars(
            select(ActionHistory).where(ActionHistory.action_id.in_(pending_ids))
            .order_by(ActionHistory.changed_at.desc(), ActionHistory.id.desc())
        ).all()
        for event in history:
            comment = (event.comment or "").strip()
            lower = comment.lower()
            if not comment or lower.startswith("action created") or lower.startswith("attachment added:"):
                continue
            followups.setdefault(event.action_id, {"at": event.changed_at, "comment": comment})

    pending = [
        action_data(db, a, followups.get(a.id))
        for a in sorted(pending_rows.values(), key=lambda x: (x.due_at is None, x.due_at or datetime.max, x.reference_date))
    ]
    return {"meeting": serialize_review(review, db), "mom_no": f"MOM-{review.review_date.isoformat()}-{review.id:03d}",
            "points": [point_data(db, p) for p in points],
            "linked_actions": other_linked_actions,
            "previous_pending_actions": pending}


def _mom_rows(db: Session, review: ReviewSession):
    points = db.scalars(select(ReviewPoint).where(ReviewPoint.review_session_id == review.id).order_by(ReviewPoint.sequence_no)).all()
    rows=[]
    point_action_ids = {p.action_id for p in points if p.action_id}
    for p in points:
        a=db.get(Action,p.action_id) if p.action_id else None; ad=action_data(db,a)
        product = db.get(Product, p.product_id) if p.product_id else None
        rows.append([p.sequence_no,p.category or "",product.name if product else "",p.discussion_point,
                     (ad or {}).get("description",""),(ad or {}).get("owner",""),
                     format_ist((ad or {}).get("due_at")),(ad or {}).get("priority",""),
                     (ad or {}).get("status",""),
                     (ad or {}).get("closure_remark","") if (ad or {}).get("status") == "CLOSED" else ""])

    # Include actions linked through other pages (Dashboard, Quality, OEE,
    # Actions) so the exported MOM contains the same complete action register
    # as the on-screen meeting page.
    linked_action_ids = set(db.scalars(
        select(ReviewActionLink.action_id).where(ReviewActionLink.review_session_id == review.id)
    ).all()) - point_action_ids
    linked_actions = db.scalars(
        select(Action).where(Action.id.in_(linked_action_ids or {-1}))
        .order_by(Action.due_at, Action.priority.desc(), Action.reference_date)
    ).all()
    for action in linked_actions:
        owner = db.get(User, action.owner_id) if action.owner_id else None
        product_names = db.scalars(
            select(Product.name).join(ActionContext, ActionContext.product_id == Product.id)
            .where(ActionContext.action_id == action.id).distinct().order_by(Product.name)
        ).all()
        rows.append([
            "", "Linked Action", ", ".join(product_names),
            action.problem_description, action.action_description,
            owner.full_name if owner else "Unassigned",
            format_ist(action.due_at), action.priority.value, action.status.value,
            action.closure_remark or "",
        ])
    return rows


@router.get("/{review_id}/pdf")
def download_mom_pdf(review_id: int, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    review=db.get(ReviewSession,review_id)
    if not review: raise HTTPException(404,"Review session not found")
    mom_no=f"MOM-{review.review_date.isoformat()}-{review.id:03d}"
    buf=BytesIO(); doc=SimpleDocTemplate(buf,pagesize=landscape(A4),rightMargin=10*mm,leftMargin=10*mm,topMargin=10*mm,bottomMargin=10*mm)
    styles=getSampleStyleSheet()
    def safe(value):
        return str(value or "").replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")

    control_style=styles["BodyText"].clone("MOMControl")
    control_style.fontName="Helvetica"
    control_style.fontSize=8
    control_style.leading=12
    title_style=styles["Title"].clone("MOMCorporateTitle")
    title_style.fontName="Helvetica-Bold"
    title_style.fontSize=19
    title_style.leading=23
    title_style.textColor=colors.HexColor("#17365D")
    title_style.alignment=2
    subtitle_style=styles["BodyText"].clone("MOMSubtitle")
    subtitle_style.fontSize=8
    subtitle_style.leading=11
    subtitle_style.textColor=colors.HexColor("#526579")
    subtitle_style.alignment=2
    control_block=Paragraph(
        "<b>Ref. No.:</b> IMS/MR/F/14<br/>"
        "<b>Issue No. / Date:</b> 02 / 16.11.2022<br/>"
        "<b>Rev. No. / Date:</b> 01 / 16.11.2022",
        control_style
    )
    title_block=Paragraph("DAILY REVIEW<br/>MINUTES OF MEETING",title_style)
    header=Table([[control_block,title_block]],colWidths=[doc.width*0.48,doc.width*0.52])
    header.setStyle(TableStyle([
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("ALIGN",(1,0),(1,0),"RIGHT"),
        ("LINEBELOW",(0,0),(-1,-1),1.2,colors.HexColor("#1F4E78")),
        ("LEFTPADDING",(0,0),(-1,-1),0),
        ("RIGHTPADDING",(0,0),(-1,-1),0),
        ("TOPPADDING",(0,0),(-1,-1),2),
        ("BOTTOMPADDING",(0,0),(-1,-1),7),
    ]))
    meta_style=styles["BodyText"].clone("MOMMeta")
    meta_style.fontSize=8
    meta_style.leading=11
    meta_style.textColor=colors.HexColor("#243447")
    meta_data=[[
        Paragraph(f"<b>MOM Reference</b><br/>{safe(mom_no)}",meta_style),
        Paragraph(f"<b>Meeting Date</b><br/>{safe(review.review_date)}",meta_style),
        Paragraph(f"<b>Participants</b><br/>{safe(review.participants or '-')}",meta_style),
    ]]
    meta=Table(meta_data,colWidths=[doc.width*0.28,doc.width*0.20,doc.width*0.52])
    meta.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#EEF3F8")),
        ("BOX",(0,0),(-1,-1),0.5,colors.HexColor("#C5D1DE")),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("LEFTPADDING",(0,0),(-1,-1),7),("RIGHTPADDING",(0,0),(-1,-1),7),
        ("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6),
    ]))
    story=[header,Spacer(1,3*mm),meta,Spacer(1,5*mm)]
    raw_data=[["Sr.","Category","Product","Discussion / MOM Point","Action Required","Owner","Target","Priority","Status","Closure Remark"]]+_mom_rows(db,review)
    cell_style=styles["BodyText"].clone("MOMCell"); cell_style.fontSize=7; cell_style.leading=9
    head_style=styles["BodyText"].clone("MOMHead"); head_style.fontSize=7; head_style.leading=9; head_style.textColor=colors.white; head_style.fontName="Helvetica-Bold"
    data=[[Paragraph(str(v or "").replace("&","&amp;").replace("<","&lt;").replace(">","&gt;"), head_style if ri==0 else cell_style) for v in row] for ri,row in enumerate(raw_data)]
    base_col_widths=[9*mm,16*mm,24*mm,40*mm,40*mm,24*mm,20*mm,14*mm,18*mm,42*mm]
    width_scale=doc.width/sum(base_col_widths)
    table=Table(data,colWidths=[w*width_scale for w in base_col_widths],repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1f4e78")),("TEXTCOLOR",(0,0),(-1,0),colors.white),
      ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("FONTSIZE",(0,0),(-1,-1),8),("VALIGN",(0,0),(-1,-1),"TOP"),
      ("GRID",(0,0),(-1,-1),0.4,colors.grey),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f4f7fa")])]))
    story.append(table)
    if review.comments: story += [Spacer(1,4*mm),Paragraph(f"<b>Meeting remarks:</b> {safe(review.comments)}",styles["BodyText"])]
    def add_page_footer(canvas, document):
        canvas.saveState()
        page_width,page_height=landscape(A4)
        canvas.setStrokeColor(colors.HexColor("#C5D1DE"))
        canvas.setLineWidth(0.5)
        canvas.line(doc.leftMargin,8*mm,page_width-doc.rightMargin,8*mm)
        canvas.setFont("Helvetica",7)
        canvas.setFillColor(colors.HexColor("#66788A"))
        canvas.drawString(doc.leftMargin,5*mm,"IMS/MR/F/14  |  Controlled document")
        canvas.drawRightString(page_width-doc.rightMargin,5*mm,f"Page {document.page}")
        canvas.restoreState()

    doc.build(story,onFirstPage=add_page_footer,onLaterPages=add_page_footer)
    return Response(buf.getvalue(),media_type="application/pdf",headers={"Content-Disposition":f'attachment; filename="{mom_no}.pdf"'})


@router.get("/{review_id}/xlsx")
def download_mom_xlsx(review_id: int, db: Session = Depends(get_db), _: User = Depends(get_current_user)):
    review=db.get(ReviewSession,review_id)
    if not review: raise HTTPException(404,"Review session not found")
    mom_no=f"MOM-{review.review_date.isoformat()}-{review.id:03d}"
    wb=Workbook(); ws=wb.active; ws.title="MOM"
    ws.append(["Daily Review - Minutes of Meeting"]); ws.append(["MOM No",mom_no]); ws.append(["Date",review.review_date.isoformat()])
    ws.append(["Participants",review.participants or ""]); ws.append([])
    headers=["Sr.","Category","Product","Discussion / MOM Point","Action Required","Owner","Target Date","Priority","Status","Closure Remark"]; ws.append(headers)
    for row in _mom_rows(db,review): ws.append(row)
    for cell in ws[6]:
        cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="1F4E78"); cell.alignment=Alignment(wrap_text=True)
    widths=[8,16,24,42,40,24,18,14,16,42]
    for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=w
    for row in ws.iter_rows(min_row=6, max_row=ws.max_row, min_col=1, max_col=9):
        for cell in row:
            cell.alignment=Alignment(wrap_text=True, vertical="top")
    for row_idx in range(7, ws.max_row + 1):
        ws.row_dimensions[row_idx].height=45
    ws.freeze_panes="A7"; ws.auto_filter.ref=f"A6:J{ws.max_row}"
    out=BytesIO(); wb.save(out)
    return Response(out.getvalue(),media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition":f'attachment; filename="{mom_no}.xlsx"'})
