from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import (
    CastingDefectDaily, CastingDefectPhenomenon, Customer, CustomerRejectionDaily,
    CustomerRejectionPhenomenon, Product, Vendor,
)


def norm(value) -> str:
    return " ".join(str(value or "").replace("\n", " ").strip().lower().split())


def code(value: str, prefix: str) -> str:
    text = re.sub(r"[^A-Za-z0-9]+", "_", str(value or "").strip().upper()).strip("_")
    return (text[:70] or prefix)[:80]


def as_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def decimal_value(value, *, required=False) -> Decimal | None:
    if value in (None, ""):
        if required:
            raise ValueError("is required")
        return Decimal("0")
    try:
        result = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError, AttributeError) as exc:
        raise ValueError("must be numeric") from exc
    if not result.is_finite() or result < 0:
        raise ValueError("must be a non-negative number")
    return result


def headers(ws) -> dict[str, int]:
    return {norm(cell.value): cell.column for cell in ws[1] if cell.value not in (None, "")}


def cell(ws, row: int, h: dict[str, int], *names: str):
    for name in names:
        column = h.get(norm(name))
        if column:
            return ws.cell(row, column).value
    return None


def lookup(rows, value):
    target = norm(value)
    if not target:
        return None
    return next((row for row in rows if norm(row.name) == target or norm(row.code) == target), None)


def create_phenomenon(db: Session, model, name: str, *, group: str | None = None, criticality: str | None = None):
    normalized = norm(name)
    if not normalized:
        raise ValueError("Phenomenon name is required")
    row = db.scalar(select(model).where(model.normalized_name == normalized))
    if not row:
        base = code(name, "PHEN")
        generated = base
        suffix = 2
        while db.scalar(select(model).where(model.code == generated)):
            generated = f"{base[:70]}_{suffix}"
            suffix += 1
        row = model(code=generated, name=" ".join(str(name).split()), normalized_name=normalized)
        db.add(row)
        db.flush()
    if group not in (None, ""):
        row.phenomenon_group = str(group).strip()
    if criticality not in (None, ""):
        row.criticality = str(criticality).strip().upper()
    return row


def _apply(existing, values: dict) -> bool:
    changed = False
    for key, value in values.items():
        if getattr(existing, key) != value:
            setattr(existing, key, value)
            changed = True
    return changed


def import_casting_defects_workbook(db: Session, path: str | Path, *, batch_id=None, entered_by_id=None) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet = "Casting_Defect_Data"
    if sheet not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Required sheet '{sheet}' was not found")
    ws = wb[sheet]
    h = headers(ws)
    required = ["Date", "Shift", "Product", "Inspected Qty", "Phenomenon", "Defect Qty"]
    missing = [name for name in required if norm(name) not in h]
    if missing:
        wb.close()
        raise ValueError("Missing columns: " + ", ".join(missing))
    products = db.scalars(select(Product)).all()
    vendors = db.scalars(select(Vendor)).all()
    phenomena = db.scalars(select(CastingDefectPhenomenon).where(CastingDefectPhenomenon.is_active.is_(True))).all()
    stats = {"rows_read": 0, "created": 0, "updated": 0, "unchanged": 0, "errors": [], "warnings": []}
    scope_denominators: dict[str, Decimal] = {}
    for number, row in enumerate(ws.iter_rows(min_row=2, values_only=False), 2):
        if all(x.value in (None, "") for x in row):
            continue
        stats["rows_read"] += 1
        on_date = as_date(cell(ws, number, h, "Date"))
        shift = str(cell(ws, number, h, "Shift") or "").strip().upper()
        if shift == "GENERAL":
            shift = "General"
        product = lookup(products, cell(ws, number, h, "Product", "Product Code"))
        vendor_raw = cell(ws, number, h, "Vendor", "Vendor Code")
        vendor = lookup(vendors, vendor_raw) if vendor_raw not in (None, "") else None
        phenomenon = lookup(phenomena, cell(ws, number, h, "Phenomenon"))
        errors = []
        if not on_date: errors.append("invalid Date")
        if shift not in {"A", "B", "C", "General"}: errors.append("Shift must be A, B, C or General")
        if not product: errors.append("Product does not match the current master")
        if vendor_raw not in (None, "") and not vendor: errors.append("Vendor does not match the current master")
        if not phenomenon: errors.append("Phenomenon does not match the approved Casting master")
        try:
            inspected = decimal_value(cell(ws, number, h, "Inspected Qty", "Received Qty"), required=True)
            defect = decimal_value(cell(ws, number, h, "Defect Qty", "Reject Qty"), required=True)
            rework = decimal_value(cell(ws, number, h, "Rework Qty"))
            scrap = decimal_value(cell(ws, number, h, "Scrap Qty"))
            if inspected <= 0: errors.append("Inspected Qty must be greater than zero")
            if defect > inspected: errors.append("Defect Qty cannot exceed Inspected Qty")
            if rework + scrap > defect: errors.append("Rework Qty plus Scrap Qty cannot exceed Defect Qty")
        except ValueError as exc:
            inspected = defect = rework = scrap = Decimal("0")
            errors.append(f"quantities {exc}")
        if errors:
            stats["errors"].append(f"Row {number}: " + "; ".join(errors))
            continue
        heat_batch = str(cell(ws, number, h, "Heat/Batch No", "Heat Batch No", "Batch No") or "").strip() or None
        scope_key = "|".join([on_date.isoformat(), norm(shift), str(product.id), str(vendor.id if vendor else 0), norm(heat_batch)])
        if scope_key not in scope_denominators:
            prior = db.scalar(select(CastingDefectDaily.inspected_qty).where(CastingDefectDaily.record_key.startswith(scope_key + "|")).limit(1))
            scope_denominators[scope_key] = prior if prior is not None else inspected
        if scope_denominators[scope_key] != inspected:
            stats["errors"].append(f"Row {number}: Inspected Qty must match other phenomena for the same date/shift/product/vendor/batch")
            continue
        key = scope_key + "|" + str(phenomenon.id)
        ppm = (defect * Decimal("1000000") / inspected).quantize(Decimal("0.001"))
        values = dict(defect_date=on_date, shift=shift, product_id=product.id, plant=product.plant,
                      vendor_id=vendor.id if vendor else None, heat_batch_no=heat_batch, phenomenon_id=phenomenon.id,
                      inspected_qty=inspected, defect_qty=defect, rework_qty=rework, scrap_qty=scrap, ppm=ppm,
                      remark=str(cell(ws, number, h, "Remark", "Remarks") or "").strip() or None,
                      source="EXCEL", import_batch_id=batch_id, entered_by_id=entered_by_id)
        existing = db.scalar(select(CastingDefectDaily).where(CastingDefectDaily.record_key == key))
        if not existing:
            db.add(CastingDefectDaily(record_key=key, **values)); stats["created"] += 1
        elif _apply(existing, values): stats["updated"] += 1
        else: stats["unchanged"] += 1
    wb.close()
    return stats


def import_customer_rejections_workbook(db: Session, path: str | Path, *, batch_id=None, entered_by_id=None) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    sheet = "Customer_Rejection_Data"
    if sheet not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Required sheet '{sheet}' was not found")
    ws = wb[sheet]
    h = headers(ws)
    required = ["Date", "Customer", "Product", "Dispatch Qty", "Phenomenon", "Reject Qty"]
    missing = [name for name in required if norm(name) not in h]
    if missing:
        wb.close()
        raise ValueError("Missing columns: " + ", ".join(missing))
    customers = db.scalars(select(Customer)).all()
    products = db.scalars(select(Product)).all()
    phenomena = db.scalars(select(CustomerRejectionPhenomenon).where(CustomerRejectionPhenomenon.is_active.is_(True))).all()
    stats = {"rows_read": 0, "created": 0, "updated": 0, "unchanged": 0, "errors": [], "warnings": []}
    scope_denominators: dict[str, Decimal] = {}
    for number, row in enumerate(ws.iter_rows(min_row=2, values_only=False), 2):
        if all(x.value in (None, "") for x in row): continue
        stats["rows_read"] += 1
        on_date = as_date(cell(ws, number, h, "Date"))
        customer = lookup(customers, cell(ws, number, h, "Customer", "Customer Code"))
        product = lookup(products, cell(ws, number, h, "Product", "Product Code"))
        phenomenon = lookup(phenomena, cell(ws, number, h, "Phenomenon"))
        errors = []
        if not on_date: errors.append("invalid Date")
        if not customer: errors.append("Customer does not match the current master")
        if not product: errors.append("Product does not match the current master")
        if customer and product and product.customer_id and product.customer_id != customer.id:
            errors.append("Product is not assigned to the selected Customer")
        if not phenomenon: errors.append("Phenomenon does not match the approved Customer Rejection master")
        try:
            dispatch = decimal_value(cell(ws, number, h, "Dispatch Qty", "Received Qty", "Supplied Qty"), required=True)
            reject = decimal_value(cell(ws, number, h, "Reject Qty", "Rejection Qty"), required=True)
            if dispatch <= 0: errors.append("Dispatch Qty must be greater than zero")
            if reject > dispatch: errors.append("Reject Qty cannot exceed Dispatch Qty")
        except ValueError as exc:
            dispatch = reject = Decimal("0"); errors.append(f"quantities {exc}")
        if errors:
            stats["errors"].append(f"Row {number}: " + "; ".join(errors)); continue
        reference = str(cell(ws, number, h, "Reference No", "Complaint No", "Invoice No") or "").strip() or None
        batch = str(cell(ws, number, h, "Batch No", "Heat/Batch No") or "").strip() or None
        scope_key = "|".join([on_date.isoformat(), str(customer.id), str(product.id), norm(reference), norm(batch)])
        if scope_key not in scope_denominators:
            prior = db.scalar(select(CustomerRejectionDaily.dispatch_qty).where(CustomerRejectionDaily.record_key.startswith(scope_key + "|")).limit(1))
            scope_denominators[scope_key] = prior if prior is not None else dispatch
        if scope_denominators[scope_key] != dispatch:
            stats["errors"].append(f"Row {number}: Dispatch Qty must match other phenomena for the same date/customer/product/reference/batch")
            continue
        key = scope_key + "|" + str(phenomenon.id)
        ppm = (reject * Decimal("1000000") / dispatch).quantize(Decimal("0.001"))
        values = dict(rejection_date=on_date, customer_id=customer.id, product_id=product.id,
                      reference_no=reference, batch_no=batch, phenomenon_id=phenomenon.id,
                      dispatch_qty=dispatch, reject_qty=reject, ppm=ppm,
                      remark=str(cell(ws, number, h, "Remark", "Remarks") or "").strip() or None,
                      source="EXCEL", import_batch_id=batch_id, entered_by_id=entered_by_id)
        existing = db.scalar(select(CustomerRejectionDaily).where(CustomerRejectionDaily.record_key == key))
        if not existing:
            db.add(CustomerRejectionDaily(record_key=key, **values)); stats["created"] += 1
        elif _apply(existing, values): stats["updated"] += 1
        else: stats["unchanged"] += 1
    wb.close()
    return stats
