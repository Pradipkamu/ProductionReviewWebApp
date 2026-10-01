"""Explicit process stages; column positions are provenance, never identity."""
import hashlib
import json
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy import select, func
from ..models import (Product, Customer, Vendor, Operation, RouteVersion, RouteOperation,
    ProcessFlowVersion, ProcessFlowStage, StageScheduleAllocation, DailyRequirement, ProcessDailySummary, DailyMIS)
from ..enums import OperationType, SourceType
from .planning import month_bounds, working_days, product_plant
from .pricing import price_for_date

ROLES={'PRODUCTION','RECEIPT','DISPATCH','VENDOR_OUT','VENDOR_IN','SUPPLY','VENDOR_PROCESS','INSPECTION','DISPATCH_DETAIL','DISPATCH_TOTAL','REFERENCE'}
SHEETS={'process-design':'Flow_Definition','stage-schedules':'Stage_Schedules','stage-daily':'Daily_Actuals'}
HEADERS={
 'process-design':['product','stage_code','stage_name','role','branch','variant','vendor','predecessors','source_column','active','parent_dispatch','alias_of','effective_from','reason','company','plant','customer','type'],
 'stage-schedules':['product','stage_code','month','effective_from','allocated_qty','reference','reason','correct_imported_plans'],
 'stage-daily':['product','stage_code','date','actual_qty','reject_qty','reason'],
}

def approved_design():
    return json.loads((Path(__file__).resolve().parents[1]/'data'/'approved_process_design.json').read_text())

def to_date(value):
    if isinstance(value,datetime):return value.date()
    if isinstance(value,date):return value
    try:return date.fromisoformat(str(value).strip())
    except ValueError:raise ValueError('Use an Excel date or YYYY-MM-DD')

def quantity(value,field):
    if value in (None,''):raise ValueError(f'{field} is missing; enter 0 explicitly when zero is intended')
    try:q=Decimal(str(value))
    except InvalidOperation:raise ValueError(f'{field} must be whole pieces')
    if not q.is_finite() or q<0 or q!=q.to_integral_value() or q>Decimal('9999999999999'):raise ValueError(f'{field} must be nonnegative whole pieces')
    return q

def boolean(value):
    if str(value or '').strip().lower() in {'true','1','yes'}:return True
    if str(value or '').strip().lower() in {'false','0','no',''}:return False
    raise ValueError('Use TRUE or FALSE')

def text(value):return str(value or '').strip()

def product_for(db,name):
    products=db.scalars(select(Product).where((func.lower(Product.name)==name.lower())|(func.lower(Product.code)==name.lower()))).all()
    if len(products)!=1:raise ValueError(f'Product {name!r} is missing or ambiguous')
    return products[0]

def active_flow(db,pid,d):
    return db.scalar(select(ProcessFlowVersion).where(ProcessFlowVersion.product_id==pid,ProcessFlowVersion.effective_from<=d).order_by(ProcessFlowVersion.effective_from.desc(),ProcessFlowVersion.revision_no.desc()).limit(1))

def stage_for(db,product,code,d):
    flow=active_flow(db,product.id,d)
    if not flow:raise ValueError('No approved process flow is effective on this date')
    stage=db.scalar(select(ProcessFlowStage).where(ProcessFlowStage.flow_id==flow.id,ProcessFlowStage.code==code))
    if not stage or not stage.is_active or not stage.route_operation_id:raise ValueError(f'Stage {code!r} is unknown, deferred or historical only')
    return stage

def read_rows(path,kind):
    sheet=SHEETS[kind];wb=load_workbook(path,data_only=False,read_only=True)
    try:
        if sheet not in wb.sheetnames:raise ValueError(f'Missing sheet {sheet}')
        ws=wb[sheet];headers=[text(v).lower() for v in next(ws.iter_rows(values_only=True))]
        if len(headers)!=len(set(headers)):raise ValueError('Duplicate workbook headers')
        missing=set(HEADERS[kind])-set(headers)
        if missing:raise ValueError('Missing headers: '+', '.join(sorted(missing)))
        rows=[]
        for number,values in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
            if all(v in (None,'') for v in values):continue
            if any(isinstance(v,str) and v.startswith('=') for v in values):raise ValueError(f'Row {number}: enter values; formulas are not imported')
            rows.append((number,dict(zip(headers,values))))
        if not rows:raise ValueError('Workbook contains no business rows')
        return rows
    finally:wb.close()

def validate_definition(rows):
    codes={r['stage_code']:r for _,r in rows}
    if len(codes)!=len(rows):raise ValueError('Duplicate stage code')
    if len({r['effective_from'] for _,r in rows})!=1:raise ValueError('All stages of a product revision require the same effective date')
    active_columns=[r['source_column'] for _,r in rows if r['active'] and r['source_column']]
    if len(active_columns)!=len(set(active_columns)):raise ValueError('Duplicate active source column')
    parents=[r for _,r in rows if r['active'] and r['parent_dispatch']]
    if any(r['active'] for _,r in rows) and len(parents)!=1:raise ValueError('Exactly one active parent dispatch stage is required')
    for _,r in rows:
        if r['role'] not in ROLES:raise ValueError('Unknown stage role '+r['role'])
        if r['active'] and r['role']=='REFERENCE':raise ValueError('Reference stages must be inactive')
        if r['parent_dispatch'] and r['role'] not in {'DISPATCH','DISPATCH_TOTAL'}:raise ValueError('Only a dispatch stage may feed parent MIS')
        if r['active'] and r['alias_of']:raise ValueError('Duplicate aliases must remain inactive')
        if r['alias_of'] and r['alias_of'] not in codes:raise ValueError('Alias points to an unknown stage')
        for pred in r['predecessors']:
            if pred not in codes or not codes[pred]['active']:raise ValueError('Predecessor is missing or inactive: '+pred)
        if r['role']=='VENDOR_IN' and r['active']:
            outs=[codes[c] for c in r['predecessors'] if codes[c]['role']=='VENDOR_OUT']
            if not outs:raise ValueError('Vendor receipt requires an outward predecessor')
            if any(o['vendor']!=r['vendor'] for o in outs):raise ValueError('Vendor receipt and outward must use the same vendor')
    visiting=set();done=set()
    def visit(code):
        if code in visiting:raise ValueError('Process dependency cycle at '+code)
        if code in done:return
        visiting.add(code)
        for pred in codes[code]['predecessors']:visit(pred)
        visiting.remove(code);done.add(code)
    for code in codes:visit(code)

def apply_definition(db,rows,stats):
    product_name=rows[0][1]['product'];validate_definition(rows)
    canonical=json.dumps([r for _,r in rows],default=str,sort_keys=True)
    digest=hashlib.sha256(canonical.encode()).hexdigest();first=rows[0][1]
    existing=db.scalars(select(Product).where(func.lower(Product.name)==product_name.lower())).all()
    if len(existing)>1:raise ValueError('Product name is ambiguous')
    if existing:product=existing[0]
    else:
        customer_name=first['customer'];customer=db.scalar(select(Customer).where(func.lower(Customer.name)==customer_name.lower())) if customer_name else None
        if customer_name and not customer:
            customer=Customer(code='FLOW_'+hashlib.sha256(customer_name.encode()).hexdigest()[:12],name=customer_name);db.add(customer);db.flush()
        product=Product(code='FLOW_'+hashlib.sha256(product_name.encode()).hexdigest()[:12],name=product_name,customer_id=customer.id if customer else None,plant=first['plant'] or None,product_group=first['type'] or None,is_active=any(r['active'] for _,r in rows));db.add(product);db.flush()
    latest=db.scalar(select(ProcessFlowVersion).where(ProcessFlowVersion.product_id==product.id).order_by(ProcessFlowVersion.revision_no.desc()).limit(1))
    if latest and latest.definition_sha256==digest:stats['unchanged']+=len(rows);return
    if latest and first['effective_from']<latest.effective_from:raise ValueError('Flow revisions must follow the latest effective date')
    revision=int(db.scalar(select(func.coalesce(func.max(RouteVersion.revision_no),-1)).where(RouteVersion.product_id==product.id)))+1
    route=RouteVersion(product_id=product.id,revision_no=revision,effective_from=first['effective_from'],description='Approved explicit stages; independent allocations',is_active=True);db.add(route);db.flush()
    flow=ProcessFlowVersion(product_id=product.id,route_version_id=route.id,effective_from=first['effective_from'],revision_no=revision,company=first['company'] or None,definition_sha256=digest,source_document='Flow_Definition import',reason=first['reason']);db.add(flow);db.flush()
    for seq,(_,r) in enumerate(rows,1):
        route_op=None
        if r['active']:
            typ={'VENDOR_OUT':OperationType.VENDOR_OUT,'VENDOR_IN':OperationType.VENDOR_IN,'VENDOR_PROCESS':OperationType.VENDOR_PROCESS,'DISPATCH':OperationType.DISPATCH,'DISPATCH_TOTAL':OperationType.DISPATCH,'DISPATCH_DETAIL':OperationType.DISPATCH,'INSPECTION':OperationType.INSPECTION}.get(r['role'],OperationType.INTERNAL)
            operation_code='FLOW_'+hashlib.sha256(f'{product.id}|{revision}|{r["stage_code"]}'.encode()).hexdigest()[:40]
            operation=db.scalar(select(Operation).where(Operation.code==operation_code))
            if not operation:
                operation=Operation(code=operation_code,name=r['stage_name'],operation_type=typ);db.add(operation);db.flush()
            vendor=db.scalar(select(Vendor).where(func.lower(Vendor.name)==r['vendor'].lower())) if r['vendor'] else None
            if r['vendor'] and not vendor:
                vendor=Vendor(code='FLOW_'+hashlib.sha256(r['vendor'].encode()).hexdigest()[:12],name=r['vendor']);db.add(vendor);db.flush()
            route_op=RouteOperation(route_version_id=route.id,operation_id=operation.id,sequence_no=seq*10,source_label=r['stage_name'],vendor_id=vendor.id if vendor else None,is_dispatch=r['parent_dispatch'],is_enabled=True)
            db.add(route_op);db.flush()
        db.add(ProcessFlowStage(flow_id=flow.id,code=r['stage_code'],name=r['stage_name'],route_operation_id=route_op.id if route_op else None,source_column=r['source_column'],role=r['role'],branch=r['branch'],variant=r['variant'],vendor_name=r['vendor'],predecessors_json=json.dumps(r['predecessors']),alias_of=r['alias_of'],is_active=r['active'],parent_dispatch=r['parent_dispatch'],sequence_no=seq*10))
        if r['active'] and r['role'].startswith('VENDOR') and not r['vendor']:stats['warnings'].append(f'{product_name} / {r["stage_name"]}: vendor assignment pending')
    db.flush();stats['new']+=len(rows)

def apply_allocation(db,r,stats):
    p=product_for(db,r['product']);d=to_date(r['effective_from']);month=to_date(r['month'])
    if month.day!=1 or month!=d.replace(day=1):raise ValueError('Month must be its first day and contain effective_from')
    s=stage_for(db,p,text(r['stage_code']),d);qty=quantity(r['allocated_qty'],'allocated_qty')
    _,last=month_bounds(month)
    reason=text(r['reason']);reference=text(r['reference'])
    if not reason or not reference:raise ValueError('Allocation reason and reference are required')
    previous=db.scalars(select(StageScheduleAllocation).where(StageScheduleAllocation.stage_id==s.id,StageScheduleAllocation.month==month).order_by(StageScheduleAllocation.revision_no.desc())).all()
    for old in previous:
        if old.reference==reference:
            if old.allocated_qty==qty and old.effective_from==d:stats['unchanged']+=1;return
            raise ValueError('Reference already exists with different values; use a new revision reference')
    if previous and d<previous[0].effective_from:raise ValueError('Schedule revisions must follow the latest effective date')
    # Stable stage codes span flow revisions; count only the applicable version on each prior date.
    prior=Decimal(0);prior_date=month
    while prior_date<d:
        prior_flow=active_flow(db,p.id,prior_date)
        prior_stage=db.scalar(select(ProcessFlowStage).where(ProcessFlowStage.flow_id==prior_flow.id,ProcessFlowStage.code==s.code)) if prior_flow else None
        if prior_stage and prior_stage.route_operation_id:
            issued=db.scalar(select(DailyRequirement.revised_plan_qty).where(DailyRequirement.route_operation_id==prior_stage.route_operation_id,DailyRequirement.req_date==prior_date))
            prior+=Decimal(str(issued or 0))
        prior_date+=timedelta(days=1)
    balance=qty-Decimal(str(prior or 0))
    if balance<0:raise ValueError('Allocation is below the plan already issued before effective_from')
    days=working_days(db,d,last,product_plant(db,p.id))
    if not days and balance:raise ValueError('No remaining working days in the plant calendar')
    q,rem=divmod(int(balance),len(days)) if days else (0,0)
    dist={day:q+(i<rem) for i,day in enumerate(days)}
    existing={x.req_date:x for x in db.scalars(select(DailyRequirement).where(DailyRequirement.route_operation_id==s.route_operation_id,DailyRequirement.req_date>=d,DailyRequirement.req_date<=last))}
    if any(x.is_frozen for x in existing.values()) and not boolean(r['correct_imported_plans']):raise ValueError('Imported plans are protected; set correct_imported_plans TRUE with a reason')
    allocation=StageScheduleAllocation(stage_id=s.id,month=month,effective_from=d,revision_no=previous[0].revision_no+1 if previous else 1,allocated_qty=qty,working_dates_json=json.dumps([x.isoformat() for x in days]),distribution_json=json.dumps({x.isoformat():int(y) for x,y in dist.items()}),reference=reference,reason=reason);db.add(allocation)
    current=d
    while current<=last:
        row=existing.get(current);daily=Decimal(dist.get(current,0))
        if not row:db.add(DailyRequirement(req_date=current,product_id=p.id,route_operation_id=s.route_operation_id,baseline_plan_qty=daily,revised_plan_qty=daily,is_frozen=False))
        else:row.revised_plan_qty=daily
        current+=timedelta(days=1)
    original_reason=db.info.get('reason')
    db.info['reason']=reason
    try:db.flush()
    finally:db.info['reason']=original_reason
    stats['new']+=1
    stats.setdefault('allocations',[]).append({'product':p.name,'stage':s.code,'monthly_qty':int(qty),'plan_before_effective':int(prior or 0),'remaining_qty':int(balance),'working_days':len(days),'daily':[{'date':x.isoformat(),'qty':int(y)} for x,y in dist.items()]})

def apply_actual(db,r,stats):
    p=product_for(db,r['product']);d=to_date(r['date']);s=stage_for(db,p,text(r['stage_code']),d)
    actual=quantity(r['actual_qty'],'actual_qty');reject=quantity(r['reject_qty'],'reject_qty')
    if reject>actual:raise ValueError('Reject quantity cannot exceed actual quantity')
    if s.role in {'DISPATCH','DISPATCH_DETAIL','DISPATCH_TOTAL'} and reject:raise ValueError('Dispatch quantity must be accepted dispatched pieces; record rejection at its detection stage')
    row=db.scalar(select(ProcessDailySummary).where(ProcessDailySummary.summary_date==d,ProcessDailySummary.route_operation_id==s.route_operation_id))
    reason=text(r['reason']) or db.info.get('request_reason',db.info.get('reason',''))
    changed=row and (row.actual_qty!=actual or row.reject_qty!=reject)
    if changed and not reason:raise ValueError('Changing an existing actual requires a correction reason')
    if row and not changed:
        parent_mis=db.scalar(select(DailyMIS).where(DailyMIS.product_id==p.id,DailyMIS.mis_date==d)) if s.parent_dispatch else None
        if not s.parent_dispatch or (parent_mis and parent_mis.actual_qty==actual):
            stats['unchanged']+=1;return
    requirement=db.scalar(select(DailyRequirement).where(DailyRequirement.route_operation_id==s.route_operation_id,DailyRequirement.req_date==d))
    if not requirement:stats['warnings'].append(f'{p.name} / {s.name} / {d}: stage allocation missing; actual retained without a supplied plan')
    if not row:
        row=ProcessDailySummary(product_id=p.id,summary_date=d,route_operation_id=s.route_operation_id,plan_qty=requirement.revised_plan_qty if requirement else 0,source=SourceType(r.get('source') or 'EXCEL'));db.add(row);stats['new']+=1
    else:stats['updated']+=1
    row.actual_qty=actual;row.reject_qty=reject;row.good_qty=actual-reject;row.remarks=reason or 'Stage daily import';row.source=SourceType(r.get('source') or 'EXCEL')
    if s.parent_dispatch:
        mis=db.scalar(select(DailyMIS).where(DailyMIS.product_id==p.id,DailyMIS.mis_date==d))
        if mis and mis.actual_qty!=actual and not reason:raise ValueError('Changing existing parent MIS dispatch requires a correction reason')
        if not mis:
            price=price_for_date(db,p.id,d)
            parent_req=db.scalar(select(DailyRequirement).where(DailyRequirement.product_id==p.id,DailyRequirement.req_date==d,DailyRequirement.route_operation_id.is_(None)))
            plan=parent_req.revised_plan_qty if parent_req else 0
            mis=DailyMIS(product_id=p.id,mis_date=d,plan_qty=plan,sales_price=price,plan_sales=plan*price,source=SourceType(r.get('source') or 'EXCEL'));db.add(mis)
        mis.actual_qty=actual;mis.actual_sales=actual*mis.sales_price
    original_reason=db.info.get('reason')
    if reason:db.info['reason']=reason
    try:db.flush()
    finally:db.info['reason']=original_reason

def check_dispatch_totals(db,contexts):
    for pid,d in contexts:
        flow=active_flow(db,pid,d)
        stages=db.scalars(select(ProcessFlowStage).where(ProcessFlowStage.flow_id==flow.id,ProcessFlowStage.is_active.is_(True))).all()
        parent=next((s for s in stages if s.parent_dispatch and s.role=='DISPATCH_TOTAL'),None)
        if not parent:continue
        detail=[s for s in stages if s.role=='DISPATCH_DETAIL']
        vals={x.route_operation_id:x.actual_qty for x in db.scalars(select(ProcessDailySummary).where(ProcessDailySummary.product_id==pid,ProcessDailySummary.summary_date==d))}
        if parent.route_operation_id in vals and detail and all(s.route_operation_id in vals for s in detail):
            if sum(vals[s.route_operation_id] for s in detail)!=vals[parent.route_operation_id]:raise ValueError(f'{d}: parent dispatch total differs from all variant dispatch details')

def import_process_workbook(db,path,kind):
    stats={'new':0,'updated':0,'unchanged':0,'errors':[],'warnings':[]};rows=read_rows(path,kind)
    db.info['request_reason']=db.info.get('reason','')
    if kind=='process-design':
        groups={}
        for n,raw in rows:
            try:
                r={k:text(v) for k,v in raw.items()};r['active']=boolean(raw['active']);r['parent_dispatch']=boolean(raw['parent_dispatch']);r['effective_from']=to_date(raw['effective_from']);r['predecessors']=[s.strip() for s in text(raw['predecessors']).split(';') if s.strip()]
                if not r['reason'] or not r['product'] or not r['stage_code'] or not r['stage_name']:raise ValueError('Product, stage code, stage name and reason are required')
                for field,limit in [('stage_code',80),('stage_name',180),('variant',120),('branch',180),('vendor',180),('company',180),('source_column',10)]:
                    if len(r[field])>limit:raise ValueError(field+' is too long')
                groups.setdefault(r['product'],[]).append((n,r))
            except ValueError as e:stats['errors'].append(f'Row {n}: {e}')
        for name,group in groups.items():
            try:apply_definition(db,group,stats)
            except ValueError as e:stats['errors'].append(f'Rows {group[0][0]}–{group[-1][0]} ({name}): {e}')
    else:
        seen=set();contexts=set()
        for n,r in rows:
            try:
                key=(text(r['product']),text(r['stage_code']),text(r.get('date') if kind=='stage-daily' else r.get('month')))
                if key in seen:raise ValueError('Duplicate product/stage/date or month row')
                seen.add(key)
                if kind=='stage-schedules':apply_allocation(db,r,stats)
                else:
                    apply_actual(db,r,stats);contexts.add((product_for(db,text(r['product'])).id,to_date(r['date'])))
            except (ValueError,KeyError) as e:stats['errors'].append(f'Row {n}: {e}')
        if kind=='stage-daily':
            try:check_dispatch_totals(db,contexts)
            except ValueError as e:stats['errors'].append(str(e))
    return stats
