from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
import hashlib
import re

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import (
    Customer, DailyMIS, Machine, MachineShiftProduction, Operation, Product,
    ProcessDailySummary, QualityPhenomenon, QualityRejectionDaily,
    QualityRejectionImportBatch, QualityRejectionMonthlyHistory,
    RouteOperation, RouteVersion,
)


def _norm(value) -> str:
    return " ".join(str(value or "").replace("\n", " ").strip().lower().split())


def _code(value: str, prefix: str = "Q") -> str:
    text = re.sub(r"[^A-Za-z0-9]+", "_", str(value or "").strip().upper()).strip("_")
    return (text[:70] or prefix)[:80]


def _to_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _next_month(value: date) -> date:
    return (value.replace(day=28) + timedelta(days=4)).replace(day=1)


def _is_dispatch_denominator(value) -> bool:
    text = str(value or "").strip().upper().replace(" ", "_")
    return "DISP" in text or "DISPATCH" in text


def _historical_mis_month_actual(db: Session, product_id: int, month: date, cache: dict[tuple[int, date], Decimal]) -> Decimal | None:
    key = (product_id, month.replace(day=1))
    if key in cache:
        return cache[key] if cache[key] > 0 else None
    start = key[1]
    end = _next_month(start)
    rows = db.scalars(select(DailyMIS).where(
        DailyMIS.product_id == product_id, DailyMIS.mis_date >= start, DailyMIS.mis_date < end
    )).all()
    total = sum((Decimal(str(x.actual_qty or 0)) for x in rows), Decimal("0"))
    cache[key] = total
    return total if total > 0 else None


def _dec(value, default: Decimal = Decimal("0")) -> Decimal:
    if value in (None, ""):
        return default
    try:
        return Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError, AttributeError):
        return default


def _headers(ws) -> dict[str, int]:
    return {_norm(c.value): c.column for c in ws[1] if c.value not in (None, "")}


def _cell(ws, row: int, headers: dict[str, int], *names: str):
    for name in names:
        c = headers.get(_norm(name))
        if c:
            return ws.cell(row, c).value
    return None


def _find_product(db: Session, name: str) -> Product | None:
    n = _norm(name)
    if not n:
        return None
    rows = db.scalars(select(Product)).all()
    return next((p for p in rows if _norm(p.name) == n or _norm(p.code) == n), None)


def _find_machine(db: Session, value: str | None) -> Machine | None:
    n = _norm(value)
    if not n:
        return None
    rows = db.scalars(select(Machine)).all()
    return next((m for m in rows if _norm(m.name) == n or _norm(m.code) == n), None)


def _route_for_date(db: Session, product_id: int, on_date: date) -> RouteVersion | None:
    q = (
        select(RouteVersion)
        .where(RouteVersion.product_id == product_id, RouteVersion.effective_from <= on_date)
        .where((RouteVersion.effective_to.is_(None)) | (RouteVersion.effective_to >= on_date))
        .order_by(RouteVersion.revision_no.desc(), RouteVersion.effective_from.desc())
    )
    route = db.scalar(q.limit(1))
    if route:
        return route
    return db.scalar(
        select(RouteVersion)
        .where(RouteVersion.product_id == product_id)
        .order_by(RouteVersion.revision_no.desc(), RouteVersion.effective_from.desc())
        .limit(1)
    )


def _resolve_route_operation(db: Session, product: Product, on_date: date, label: str | None) -> RouteOperation | None:
    n = _norm(label)
    if not n:
        return None
    route = _route_for_date(db, product.id, on_date)
    if not route:
        return None
    rows = db.execute(
        select(RouteOperation, Operation)
        .join(Operation, Operation.id == RouteOperation.operation_id)
        .where(RouteOperation.route_version_id == route.id)
        .order_by(RouteOperation.sequence_no)
    ).all()
    if n in {"disp_done", "disp done", "dispatch", "dispatch done", "dispatched"}:
        return next((ro for ro, _ in rows if ro.is_dispatch), None)
    for ro, op in rows:
        if n in {_norm(op.name), _norm(op.code), _norm(ro.source_label)}:
            return ro
    return None


def _phenomenon(db: Session, value: str | None) -> QualityPhenomenon | None:
    n = _norm(value)
    if not n:
        return None
    return db.scalar(select(QualityPhenomenon).where(QualityPhenomenon.normalized_name == n))


def upsert_phenomenon(db: Session, name: str, *, group: str | None = None, default_team: str | None = None,
                      criticality: str | None = None) -> QualityPhenomenon:
    n = _norm(name)
    if not n:
        raise ValueError("Phenomenon name is required")
    row = db.scalar(select(QualityPhenomenon).where(QualityPhenomenon.normalized_name == n))
    if not row:
        code = _code(name, "PHEN")
        base = code
        i = 2
        while db.scalar(select(QualityPhenomenon).where(QualityPhenomenon.code == code)):
            code = f"{base[:70]}_{i}"
            i += 1
        row = QualityPhenomenon(code=code, name=" ".join(str(name).split()), normalized_name=n)
        db.add(row)
        # SessionLocal uses autoflush=False. Flush immediately so later rows in the
        # same workbook can see this pending phenomenon when checking both
        # normalized_name and generated code uniqueness. This prevents duplicate
        # master rows such as DAMAGE/Damage and safely suffixes true code collisions
        # such as OD - vs OD (-) (+).
        db.flush()
    if group not in (None, ""):
        row.phenomenon_group = str(group).strip()
    if default_team not in (None, ""):
        row.default_responsible_team = str(default_team).strip()
    if criticality not in (None, ""):
        row.criticality = str(criticality).strip().upper()
    return row



def import_phenomenon_master_workbook(db: Session, path) -> dict:
    """Import approved rejection phenomena from the standard master workbook."""
    wb = load_workbook(path, data_only=True, read_only=True)
    if "Phenomenon_Master" not in wb.sheetnames:
        return {"created": 0, "updated": 0, "unchanged": 0, "errors": ["Sheet 'Phenomenon_Master' is required"], "warnings": []}
    ws = wb["Phenomenon_Master"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"created": 0, "updated": 0, "unchanged": 0, "errors": ["Phenomenon_Master is empty"], "warnings": []}
    headers = {_norm(v).replace(" ", "_"): i for i, v in enumerate(rows[0]) if v not in (None, "")}
    aliases = {
        "name": ("phenomenon_name", "phenomenon"),
        "group": ("phenomenon_group", "group"),
        "team": ("responsible_team", "default_responsible_team"),
        "criticality": ("criticality",),
        "active": ("active", "is_active"),
    }
    def col(key):
        for name in aliases[key]:
            if name in headers:
                return headers[name]
        return None
    name_col = col("name")
    if name_col is None:
        return {"created": 0, "updated": 0, "unchanged": 0, "errors": ["Phenomenon Name column is required"], "warnings": []}
    group_col, team_col, criticality_col, active_col = col("group"), col("team"), col("criticality"), col("active")
    created = updated = unchanged = 0
    errors, warnings, seen = [], [], set()
    allowed_criticality = {"LOW", "NORMAL", "HIGH", "CRITICAL"}
    for excel_row, values in enumerate(rows[1:], 2):
        name = str(values[name_col] or "").strip() if name_col < len(values) else ""
        if not name:
            if any(v not in (None, "") for v in values):
                errors.append(f"Row {excel_row}: Phenomenon Name is required")
            continue
        normalized = _norm(name)
        if normalized in seen:
            errors.append(f"Row {excel_row}: duplicate Phenomenon Name '{name}' in workbook")
            continue
        seen.add(normalized)
        group = str(values[group_col] or "").strip() if group_col is not None and group_col < len(values) else ""
        team = str(values[team_col] or "").strip() if team_col is not None and team_col < len(values) else ""
        criticality = str(values[criticality_col] or "NORMAL").strip().upper() if criticality_col is not None and criticality_col < len(values) else "NORMAL"
        if criticality not in allowed_criticality:
            errors.append(f"Row {excel_row}: Criticality must be LOW, NORMAL, HIGH or CRITICAL")
            continue
        active_text = str(values[active_col] or "Yes").strip().lower() if active_col is not None and active_col < len(values) else "yes"
        if active_text not in {"yes", "no", "true", "false", "1", "0", "active", "inactive"}:
            errors.append(f"Row {excel_row}: Active must be Yes or No")
            continue
        is_active = active_text in {"yes", "true", "1", "active"}
        existing = db.scalar(select(QualityPhenomenon).where(QualityPhenomenon.normalized_name == normalized))
        before = None if not existing else (existing.name, existing.phenomenon_group, existing.default_responsible_team, existing.criticality, existing.is_active)
        row = upsert_phenomenon(db, name, group=group or None, default_team=team or None, criticality=criticality)
        row.is_active = is_active
        after = (row.name, row.phenomenon_group, row.default_responsible_team, row.criticality, row.is_active)
        if existing is None:
            created += 1
        elif before != after:
            updated += 1
        else:
            unchanged += 1
    return {"created": created, "updated": updated, "unchanged": unchanged, "errors": errors, "warnings": warnings}


def denominator_for_rejection(db: Session, *, on_date: date, shift: str, product: Product,
                              detection_ro: RouteOperation | None, machine: Machine | None) -> tuple[str, Decimal | None]:
    if detection_ro and machine:
        q = select(func.sum(MachineShiftProduction.total_count)).where(
            MachineShiftProduction.production_date == on_date,
            MachineShiftProduction.product_id == product.id,
            MachineShiftProduction.route_operation_id == detection_ro.id,
            MachineShiftProduction.machine_id == machine.id,
        )
        if shift and _norm(shift) not in {"general", "all", ""}:
            q = q.where(func.lower(MachineShiftProduction.shift) == shift.lower())
        total = db.scalar(q)
        if total is not None and Decimal(str(total)) > 0:
            return "MACHINE_PRODUCTION", Decimal(str(total))

    if detection_ro and not detection_ro.is_dispatch:
        pd = db.scalar(select(ProcessDailySummary).where(
            ProcessDailySummary.summary_date == on_date,
            ProcessDailySummary.product_id == product.id,
            ProcessDailySummary.route_operation_id == detection_ro.id,
        ))
        if pd and pd.actual_qty is not None and Decimal(str(pd.actual_qty)) > 0:
            return "OPERATION_ACTUAL", Decimal(str(pd.actual_qty))

    # Dispatch Done is the authoritative denominator for final rejection.
    if detection_ro is None or detection_ro.is_dispatch:
        mis = db.scalar(select(DailyMIS).where(DailyMIS.mis_date == on_date, DailyMIS.product_id == product.id))
        if mis and mis.actual_qty is not None and Decimal(str(mis.actual_qty)) > 0:
            return "DISP_DONE", Decimal(str(mis.actual_qty))

    return "NOT_AVAILABLE", None


def _daily_record_key(on_date: date, shift: str, product_id: int, det_id: int | None, resp_id: int | None,
                      machine_id: int | None, phenomenon_id: int) -> str:
    return "|".join([
        on_date.isoformat(), _norm(shift) or "general", str(product_id), str(det_id or 0),
        str(resp_id or 0), str(machine_id or 0), str(phenomenon_id),
    ])


def import_daily_rejection_workbook(db: Session, path: str | Path, *, batch_id: int | None = None,
                                    entered_by_id: int | None = None) -> dict:
    wb = load_workbook(path, data_only=False, read_only=False)
    sheet_name = "Daily_Rejection_Data"
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Required sheet '{sheet_name}' was not found")
    ws = wb[sheet_name]
    h = _headers(ws)
    required = ["Date", "Shift", "Product", "Detection Process", "Phenomenon", "Reject Qty"]
    missing = [x for x in required if _norm(x) not in h]
    if missing:
        raise ValueError("Missing columns: " + ", ".join(missing))

    stats = {"rows_read": 0, "created": 0, "updated": 0, "unchanged": 0, "errors": [], "warnings": [], "ppm_pending": 0}
    workbook_keys: dict[str, int] = {}
    direct_input_columns = [h.get(_norm(x)) for x in [
        "Date", "Shift", "Product", "Detection Process", "Responsible Process", "Machine",
        "Phenomenon", "Reject Qty", "Rework Qty", "Scrap Qty", "Remark", "Raise Action",
    ]]
    direct_input_columns = [c for c in direct_input_columns if c]
    for r in range(2, ws.max_row + 1):
        # Formula-assisted Plant/Customer/Type cells exist in every template row.
        # They do not make an otherwise blank row an import record.
        if all(ws.cell(r, c).value in (None, "") for c in direct_input_columns):
            continue
        stats["rows_read"] += 1
        on_date = _to_date(_cell(ws, r, h, "Date"))
        shift_raw = _cell(ws, r, h, "Shift")
        shift_map = {"a": "A", "b": "B", "c": "C", "general": "General"}
        shift = shift_map.get(_norm(shift_raw))
        product_name = _cell(ws, r, h, "Product", "Part Name")
        plant = _cell(ws, r, h, "Plant")
        if isinstance(plant, str) and plant.startswith("="):
            # Standard template auto-fills Plant from the Product Master. The
            # importer treats formula text as display assistance, not as a new master value.
            plant = None
        det_label = _cell(ws, r, h, "Detection Process", "Detection Operation")
        resp_label = _cell(ws, r, h, "Responsible Process", "Responsible Operation")
        machine_label = _cell(ws, r, h, "Machine")
        phen_name = _cell(ws, r, h, "Phenomenon", "Rejection Phenomenon")
        reject_raw = _cell(ws, r, h, "Reject Qty", "Rejection Qty")
        rework_raw = _cell(ws, r, h, "Rework Qty")
        scrap_raw = _cell(ws, r, h, "Scrap Qty")
        try:
            reject_qty = Decimal(str(reject_raw).replace(",", "").strip()) if reject_raw not in (None, "") else None
            rework_qty = Decimal(str(rework_raw).replace(",", "").strip()) if rework_raw not in (None, "") else Decimal("0")
            scrap_qty = Decimal(str(scrap_raw).replace(",", "").strip()) if scrap_raw not in (None, "") else Decimal("0")
        except (InvalidOperation, ValueError, AttributeError):
            stats["errors"].append(f"Row {r}: Reject, Rework and Scrap quantities must be numeric")
            continue
        remark = _cell(ws, r, h, "Remark", "Remarks")
        action_required_raw = _norm(_cell(ws, r, h, "Raise Action", "Action Required"))

        if not on_date:
            stats["errors"].append(f"Row {r}: invalid Date")
            continue
        if not shift:
            stats["errors"].append(f"Row {r}: Shift is required and must be A, B, C or General")
            continue
        product = _find_product(db, str(product_name or ""))
        if not product:
            stats["errors"].append(f"Row {r}: Product '{product_name}' is not in Product Master")
            continue
        if plant not in (None, "") and product.plant not in (None, "") and _norm(plant) != _norm(product.plant):
            stats["errors"].append(f"Row {r}: Plant '{plant}' does not match Product Master plant '{product.plant}' for {product.name}")
            continue
        detection_ro = _resolve_route_operation(db, product, on_date, str(det_label or ""))
        if not detection_ro:
            stats["errors"].append(f"Row {r}: Detection Process '{det_label}' is not in the active route for {product.name}")
            continue
        responsible_ro = None
        if resp_label not in (None, ""):
            responsible_ro = _resolve_route_operation(db, product, on_date, str(resp_label))
            if not responsible_ro:
                stats["errors"].append(f"Row {r}: Responsible Process '{resp_label}' is not in the active route for {product.name}")
                continue
        machine = _find_machine(db, str(machine_label or "")) if machine_label not in (None, "") else None
        if machine_label not in (None, "") and not machine:
            stats["errors"].append(f"Row {r}: Machine '{machine_label}' is not in Machine Master")
            continue
        phenomenon = _phenomenon(db, str(phen_name or ""))
        if not phenomenon:
            stats["errors"].append(f"Row {r}: Phenomenon '{phen_name}' is not in Quality Phenomenon Master. Add it once in the web app, then re-download the template.")
            continue
        if reject_qty is None or reject_qty <= 0:
            stats["errors"].append(f"Row {r}: Reject Qty is required and must be greater than zero")
            continue
        if rework_qty < 0 or scrap_qty < 0:
            stats["errors"].append(f"Row {r}: quantities cannot be negative")
            continue
        if action_required_raw not in {"", "yes", "no", "y", "n", "true", "false", "1", "0"}:
            stats["errors"].append(f"Row {r}: Raise Action must be Yes or No")
            continue

        denom_source, denom_qty = denominator_for_rejection(
            db, on_date=on_date, shift=shift, product=product, detection_ro=detection_ro, machine=machine
        )
        ppm = None
        if denom_qty is not None and denom_qty > 0:
            ppm = (reject_qty / denom_qty * Decimal("1000000")).quantize(Decimal("0.001"))
        else:
            stats["ppm_pending"] += 1
            stats["warnings"].append(f"Row {r}: denominator unavailable for {product.name} / {det_label} on {on_date}; rejection imported and PPM left pending")

        key = _daily_record_key(on_date, shift, product.id, detection_ro.id, responsible_ro.id if responsible_ro else None,
                                machine.id if machine else None, phenomenon.id)
        first_row = workbook_keys.get(key)
        if first_row is not None:
            stats["errors"].append(
                f"Rows {first_row} and {r}: duplicate daily rejection combination for "
                f"{product.name} / {det_label} / {phenomenon.name}. Combine the quantities into one row."
            )
            continue
        workbook_keys[key] = r
        row = db.scalar(select(QualityRejectionDaily).where(QualityRejectionDaily.record_key == key))
        desired = (
            reject_qty, rework_qty, scrap_qty, denom_source, denom_qty, ppm,
            str(remark).strip() if remark not in (None, "") else None,
            action_required_raw in {"yes", "y", "true", "1"},
        )
        if not row:
            row = QualityRejectionDaily(
                record_key=key, rejection_date=on_date, shift=shift, product_id=product.id,
                plant=product.plant, detection_route_operation_id=detection_ro.id,
                responsible_route_operation_id=responsible_ro.id if responsible_ro else None,
                responsible_team="Operation", machine_id=machine.id if machine else None,
                phenomenon_id=phenomenon.id, source="EXCEL", import_batch_id=batch_id, entered_by_id=entered_by_id,
            )
            db.add(row)
            stats["created"] += 1
        else:
            current = (row.reject_qty, row.rework_qty, row.scrap_qty, row.denominator_source, row.denominator_qty, row.ppm,
                       row.remark, row.action_required)
            if current == desired:
                stats["unchanged"] += 1
                continue
            stats["updated"] += 1
        row.reject_qty = reject_qty
        row.rework_qty = rework_qty
        row.scrap_qty = scrap_qty
        row.denominator_source = denom_source
        row.denominator_qty = denom_qty
        row.ppm = ppm
        row.remark = str(remark).strip() if remark not in (None, "") else None
        row.action_required = action_required_raw in {"yes", "y", "true", "1"}
        row.import_batch_id = batch_id or row.import_batch_id
        row.entered_by_id = entered_by_id or row.entered_by_id
    db.flush()
    return stats


def _create_product_from_additions(db: Session, wb, stats: dict) -> None:
    if "Product_Master_Additions" not in wb.sheetnames:
        return
    ws = wb["Product_Master_Additions"]
    h = _headers(ws)
    for r in range(2, ws.max_row + 1):
        name = _cell(ws, r, h, "Product")
        if name in (None, ""):
            continue
        if _find_product(db, str(name)):
            continue
        action = _norm(_cell(ws, r, h, "Action"))
        if "create" not in action:
            continue
        plant = _cell(ws, r, h, "Plant_INPUT", "Plant")
        group = _cell(ws, r, h, "Type_INPUT", "Type")
        customer_name = _cell(ws, r, h, "Customer")
        if plant in (None, "") or group in (None, ""):
            stats["warnings"].append(f"Product '{name}' not created: Plant and Type are required in Product_Master_Additions")
            continue
        customer = None
        if customer_name not in (None, ""):
            customer = next((c for c in db.scalars(select(Customer)).all() if _norm(c.name) == _norm(customer_name) or _norm(c.code) == _norm(customer_name)), None)
        code = _code(str(name), "PROD")[:60]
        base = code
        i = 2
        while db.scalar(select(Product).where(Product.code == code)):
            code = f"{base[:52]}_{i}"
            i += 1
        product = Product(code=code, name=" ".join(str(name).split()), customer_id=customer.id if customer else None,
                          plant=str(plant).strip(), product_group=str(group).strip(), sort_order=999, is_active=True)
        db.add(product)
        # Same-batch master additions must be visible before generating the next
        # product code because the application session intentionally disables
        # autoflush.
        db.flush()
        stats["products_created"] += 1
    db.flush()


def _import_phenomenon_master(db: Session, wb, stats: dict) -> None:
    name = "Phenomenon_Master_Import"
    if name not in wb.sheetnames:
        return
    ws = wb[name]
    h = _headers(ws)
    for r in range(2, ws.max_row + 1):
        phen = _cell(ws, r, h, "Phenomenon_WebApp", "Phenomenon_Source")
        if phen in (None, ""):
            continue
        before = _phenomenon(db, str(phen))
        upsert_phenomenon(
            db, str(phen), group=_cell(ws, r, h, "Phenomenon_Group"),
            default_team=_cell(ws, r, h, "Default_Responsible_Team"),
            criticality=_cell(ws, r, h, "Criticality"),
        )
        if not before:
            stats["phenomena_created"] += 1
    db.flush()


def import_historical_rejection_workbook(db: Session, path: str | Path, *, batch_id: int | None = None) -> dict:
    wb = load_workbook(path, data_only=False, read_only=False)
    if "Historical_Rejection_Import" not in wb.sheetnames:
        raise ValueError("Required sheet 'Historical_Rejection_Import' was not found")

    # Preflight the workbook business keys BEFORE writing any master or history rows.
    # SessionLocal intentionally uses autoflush=False, so two rows with the same
    # record_key in one upload would otherwise both appear absent until flush and
    # PostgreSQL would raise a UniqueViolation. More importantly, duplicate keys
    # with different quantities are ambiguous source data and must not be silently
    # overwritten. Give the user a clear validation error instead.
    ws_preflight = wb["Historical_Rejection_Import"]
    hp = _headers(ws_preflight)
    seen_keys: dict[str, int] = {}
    duplicate_keys: list[tuple[str, int, int]] = []
    for rr in range(2, ws_preflight.max_row + 1):
        raw_key = _cell(ws_preflight, rr, hp, "Record_Key")
        if raw_key in (None, ""):
            continue
        key = str(raw_key).strip()
        first_row = seen_keys.get(key)
        if first_row is None:
            seen_keys[key] = rr
        else:
            duplicate_keys.append((key, first_row, rr))
    if duplicate_keys:
        sample = "; ".join(
            f"rows {first}/{second}: {key}" for key, first, second in duplicate_keys[:8]
        )
        extra = len(duplicate_keys) - min(len(duplicate_keys), 8)
        suffix = f"; plus {extra} more" if extra else ""
        raise ValueError(
            f"Historical_Rejection_Import contains {len(duplicate_keys)} duplicate Record_Key values. "
            f"Each Record_Key must be unique in one workbook. {sample}{suffix}"
        )

    stats = {"rows_read": 0, "created": 0, "updated": 0, "unchanged": 0, "products_created": 0,
             "phenomena_created": 0, "dispatch_from_mis": 0, "ppm_pending": 0, "errors": [], "warnings": []}
    _create_product_from_additions(db, wb, stats)
    _import_phenomenon_master(db, wb, stats)

    dispatch_map: dict[str, Decimal] = {}
    if "Monthly_Dispatch_Qty_Input" in wb.sheetnames:
        ws_d = wb["Monthly_Dispatch_Qty_Input"]
        hd = _headers(ws_d)
        for r in range(2, ws_d.max_row + 1):
            key = _cell(ws_d, r, hd, "Denominator_Key")
            qty = _cell(ws_d, r, hd, "Dispatch_Qty_INPUT")
            if key not in (None, "") and qty not in (None, ""):
                dispatch_map[str(key).strip()] = _dec(qty)

    ws = wb["Historical_Rejection_Import"]
    h = _headers(ws)
    mis_dispatch_cache: dict[tuple[int, date], Decimal] = {}
    for r in range(2, ws.max_row + 1):
        record_key = _cell(ws, r, h, "Record_Key")
        if record_key in (None, ""):
            continue
        stats["rows_read"] += 1
        month = _to_date(_cell(ws, r, h, "Month"))
        product_name = _cell(ws, r, h, "WebApp_Product")
        phen_name = _cell(ws, r, h, "Phenomenon_WebApp", "Phenomenon_Source")
        if not month:
            stats["errors"].append(f"Row {r}: invalid Month")
            continue
        month = month.replace(day=1)
        product = _find_product(db, str(product_name or ""))
        if not product:
            stats["errors"].append(f"Row {r}: Product '{product_name}' is not in Product Master")
            continue
        phenomenon = _phenomenon(db, str(phen_name or ""))
        if not phenomenon:
            stats["errors"].append(f"Row {r}: Phenomenon '{phen_name}' is not in Quality Phenomenon Master")
            continue
        source = str(_cell(ws, r, h, "Source") or "").strip() or None
        source_sheet = str(_cell(ws, r, h, "Source_Sheet") or "").strip() or None
        record_scope = str(_cell(ws, r, h, "Record_Scope") or "PRODUCT_TOTAL").strip()
        include_raw = _norm(_cell(ws, r, h, "Include_In_Overall_Aggregate"))
        include = include_raw not in {"no", "n", "false", "0"}
        denominator_source = str(_cell(ws, r, h, "Denominator_Source") or "DISP_DONE").strip()

        dispatch_qty = _cell(ws, r, h, "Dispatch_Qty")
        if dispatch_qty in (None, "") and source and source_sheet:
            dkey = f"{source}|{source_sheet}|{month:%Y-%m}"
            dispatch_qty = dispatch_map.get(dkey)
        dispatch_dec = _dec(dispatch_qty) if dispatch_qty not in (None, "") else None

        # When the historical rejection denominator is Dispatch Done and the row is an
        # aggregate/product-total record, use the matching Historical Daily MIS Actual Qty
        # automatically. This removes the need for a second manual Dispatch Qty upload.
        scope_key = record_scope.upper()
        if (dispatch_dec is None or dispatch_dec <= 0) and include and scope_key in {"PRODUCT_TOTAL", "AGGREGATE_TOTAL", "TOTAL", ""} and _is_dispatch_denominator(denominator_source):
            mis_qty = _historical_mis_month_actual(db, product.id, month, mis_dispatch_cache)
            if mis_qty is not None and mis_qty > 0:
                dispatch_dec = mis_qty
                stats["dispatch_from_mis"] += 1

        reject_qty = _dec(_cell(ws, r, h, "Reject_Qty"))
        ppm = None
        if dispatch_dec is not None and dispatch_dec > 0:
            ppm = (reject_qty / dispatch_dec * Decimal("1000000")).quantize(Decimal("0.001"))
        else:
            stats["ppm_pending"] += 1
        row = db.scalar(select(QualityRejectionMonthlyHistory).where(QualityRejectionMonthlyHistory.record_key == str(record_key)))
        desired = (
            reject_qty, dispatch_dec, ppm, include,
            str(_cell(ws, r, h, "Data_Quality_Note") or "").strip() or None,
        )
        if not row:
            row = QualityRejectionMonthlyHistory(record_key=str(record_key), month=month, product_id=product.id,
                                                 phenomenon_id=phenomenon.id, import_batch_id=batch_id)
            db.add(row)
            stats["created"] += 1
        else:
            current = (row.reject_qty, row.dispatch_qty, row.ppm, row.include_in_aggregate, row.data_quality_note)
            if current == desired:
                stats["unchanged"] += 1
                continue
            stats["updated"] += 1
        row.source = source
        row.source_sheet = source_sheet
        row.record_scope = record_scope
        row.include_in_aggregate = include
        row.plant = str(_cell(ws, r, h, "Plant") or product.plant or "").strip() or None
        row.detection_operation = str(_cell(ws, r, h, "Detection_Operation") or "Disp_Done").strip()
        row.responsible_team = str(_cell(ws, r, h, "Responsible_Team") or "Operation").strip()
        row.responsible_operation = str(_cell(ws, r, h, "Responsible_Operation") or "").strip() or None
        row.machine = str(_cell(ws, r, h, "Machine") or "").strip() or None
        row.reject_qty = reject_qty
        row.denominator_source = denominator_source
        row.dispatch_qty = dispatch_dec
        row.ppm = ppm
        row.source_production_qty = _dec(_cell(ws, r, h, "Source_Production_Qty")) if _cell(ws, r, h, "Source_Production_Qty") not in (None, "") else None
        row.source_inspection_qty = _dec(_cell(ws, r, h, "Source_Inspection_Qty")) if _cell(ws, r, h, "Source_Inspection_Qty") not in (None, "") else None
        row.source_total_reject_qty = _dec(_cell(ws, r, h, "Source_Total_Reject_Qty")) if _cell(ws, r, h, "Source_Total_Reject_Qty") not in (None, "") else None
        row.phenomenon_sum_reject_qty = _dec(_cell(ws, r, h, "Phenomenon_Sum_Reject_Qty")) if _cell(ws, r, h, "Phenomenon_Sum_Reject_Qty") not in (None, "") else None
        row.source_reported_ppm = _dec(_cell(ws, r, h, "Source_Reported_PPM")) if _cell(ws, r, h, "Source_Reported_PPM") not in (None, "") else None
        row.data_quality_note = str(_cell(ws, r, h, "Data_Quality_Note") or "").strip() or None
        row.import_batch_id = batch_id or row.import_batch_id
    db.flush()
    return stats


def sha256_file(path: str | Path) -> str:
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()
