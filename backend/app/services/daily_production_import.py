"""One atomic daily import for customer MIS and explicit stage actuals."""
from openpyxl import load_workbook

from .historical_mis_import import SHEET_NAME, REQUIRED_HEADERS, import_historical_daily_mis
from .process_flows import read_rows, product_for, stage_for, to_date, quantity, text, import_process_workbook


def import_daily_production(db, path):
    # Read and reconcile both sheets before any business writes. The API owns
    # rollback/commit, including month authorization and audit evidence.
    errors = []
    stages = read_rows(path, 'stage-daily')
    parents = {}
    for number, row in stages:
        try:
            product = product_for(db, text(row['product']))
            day = to_date(row['date'])
            stage = stage_for(db, product, text(row['stage_code']), day)
            actual = quantity(row['actual_qty'], 'actual_qty')
            reject = quantity(row['reject_qty'], 'reject_qty')
            if reject > actual:
                raise ValueError('Reject quantity cannot exceed actual quantity')
            if stage.role in {'DISPATCH', 'DISPATCH_TOTAL', 'DISPATCH_DETAIL'} and reject:
                raise ValueError('Dispatch must contain accepted pieces, with zero rejects')
            if stage.parent_dispatch:
                key = (product.id, day)
                if key in parents:
                    raise ValueError('Duplicate parent dispatch product/date')
                parents[key] = (actual, text(row['reason']), number)
        except (ValueError, KeyError) as exc:
            errors.append(f'Daily_Actuals row {number}: {exc}')

    wb = load_workbook(path, data_only=False, read_only=True)
    mis_keys = set()
    reasons = {}
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError(f'Missing sheet {SHEET_NAME}; export both daily tabs together')
        ws = wb[SHEET_NAME]
        headers = [str(v).strip() if v is not None else '' for v in next(ws.iter_rows(values_only=True))]
        populated = [h for h in headers if h]
        if len(populated) != len(set(populated)):
            raise ValueError(f'{SHEET_NAME}: duplicate headers')
        if missing := REQUIRED_HEADERS - set(headers):
            raise ValueError(f'{SHEET_NAME}: missing headers: {", ".join(sorted(missing))}')
        for number, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if all(v in (None, '') for v in values):
                continue
            try:
                if any(isinstance(v, str) and v.startswith('=') for v in values):
                    raise ValueError('Use the exported values-only workbook, not the formula working workbook')
                row = dict(zip(headers, values))
                product = product_for(db, text(row['Product']))
                if text(row['Product']).lower() != product.name.strip().lower():
                    raise ValueError('Use the exact Product master name for customer MIS')
                day = to_date(row['Date'])
                key = (product.id, day)
                quantity(row['Plan_Qty'], 'Plan_Qty')
                actual = quantity(row['Actual_Qty'], 'Actual_Qty')
                if key in mis_keys:
                    raise ValueError('Duplicate customer MIS product/date')
                mis_keys.add(key)
                if key not in parents:
                    raise ValueError('Matching parent dispatch is missing from Daily_Actuals')
                expected, reason, stage_number = parents[key]
                if actual != expected:
                    raise ValueError(f'Actual_Qty {actual} differs from parent dispatch {expected} in Daily_Actuals row {stage_number}')
                reasons[key] = reason
            except (ValueError, KeyError) as exc:
                errors.append(f'{SHEET_NAME} row {number}: {exc}')
        if not mis_keys:
            errors.append(f'{SHEET_NAME}: no customer MIS rows')
        for key, (_, _, number) in parents.items():
            if key not in mis_keys:
                errors.append(f'Daily_Actuals row {number}: matching customer MIS row is missing')
    finally:
        wb.close()
    if errors:
        return {'created': 0, 'updated': 0, 'unchanged': 0, 'errors': errors, 'warnings': []}

    mis = import_historical_daily_mis(db, path, row_reasons=reasons)
    stage_stats = import_process_workbook(db, path, 'stage-daily')
    return {
        'created': mis['mis_created'] + stage_stats['new'],
        'updated': mis['mis_updated'] + stage_stats['updated'],
        'unchanged': mis['mis_unchanged'] + stage_stats['unchanged'],
        'errors': stage_stats['errors'],
        'warnings': list(dict.fromkeys(mis['warnings'] + stage_stats['warnings'])),
        'customer_mis': mis,
        'stage_actuals': stage_stats,
    }
