"""Collect row errors before invoking historical importers."""
from decimal import Decimal, InvalidOperation
from openpyxl import load_workbook
from sqlalchemy import select
from ..models import Product
from .historical_mis_import import _to_date

def validate_historical_rows(db,path,kind):
    if kind not in {'historical-daily-mis','historical-sales-prices'}:return []
    wb=load_workbook(path,data_only=True,read_only=True)
    sheet='Historical_Daily_MIS_Import' if kind=='historical-daily-mis' else 'Historical_Sales_Price_Import'
    if sheet not in wb.sheetnames:wb.close();return [f'Missing sheet: {sheet}']
    ws=wb[sheet];headers={str(v).strip().lower().replace(' ','_'):i for i,v in enumerate(next(ws.iter_rows(values_only=True))) if v is not None}
    names={p.name.strip().lower():p.id for p in db.scalars(select(Product))};codes={p.code.strip().lower():p.id for p in db.scalars(select(Product))}
    errors=[];seen={}
    def get(row,*aliases):
        for alias in aliases:
            idx=headers.get(alias)
            if idx is not None:return row[idx] if idx<len(row) else None
    for number,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
        if all(x in (None,'') for x in row):continue
        name=str(get(row,'product') or '').strip().lower();pid=names.get(name) or (codes.get(name) if kind.endswith('prices') else None)
        if not pid:errors.append(f'Row {number}: Product is missing or does not match the current master')
        d=_to_date(get(row,'date') if kind.endswith('mis') else get(row,'effective_from'))
        if not d:errors.append(f'Row {number}: Invalid date')
        values=[]
        columns=[('plan_qty',),('actual_qty',)] if kind.endswith('mis') else [('sales_price_rs_per_pc','price')]
        for aliases in columns:
            raw=get(row,*aliases)
            try:
                value=Decimal(str(raw if raw not in (None,'') else 0))
                if not value.is_finite() or value<0 or (kind.endswith('prices') and value==0):raise ValueError()
                values.append(value)
            except (ValueError,InvalidOperation):errors.append(f'Row {number}: Invalid {aliases[0]}')
        key=(pid,d)
        if pid and d and key in seen:
            errors.append(f'Row {number}: Duplicate product/date key; conflicts with row {seen[key]}')
        elif pid and d:seen[key]=number
    wb.close()
    return errors
