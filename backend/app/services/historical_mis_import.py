from __future__ import annotations

from collections import defaultdict
from contextlib import closing
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import DailyMIS, DailyRequirement, Product, SourceType
from .pricing import price_for_date

SHEET_NAME = "Historical_Daily_MIS_Import"
REQUIRED_HEADERS = {"Date", "Product", "Plan_Qty", "Actual_Qty"}


def _to_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    return None


def _decimal(value) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    return Decimal(str(value))


def import_historical_daily_mis(db: Session, file_path: str | Path, *, row_reasons: dict | None = None) -> dict:
    original_reason = db.info.get('reason')
    try:
        with closing(load_workbook(file_path, data_only=True, read_only=True)) as wb:
            return _import_historical_daily_mis(db, wb, row_reasons=row_reasons, request_reason=original_reason)
    finally:
        db.info['reason'] = original_reason


def _import_historical_daily_mis(db: Session, wb, *, row_reasons=None, request_reason=None) -> dict:
    """Import the one-time normalized May-Aug historical Daily MIS sheet.

    The sheet deliberately uses current Product master names. Unknown products are
    rejected rather than auto-created so historical imports cannot create duplicate
    masters through spelling differences.
    """
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Workbook must contain sheet '{SHEET_NAME}'.")
    ws = wb[SHEET_NAME]
    headers = {str(cell.value).strip(): idx for idx, cell in enumerate(ws[1]) if cell.value not in (None, "")}
    missing = REQUIRED_HEADERS - set(headers)
    if missing:
        raise ValueError(f"{SHEET_NAME} is missing required columns: {', '.join(sorted(missing))}")

    products = {p.name.strip().lower(): p for p in db.scalars(select(Product)).all()}
    stats = {
        "rows_read": 0,
        "mis_created": 0,
        "mis_updated": 0,
        "mis_unchanged": 0,
        "requirements_created": 0,
        "requirements_updated": 0,
        "requirements_unchanged": 0,
        "unknown_products": [],
        "price_warnings": [],
        "min_date": None,
        "max_date": None,
    }
    price_warned: set[int] = set()

    def val(row, name):
        idx = headers.get(name)
        return row[idx] if idx is not None and idx < len(row) else None

    for raw in ws.iter_rows(min_row=2, values_only=True):
        d = _to_date(val(raw, "Date"))
        product_name = str(val(raw, "Product") or "").strip()
        if not d and not product_name:
            continue
        if not d:
            raise ValueError(f"Invalid Date in row {stats['rows_read'] + 2}.")
        if not product_name:
            raise ValueError(f"Missing Product in row {stats['rows_read'] + 2}.")
        stats["rows_read"] += 1

        product = products.get(product_name.lower())
        if product is None:
            if product_name not in stats["unknown_products"]:
                stats["unknown_products"].append(product_name)
            continue

        plan = _decimal(val(raw, "Plan_Qty"))
        actual = _decimal(val(raw, "Actual_Qty"))
        if plan < 0 or actual < 0:
            raise ValueError(f"Negative quantity is not allowed for {product_name} on {d}.")

        if row_reasons is not None:
            # Set before queries: SQLAlchemy may autoflush while finding the
            # matching parent requirement, before the explicit flush below.
            db.info['reason'] = row_reasons.get((product.id, d)) or request_reason

        price = price_for_date(db, product.id, d, Decimal("0"))
        if price == 0 and product.id not in price_warned:
            stats["price_warnings"].append(
                f"{product.name}: no effective sales price exists for part of the historical period; sales values are 0 until a historical Sales Price Revision is entered."
            )
            price_warned.add(product.id)

        mis = db.scalar(select(DailyMIS).where(DailyMIS.mis_date == d, DailyMIS.product_id == product.id))
        desired = (plan, actual, price, plan * price, actual * price)
        if mis is None:
            mis = DailyMIS(
                mis_date=d,
                product_id=product.id,
                plan_qty=plan,
                actual_qty=actual,
                sales_price=price,
                plan_sales=plan * price,
                actual_sales=actual * price,
                source=SourceType.EXCEL,
                remark="Historical Daily MIS one-time import",
            )
            db.add(mis)
            stats["mis_created"] += 1
        else:
            current = (
                Decimal(str(mis.plan_qty or 0)), Decimal(str(mis.actual_qty or 0)), Decimal(str(mis.sales_price or 0)),
                Decimal(str(mis.plan_sales or 0)), Decimal(str(mis.actual_sales or 0)),
            )
            if current == desired:
                stats["mis_unchanged"] += 1
            else:
                if row_reasons is not None and not db.info.get('reason'):
                    raise ValueError(f'{product.name} on {d}: changing an existing customer MIS plan, dispatch or price requires a correction reason')
                mis.plan_qty, mis.actual_qty, mis.sales_price, mis.plan_sales, mis.actual_sales = desired
                mis.source = SourceType.EXCEL
                mis.remark = "Historical Daily MIS one-time import"
                stats["mis_updated"] += 1

        req = db.scalar(select(DailyRequirement).where(
            DailyRequirement.req_date == d,
            DailyRequirement.product_id == product.id,
            DailyRequirement.route_operation_id.is_(None),
        ))
        if req is None:
            req = DailyRequirement(
                req_date=d,
                product_id=product.id,
                route_operation_id=None,
                baseline_plan_qty=plan,
                revised_plan_qty=plan,
                is_frozen=True,
            )
            db.add(req)
            stats["requirements_created"] += 1
        else:
            current_req = (Decimal(str(req.baseline_plan_qty or 0)), Decimal(str(req.revised_plan_qty or 0)), bool(req.is_frozen))
            desired_req = (plan, plan, True)
            if current_req == desired_req:
                stats["requirements_unchanged"] += 1
            else:
                if row_reasons is not None and not db.info.get('reason'):
                    raise ValueError(f'{product.name} on {d}: changing an existing customer daily requirement requires a correction reason')
                req.baseline_plan_qty = plan
                req.revised_plan_qty = plan
                req.is_frozen = True
                stats["requirements_updated"] += 1

        stats["min_date"] = d if stats["min_date"] is None or d < stats["min_date"] else stats["min_date"]
        stats["max_date"] = d if stats["max_date"] is None or d > stats["max_date"] else stats["max_date"]

        if row_reasons is not None:
            db.flush()

    if stats["unknown_products"]:
        raise ValueError(
            "Unknown Product master names: " + ", ".join(stats["unknown_products"]) + ". Use current Product master names; no new product was created."
        )

    db.flush()
    stats["warnings"] = stats.pop("price_warnings")
    return stats
