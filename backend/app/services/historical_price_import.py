from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Product, SalesPriceHistory
from .pricing import normalize_price_ranges, recalculate_mis_sales_from


SHEET_NAME = "Historical_Sales_Price_Import"


def _norm(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _headers(ws) -> dict[str, int]:
    out: dict[str, int] = {}
    for idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1, values_only=True))):
        key = _norm(cell).replace(" ", "_")
        if key:
            out[key] = idx
    return out


def _pick(row: tuple, hdr: dict[str, int], *names: str):
    for name in names:
        idx = hdr.get(_norm(name).replace(" ", "_"))
        if idx is not None and idx < len(row):
            return row[idx]
    return None


def _date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d-%B-%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    raise ValueError(f"Invalid Effective_From date: {value!r}")


def _price(value: object) -> Decimal:
    if value is None or value == "":
        raise ValueError("Sales price is blank")
    try:
        p = Decimal(str(value).replace(",", "").replace("₹", "").strip())
    except (InvalidOperation, AttributeError) as exc:
        raise ValueError(f"Invalid sales price: {value!r}") from exc
    if p <= 0:
        raise ValueError(f"Sales price must be greater than zero: {value!r}")
    return p


def import_historical_sales_prices(db: Session, path: Path, entered_by_id: int | None = None) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Workbook must contain sheet '{SHEET_NAME}'")
    ws = wb[SHEET_NAME]
    hdr = _headers(ws)
    required = {
        "product": ["Product"],
        "effective_from": ["Effective_From", "Effective From"],
        "price": ["Sales_Price_Rs_Per_Pc", "Sales Price Rs Per Pc", "Price"],
    }
    missing = [label for label, aliases in required.items() if not any(_norm(a).replace(" ", "_") in hdr for a in aliases)]
    if missing:
        raise ValueError("Missing required columns: " + ", ".join(missing))

    products = db.scalars(select(Product)).all()
    by_name = {_norm(p.name): p for p in products}
    by_code = {_norm(p.code): p for p in products}

    stats = {
        "price_rows_created": 0,
        "price_rows_updated": 0,
        "price_rows_unchanged": 0,
        "rows_skipped_blank": 0,
        "rows_read": 0,
        "products_touched": 0,
        "mis_rows_recalculated": 0,
        "warnings": [],
    }
    seen: dict[tuple[int, date], Decimal] = {}
    touched: dict[int, date] = {}

    for excel_row, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        product_name = _pick(row, hdr, "Product")
        eff_raw = _pick(row, hdr, "Effective_From", "Effective From")
        price_raw = _pick(row, hdr, "Sales_Price_Rs_Per_Pc", "Sales Price Rs Per Pc", "Price")
        if all(v in (None, "") for v in (product_name, eff_raw, price_raw)):
            stats["rows_skipped_blank"] += 1
            continue
        stats["rows_read"] += 1

        product = by_name.get(_norm(product_name)) or by_code.get(_norm(product_name))
        if product is None:
            raise ValueError(f"Row {excel_row}: Product '{product_name}' is not in the current Product Master")
        eff = _date(eff_raw)
        price = _price(price_raw)
        key = (product.id, eff)
        if key in seen:
            if seen[key] != price:
                raise ValueError(
                    f"Row {excel_row}: conflicting duplicate for {product.name} on {eff}: {seen[key]} vs {price}"
                )
            stats["price_rows_unchanged"] += 1
            continue
        seen[key] = price

        rev = str(_pick(row, hdr, "Revision_Reference", "Revision Reference") or "").strip() or None
        reason = str(_pick(row, hdr, "Reason") or "").strip() or None
        source_doc = str(_pick(row, hdr, "Source_Document", "Source Document") or "").strip() or None

        existing = db.scalar(
            select(SalesPriceHistory)
            .where(SalesPriceHistory.product_id == product.id, SalesPriceHistory.effective_from == eff)
            .order_by(SalesPriceHistory.id.desc())
            .limit(1)
        )
        if existing is None:
            existing = SalesPriceHistory(
                product_id=product.id,
                effective_from=eff,
                price=price,
                reason=reason,
                revision_reference=rev,
                source_document=source_doc,
                source="HISTORICAL_IMPORT",
                entered_by_id=entered_by_id,
            )
            db.add(existing)
            db.flush()
            stats["price_rows_created"] += 1
        else:
            same = (
                Decimal(str(existing.price)) == price
                and (existing.reason or None) == reason
                and (existing.revision_reference or None) == rev
                and (existing.source_document or None) == source_doc
                and existing.source == "HISTORICAL_IMPORT"
            )
            if same:
                stats["price_rows_unchanged"] += 1
            else:
                existing.price = price
                existing.reason = reason
                existing.revision_reference = rev
                existing.source_document = source_doc
                existing.source = "HISTORICAL_IMPORT"
                existing.entered_by_id = entered_by_id
                stats["price_rows_updated"] += 1

        touched[product.id] = min(touched.get(product.id, eff), eff)

    for product_id, earliest in touched.items():
        normalize_price_ranges(db, product_id)
        stats["mis_rows_recalculated"] += recalculate_mis_sales_from(db, product_id, earliest)

    stats["products_touched"] = len(touched)
    db.flush()
    return stats
