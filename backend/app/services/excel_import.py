from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import get_settings
from ..enums import OperationType, SourceType
from .planning import rebuild_process_requirements
from .pricing import price_for_date

from ..models import (
    Customer, DailyMIS, DailyRequirement, Operation, Product, ProcessDailySummary,
    RouteOperation, RouteVersion, SalesPriceHistory, ScheduleRevision, Vendor,
    WorkingCalendar,
)


VENDOR_HINTS = ["jay", "aryan", "saptrishi", "saptarshi", "satara auto", "joytirling", "siddharth", "kr"]


def _norm(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())


def _code(value: str, max_len: int = 55) -> str:
    txt = re.sub(r"[^A-Z0-9]+", "_", value.upper()).strip("_")
    return txt[:max_len] or "ITEM"


def _to_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        # Excel serial date system used by this workbook.
        return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
    try:
        return datetime.fromisoformat(str(value)).date()
    except Exception:
        return None

def _value_decimal(value: Any) -> Decimal:
    return Decimal(str(value or 0))


def _header_map(ws) -> dict[str, int]:
    """Return normalized PBI_Products header -> zero-based tuple index."""
    return {_norm(cell.value): idx for idx, cell in enumerate(ws[1]) if cell.value not in (None, "")}


def _find_header(headers: dict[str, int], *aliases: str) -> int | None:
    for alias in aliases:
        key = _norm(alias)
        if key in headers:
            return headers[key]
    return None


def _row_value(row, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def classify_operation(label: str) -> OperationType:
    n = _norm(label)
    if any(x in n for x in ["dispatch", "dispdone"]):
        return OperationType.DISPATCH
    if any(x in n for x in ["inspection", "insp", "rt", "quality"]):
        return OperationType.INSPECTION
    if any(x in n for x in ["outward", "cpgout", "topowdercoating", "forcoating"]):
        return OperationType.VENDOR_OUT
    if any(x in n for x in ["inward", "cpgin", "fromcoating"]):
        return OperationType.VENDOR_IN
    if any(v.replace(" ", "") in n for v in VENDOR_HINTS):
        return OperationType.VENDOR_PROCESS
    if "powdercoating" in n or "coatingdone" in n:
        return OperationType.VENDOR_PROCESS
    if "packing" in n:
        return OperationType.PACKING
    if "customer" in n or "tvsin" in n:
        return OperationType.CUSTOMER_RECEIPT
    return OperationType.INTERNAL


def vendor_from_label(label: str) -> str | None:
    low = str(label or "").lower().replace("\n", " ")
    for hint in VENDOR_HINTS:
        if hint in low:
            return hint.title()
    if "powder coating" in low or "coating" in low:
        return "Coating Vendor"
    if "cpg" in low:
        return "CPG"
    return None


def get_or_create_customer(db: Session, name: str | None) -> Customer | None:
    if not name:
        return None
    code = _code(name, 35)
    row = db.scalar(select(Customer).where(Customer.code == code))
    if not row:
        row = Customer(code=code, name=name.strip())
        db.add(row)
        db.flush()
    return row


def get_or_create_product(db: Session, name: str, customer: Customer | None, actual_measure: str | None, sort_order: int) -> Product:
    code = _code(name)
    row = db.scalar(select(Product).where(Product.code == code))
    if not row:
        row = Product(code=code, name=" ".join(name.split()), customer_id=customer.id if customer else None,
                      actual_measure=actual_measure, sort_order=sort_order)
        db.add(row)
        db.flush()
    else:
        row.name = " ".join(name.split())
        row.customer_id = customer.id if customer else row.customer_id
        row.actual_measure = actual_measure or row.actual_measure
        row.sort_order = sort_order or row.sort_order
    return row


def get_or_create_operation(db: Session, label: str) -> Operation:
    code = _code(label, 70)
    row = db.scalar(select(Operation).where(Operation.code == code))
    if not row:
        row = Operation(code=code, name=" ".join(str(label).replace("\n", " ").split()), operation_type=classify_operation(label))
        db.add(row)
        db.flush()
    return row


def get_or_create_vendor(db: Session, name: str | None) -> Vendor | None:
    if not name:
        return None
    code = _code(name, 50)
    row = db.scalar(select(Vendor).where(Vendor.code == code))
    if not row:
        row = Vendor(code=code, name=name)
        db.add(row)
        db.flush()
    return row


def import_daily_production_workbook(db: Session, file_path: str | Path) -> dict:
    """Import the current Daily Production Review workbook.

    Supported sheets are the Power BI staging sheets created for this workbook:
    PBI_Products, PBI_Fact_Daily, PBI_Control and the existing Daily MIS.
    The importer is idempotent for product/date records and creates a route draft
    from the operation headers already present in Daily MIS.
    """
    wbv = load_workbook(file_path, data_only=True, read_only=False)
    required = {"PBI_Products", "PBI_Fact_Daily"}
    missing = required - set(wbv.sheetnames)
    if missing:
        raise ValueError(f"Workbook is missing required sheets: {', '.join(sorted(missing))}")

    stats = {
        "products_created_or_updated": 0,
        "mis_rows_imported": 0,
        "mis_created": 0,
        "mis_updated": 0,
        "mis_unchanged": 0,
        "schedule_revisions_created": 0,
        "routes_created": 0,
        "process_rows_imported": 0,
        "process_created": 0,
        "process_updated": 0,
        "process_unchanged": 0,
        "off_days_imported": 0,
        "warnings": [],
    }

    product_ws = wbv["PBI_Products"]
    as_of_date = None
    if "PBI_Control" in wbv.sheetnames:
        as_of_date = _to_date(wbv["PBI_Control"]["A2"].value)
    if as_of_date is None and "Daily MIS" in wbv.sheetnames:
        as_of_date = _to_date(wbv["Daily MIS"]["B2"].value)

    products_by_name: dict[str, Product] = {}
    product_meta: dict[str, dict] = {}

    # PBI_Products is treated as a master table by header name rather than fixed
    # column positions. This means Plant / Product Group can be inserted or moved
    # without changing the importer.
    headers = _header_map(product_ws)

    # Warn when a column contains values but has no header. It is safer to ignore
    # such a column than to guess that it is Plant / Group and corrupt filters.
    for col in range(1, product_ws.max_column + 1):
        if product_ws.cell(1, col).value in (None, ""):
            populated = [product_ws.cell(r, col).value for r in range(2, min(product_ws.max_row, 200) + 1)]
            if any(v not in (None, "") for v in populated):
                stats["warnings"].append(
                    f"PBI_Products column {product_ws.cell(1, col).column_letter} has data but no header; it was not imported as a master attribute."
                )

    idx_product = _find_header(headers, "Product", "Part", "Part Number", "Product Name")
    idx_customer = _find_header(headers, "Customer", "Customer Name")
    idx_actual_measure = _find_header(headers, "ActualMeasure", "Actual Measure")
    idx_plan_col = _find_header(headers, "PlanSourceCol", "Plan Source Col")
    idx_actual_col = _find_header(headers, "ActualSourceCol", "Actual Source Col")
    idx_sort = _find_header(headers, "SortOrder", "Sort Order")
    idx_price = _find_header(headers, "Sales Price (Rs/pc)", "Sales Price", "SalesPrice", "Price")
    idx_weight = _find_header(headers, "Finish Weight (Kg/pc)", "Finish Weight", "FinishWeight", "Weight Kg/pc")
    idx_plant = _find_header(headers, "Plant", "Plant Name", "Manufacturing Plant", "Factory", "Unit", "Location")
    idx_group = _find_header(headers, "Product Group", "ProductGroup", "Product Category", "Category", "Part Group", "Segment", "Business Group", "Bifurcation", "Type", "Material Type")

    if idx_product is None:
        raise ValueError("PBI_Products must contain a Product column.")
    if idx_plant is None:
        stats["warnings"].append("PBI_Products has no Plant header; plant filters will remain blank until that column is added.")
    if idx_group is None:
        stats["warnings"].append("PBI_Products has no Type / Product Group / Category header; group filters will remain blank until that column is added.")

    for row in product_ws.iter_rows(min_row=2, values_only=True):
        product_name = _row_value(row, idx_product)
        if not product_name:
            continue
        customer_name = _row_value(row, idx_customer)
        customer = get_or_create_customer(db, customer_name)
        actual_measure = _row_value(row, idx_actual_measure)
        sort_raw = _row_value(row, idx_sort)
        try:
            sort_order = int(sort_raw or 999)
        except Exception:
            sort_order = 999
        product = get_or_create_product(db, str(product_name), customer, actual_measure, sort_order)

        weight_raw = _row_value(row, idx_weight)
        if weight_raw not in (None, ""):
            try:
                product.finish_weight_kg = Decimal(str(weight_raw))
            except Exception:
                stats["warnings"].append(f"Invalid finish weight for {product.name}: {weight_raw}")

        plant_raw = _row_value(row, idx_plant)
        group_raw = _row_value(row, idx_group)
        product.plant = str(plant_raw).strip() if plant_raw not in (None, "") else None
        product.product_group = str(group_raw).strip() if group_raw not in (None, "") else None

        price_raw = _row_value(row, idx_price)
        try:
            price = Decimal(str(price_raw or 0))
        except Exception:
            price = Decimal("0")
            stats["warnings"].append(f"Invalid sales price for {product.name}: {price_raw}")

        product_meta[_norm(product_name)] = {
            "product": product,
            "plan_col": _row_value(row, idx_plan_col),
            "actual_col": _row_value(row, idx_actual_col),
            "price": price,
        }
        products_by_name[_norm(product_name)] = product
        stats["products_created_or_updated"] += 1
    db.flush()

    fact_ws = wbv["PBI_Fact_Daily"]
    min_date: date | None = None
    max_date: date | None = None
    for row in fact_ws.iter_rows(min_row=2, values_only=True):
        d = _to_date(row[0])
        pname = row[1]
        if not d or not pname:
            continue
        product = products_by_name.get(_norm(pname))
        if not product:
            stats["warnings"].append(f"Unknown product in PBI_Fact_Daily: {pname}")
            continue
        workbook_price = product_meta[_norm(pname)]["price"]
        # Once an effective-dated price history exists, it becomes authoritative.
        # Re-importing Excel therefore cannot silently undo a web-entered price revision.
        price = price_for_date(db, product.id, d, workbook_price)
        plan = Decimal(str(row[4] or 0))
        actual = Decimal(str(row[5] or 0))
        mis = db.scalar(select(DailyMIS).where(DailyMIS.mis_date == d, DailyMIS.product_id == product.id))
        desired = (plan, actual, price, plan * price, actual * price)
        if not mis:
            mis = DailyMIS(mis_date=d, product_id=product.id, source=SourceType.EXCEL)
            db.add(mis)
            stats["mis_created"] += 1
        else:
            current = (_value_decimal(mis.plan_qty), _value_decimal(mis.actual_qty), _value_decimal(mis.sales_price), _value_decimal(mis.plan_sales), _value_decimal(mis.actual_sales))
            if current == desired:
                stats["mis_unchanged"] += 1
            else:
                stats["mis_updated"] += 1
        mis.plan_qty = plan
        mis.actual_qty = actual
        mis.sales_price = price
        mis.plan_sales = plan * price
        mis.actual_sales = actual * price
        min_date = d if min_date is None or d < min_date else min_date
        max_date = d if max_date is None or d > max_date else max_date
        stats["mis_rows_imported"] += 1
    db.flush()

    if min_date:
        month = min_date.replace(day=1)
        # The workbook seeds the initial price only. After that, effective-dated
        # price history is authoritative so future Excel imports cannot overwrite
        # a manually approved price revision.
        for meta in product_meta.values():
            product = meta["product"]
            price = meta["price"]
            if price > 0:
                any_price = db.scalar(select(SalesPriceHistory).where(
                    SalesPriceHistory.product_id == product.id,
                ).order_by(SalesPriceHistory.effective_from).limit(1))
                if not any_price:
                    db.add(SalesPriceHistory(
                        product_id=product.id, effective_from=month, price=price,
                        reason="Initial price imported from PBI_Products", source="EXCEL",
                    ))
                else:
                    active = price_for_date(db, product.id, month, price)
                    if active != price:
                        stats["warnings"].append(
                            f"Workbook price for {product.name} ({price}) differs from effective price history ({active}); price history was kept. Use Sales Price Revision in the web app to change it."
                        )

    daily_ws = wbv["Daily MIS"] if "Daily MIS" in wbv.sheetnames else None
    if daily_ws and min_date:
        # Import explicit off-days from D2:M2; create the rest of the month's calendar
        # with Sun off / Mon-Sat working default first.
        first = min_date.replace(day=1)
        next_month = (first.replace(day=28) + timedelta(days=4)).replace(day=1)
        last = next_month - timedelta(days=1)
        off_days = set()
        for c in range(4, 14):
            od = _to_date(daily_ws.cell(2, c).value)
            if od:
                off_days.add(od)
        calendar_plants = {get_settings().plant_name}
        calendar_plants.update({
            str(meta["product"].plant).strip()
            for meta in product_meta.values()
            if meta["product"].plant not in (None, "")
        })
        for plant_name in sorted(calendar_plants):
            d = first
            while d <= last:
                row = db.scalar(select(WorkingCalendar).where(
                    WorkingCalendar.work_date == d,
                    WorkingCalendar.plant == plant_name,
                ))
                if not row:
                    is_working = d.weekday() != 6 and d not in off_days
                    row = WorkingCalendar(work_date=d, plant=plant_name,
                                          is_working_day=is_working,
                                          holiday_name="Imported Off Day" if d in off_days else None)
                    db.add(row)
                    if d in off_days:
                        stats["off_days_imported"] += 1
                d += timedelta(days=1)
        db.flush()

        # Daily requirement baseline/revised from the existing PBI plan. This freezes
        # the workbook's historical plan instead of re-creating history from today's target.
        for mis in db.scalars(select(DailyMIS).where(DailyMIS.mis_date >= first, DailyMIS.mis_date <= last)).all():
            req = db.scalar(select(DailyRequirement).where(
                DailyRequirement.req_date == mis.mis_date,
                DailyRequirement.product_id == mis.product_id,
                DailyRequirement.route_operation_id.is_(None),
            ))
            if not req:
                db.add(DailyRequirement(
                    req_date=mis.mis_date,
                    product_id=mis.product_id,
                    route_operation_id=None,
                    baseline_plan_qty=mis.plan_qty,
                    revised_plan_qty=mis.plan_qty,
                    is_frozen=bool(as_of_date and mis.mis_date < as_of_date),
                ))
        db.flush()

        # Build route drafts by finding each product name in row 3. Operations start
        # one column before the product-name cell and run through the configured
        # ActualSourceCol from PBI_Products.
        for pnorm, meta in product_meta.items():
            product: Product = meta["product"]
            product_col = None
            for c in range(1, daily_ws.max_column + 1):
                if _norm(daily_ws.cell(3, c).value) == pnorm:
                    product_col = c
                    break
            if not product_col:
                # Fuzzy contains handles minor punctuation/spacing variants such as KS.
                for c in range(1, daily_ws.max_column + 1):
                    valn = _norm(daily_ws.cell(3, c).value)
                    if valn and (valn in pnorm or pnorm in valn):
                        product_col = c
                        break
            if not product_col or not meta["actual_col"]:
                stats["warnings"].append(f"Could not infer route for {product.name}")
                continue
            try:
                end_col = column_index_from_string(str(meta["actual_col"]))
            except Exception:
                stats["warnings"].append(f"Invalid ActualSourceCol for {product.name}: {meta['actual_col']}")
                continue
            start_col = max(1, product_col - 1)
            labels: list[tuple[int, str]] = []
            for c in range(start_col, end_col + 1):
                label = daily_ws.cell(5, c).value
                if label is None or isinstance(label, (int, float)):
                    continue
                label = " ".join(str(label).replace("\n", " ").split())
                if not label:
                    continue
                # Stock and reason columns are not process operations.
                n = _norm(label)
                if "stock" in n or n.startswith("reasonfor"):
                    continue
                labels.append((c, label))
            if not labels:
                continue

            route = db.scalar(select(RouteVersion).where(RouteVersion.product_id == product.id, RouteVersion.revision_no == 0))
            if not route:
                route = RouteVersion(product_id=product.id, revision_no=0, effective_from=first,
                                     description="Imported route draft from Daily MIS")
                db.add(route)
                db.flush()
                for seq, (source_col, label) in enumerate(labels, start=10):
                    op = get_or_create_operation(db, label)
                    vname = vendor_from_label(label) if op.operation_type in {
                        OperationType.VENDOR_OUT, OperationType.VENDOR_PROCESS, OperationType.VENDOR_IN
                    } else None
                    vendor = get_or_create_vendor(db, vname)
                    ro = RouteOperation(
                        route_version_id=route.id,
                        operation_id=op.id,
                        sequence_no=seq * 10,
                        source_label=label,
                        vendor_id=vendor.id if vendor else None,
                        standard_yield=Decimal("1"),
                        standard_lead_time_days=0,
                        is_dispatch=(source_col == end_col or op.operation_type == OperationType.DISPATCH),
                    )
                    db.add(ro)
                db.flush()
                stats["routes_created"] += 1

            # Initial monthly target from row 4 at the configured final actual column.
            target_val = daily_ws.cell(4, end_col).value
            try:
                target = Decimal(str(target_val or 0))
            except Exception:
                target = Decimal("0")
            if target > 0:
                sched = db.scalar(select(ScheduleRevision).where(
                    ScheduleRevision.product_id == product.id,
                    ScheduleRevision.month == first,
                    ScheduleRevision.revision_no == 0,
                ))
                if not sched:
                    sched = ScheduleRevision(product_id=product.id, month=first, revision_no=0,
                                             effective_from=first, monthly_target_qty=target,
                                             reason="Imported opening monthly schedule")
                    db.add(sched)
                    db.flush()
                    stats["schedule_revisions_created"] += 1
                # Attach current revision to daily MIS/requirements where missing.
                for mis in db.scalars(select(DailyMIS).where(
                    DailyMIS.product_id == product.id,
                    DailyMIS.mis_date >= first,
                    DailyMIS.mis_date <= last,
                )).all():
                    mis.schedule_revision_id = mis.schedule_revision_id or sched.id
                for req in db.scalars(select(DailyRequirement).where(
                    DailyRequirement.product_id == product.id,
                    DailyRequirement.req_date >= first,
                    DailyRequirement.req_date <= last,
                    DailyRequirement.route_operation_id.is_(None),
                )).all():
                    req.schedule_revision_id = req.schedule_revision_id or sched.id

            # Import operation actuals from Daily MIS rows 7:37 using the route source labels.
            route_ops = db.scalars(select(RouteOperation).where(RouteOperation.route_version_id == route.id).order_by(RouteOperation.sequence_no)).all()
            label_map = {_norm(ro.source_label): ro for ro in route_ops if ro.source_label}
            for source_col, label in labels:
                ro = label_map.get(_norm(label))
                if not ro:
                    continue
                for r in range(7, 38):
                    d = _to_date(daily_ws.cell(r, 2).value)
                    if not d:
                        continue
                    val = daily_ws.cell(r, source_col).value
                    if val in (None, ""):
                        continue
                    try:
                        actual = Decimal(str(val))
                    except Exception:
                        continue
                    pd = db.scalar(select(ProcessDailySummary).where(
                        ProcessDailySummary.summary_date == d,
                        ProcessDailySummary.product_id == product.id,
                        ProcessDailySummary.route_operation_id == ro.id,
                    ))
                    if not pd:
                        pd = ProcessDailySummary(summary_date=d, product_id=product.id,
                                                 route_operation_id=ro.id, source=SourceType.EXCEL)
                        db.add(pd)
                        stats["process_created"] += 1
                    elif _value_decimal(pd.actual_qty) == actual and _value_decimal(pd.good_qty) == actual:
                        stats["process_unchanged"] += 1
                    else:
                        stats["process_updated"] += 1
                    pd.actual_qty = actual
                    pd.good_qty = actual
                    stats["process_rows_imported"] += 1

        # Generate process-level requirements from the imported dispatch requirement
        # and the newly inferred route. Imported dispatch history remains frozen;
        # only derived process requirement rows are created here.
        for meta in product_meta.values():
            try:
                rebuild_process_requirements(db, meta["product"].id, first, start_date=first)
            except Exception as exc:
                stats["warnings"].append(f"Process requirement build failed for {meta['product'].name}: {exc}")

    db.flush()
    return stats
