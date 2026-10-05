from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.main import app
from app.models import (
    CastingDefectDaily, CastingDefectPhenomenon, Customer, CustomerRejectionDaily,
    CustomerRejectionPhenomenon, Product, ProductValueAdditionHistory, Vendor,
)
from app.services.special_quality_import import (
    create_phenomenon, import_casting_defects_workbook, import_customer_rejections_workbook,
)


def auth(client):
    login = client.post('/api/auth/login', json={'username': 'admin', 'password': 'ChangeMe123!'})
    assert login.status_code == 200
    return {'Authorization': f"Bearer {login.json()['access_token']}"}


def seed():
    with SessionLocal() as db:
        customer = Customer(code='CUST-Q', name='Customer Quality Ltd')
        vendor = Vendor(code='FOUNDRY-Q', name='Approved Foundry')
        db.add_all([customer, vendor]); db.flush()
        product = Product(code='CAST-P', name='Casting Part', customer_id=customer.id, plant='Plant 1')
        db.add(product); db.flush()
        casting = create_phenomenon(db, CastingDefectPhenomenon, 'Blow Hole', group='Porosity')
        customer_ph = create_phenomenon(db, CustomerRejectionPhenomenon, 'Leakage', group='Functional')
        db.commit()
        return product.id, customer.id, vendor.id, casting.id, customer_ph.id


def test_manual_casting_and_customer_ppm_are_kept_separate():
    product_id, customer_id, vendor_id, casting_id, customer_ph_id = seed()
    client = TestClient(app); headers = auth(client)
    casting = client.post('/api/casting-quality/daily', headers=headers, json={
        'defect_date': '2026-10-05', 'shift': 'A', 'product_id': product_id, 'vendor_id': vendor_id,
        'heat_batch_no': 'H-100', 'phenomenon_id': casting_id, 'inspected_qty': 100,
        'defect_qty': 2, 'rework_qty': 1, 'scrap_qty': 1, 'remark': 'Incoming inspection',
    })
    assert casting.status_code == 200, casting.text
    assert casting.json()['ppm'] == 20000.0
    customer = client.post('/api/customer-quality/daily', headers=headers, json={
        'rejection_date': '2026-10-05', 'customer_id': customer_id, 'product_id': product_id,
        'reference_no': 'COM-01', 'batch_no': 'H-100', 'phenomenon_id': customer_ph_id,
        'dispatch_qty': 400, 'reject_qty': 2, 'remark': 'Customer complaint',
    })
    assert customer.status_code == 200, customer.text
    assert customer.json()['ppm'] == 5000.0
    with SessionLocal() as db:
        second = create_phenomenon(db, CastingDefectPhenomenon, 'Cold Shut', group='Filling')
        db.commit(); second_id = second.id
    second_casting = client.post('/api/casting-quality/daily', headers=headers, json={
        'defect_date': '2026-10-05', 'shift': 'A', 'product_id': product_id, 'vendor_id': vendor_id,
        'heat_batch_no': 'H-100', 'phenomenon_id': second_id, 'inspected_qty': 100,
        'defect_qty': 1, 'rework_qty': 1, 'scrap_qty': 0,
    })
    assert second_casting.status_code == 200, second_casting.text
    casting_report = client.get('/api/casting-quality/report?from_date=2026-10-01&to_date=2026-10-31', headers=headers).json()
    customer_report = client.get('/api/customer-quality/report?from_date=2026-10-01&to_date=2026-10-31', headers=headers).json()
    assert casting_report['summary']['inspected_qty'] == 100
    # The same inspection denominator is entered once per phenomenon row but
    # must be counted once in the weighted PPM total.
    assert casting_report['summary']['inspected_qty'] == 100
    assert casting_report['summary']['defect_qty'] == 3
    assert casting_report['summary']['ppm'] == 30000
    assert customer_report['summary']['dispatch_qty'] == 400
    assert customer_report['summary']['ppm'] == 5000
    with SessionLocal() as db:
        assert db.query(CastingDefectDaily).count() == 2
        assert db.query(CustomerRejectionDaily).count() == 1


def test_special_quality_quantity_validation():
    product_id, _, vendor_id, casting_id, _ = seed()
    client = TestClient(app); headers = auth(client)
    response = client.post('/api/casting-quality/daily', headers=headers, json={
        'defect_date': '2026-10-05', 'shift': 'General', 'product_id': product_id, 'vendor_id': vendor_id,
        'phenomenon_id': casting_id, 'inspected_qty': 10, 'defect_qty': 3, 'rework_qty': 2, 'scrap_qty': 2,
    })
    assert response.status_code == 422
    assert 'Rework Qty plus Scrap Qty' in response.text


def test_controlled_templates_and_workbook_imports(tmp_path):
    seed()
    client = TestClient(app); headers = auth(client)
    for endpoint, sheet in (('/api/casting-quality/template', 'Casting_Defect_Data'), ('/api/customer-quality/template', 'Customer_Rejection_Data')):
        response = client.get(endpoint, headers=headers)
        assert response.status_code == 200
        wb = load_workbook(BytesIO(response.content), data_only=False)
        assert wb['Masters'].sheet_state == 'hidden'
        assert wb[sheet].protection.sheet is True
        assert wb[sheet]['A2'].protection.locked is False
        assert wb[sheet]['A2'].fill.fgColor.rgb.endswith('FFF2CC')

    casting_path = tmp_path / 'casting.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Casting_Defect_Data'
    ws.append(['Date','Shift','Product','Vendor','Heat/Batch No','Inspected Qty','Phenomenon','Defect Qty','Rework Qty','Scrap Qty','Remark'])
    ws.append([date(2026,10,5),'B','Casting Part','Approved Foundry','H-200',200,'Blow Hole',4,3,1,'Checked'])
    wb.save(casting_path)
    customer_path = tmp_path / 'customer.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Customer_Rejection_Data'
    ws.append(['Date','Customer','Product','Reference No','Batch No','Dispatch Qty','Phenomenon','Reject Qty','Remark'])
    ws.append([date(2026,10,5),'Customer Quality Ltd','Casting Part','COM-02','H-200',1000,'Leakage',1,'Reported'])
    wb.save(customer_path)
    with SessionLocal() as db:
        assert import_casting_defects_workbook(db, casting_path)['created'] == 1
        assert import_customer_rejections_workbook(db, customer_path)['created'] == 1
        db.commit()
        assert db.scalar(select(CastingDefectDaily)).ppm == Decimal('20000.000')
        assert db.scalar(select(CustomerRejectionDaily)).ppm == Decimal('1000.000')


def test_value_addition_is_effective_dated_and_current_snapshot_is_safe():
    product_id, *_ = seed()
    client = TestClient(app); headers = auth(client)
    earlier_date = date.today() - timedelta(days=30)
    current_date = date.today() - timedelta(days=1)
    first = client.post(f'/api/masters/products/{product_id}/value-addition-revisions', headers=headers, json={
        'effective_from': earlier_date.isoformat(), 'value_addition_per_piece': 110, 'reason': 'Approved costing baseline',
    })
    assert first.status_code == 200, first.text
    assert first.json()['current_value_addition_per_piece'] == 110
    current = client.post(f'/api/masters/products/{product_id}/value-addition-revisions', headers=headers, json={
        'effective_from': current_date.isoformat(), 'value_addition_per_piece': 125.5, 'reason': 'Current commercial revision',
    })
    assert current.status_code == 200, current.text
    assert current.json()['current_value_addition_per_piece'] == 125.5
    products = client.get('/api/masters/products', headers=headers).json()
    assert next(p for p in products if p['id'] == product_id)['value_addition_per_piece'] == 125.5
    history = client.get(f'/api/masters/products/{product_id}/value-addition-history', headers=headers).json()
    assert [x['value_addition_per_piece'] for x in history] == [125.5, 110.0]
    with SessionLocal() as db:
        assert db.query(ProductValueAdditionHistory).count() == 2


def test_casting_workbook_uses_atomic_preview_and_confirm():
    seed()
    wb = Workbook(); ws = wb.active; ws.title = 'Casting_Defect_Data'
    ws.append(['Date','Shift','Product','Vendor','Heat/Batch No','Inspected Qty','Phenomenon','Defect Qty','Rework Qty','Scrap Qty'])
    ws.append([date(2026,10,5),'C','Casting Part','Approved Foundry','H-300',500,'Blow Hole',5,4,1])
    content = BytesIO(); wb.save(content)
    client = TestClient(app); headers = auth(client)
    preview = client.post('/api/import/preview/casting-daily', headers=headers,
                          files={'file': ('casting.xlsx', content.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert preview.status_code == 200, preview.text
    assert preview.json()['counts']['new'] == 1
    with SessionLocal() as db:
        assert db.query(CastingDefectDaily).count() == 0
    confirmed = client.post('/api/import/confirm', headers=headers, json={'preview_token': preview.json()['preview_token']})
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()['counts']['new'] == 1
    with SessionLocal() as db:
        assert db.query(CastingDefectDaily).count() == 1
