from datetime import date
from decimal import Decimal
from pathlib import Path
from io import BytesIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import event, select

from app.db import SessionLocal
from app.enums import OperationType, SourceType
from app.models import (
    Customer, DailyMIS, Operation, Product, QualityPhenomenon, QualityRejectionDaily,
    QualityRejectionMonthlyHistory, QualityHistoricalPpmProduction, RouteOperation, RouteVersion,
)
from app.services.quality_import import import_daily_rejection_workbook, import_historical_ppm_production_workbook, upsert_phenomenon
from app.api.quality import _daily_context, _serialize_daily, _history_dispatch_resolution
from app.api.import_preview import preview_error_message


def _headers_for_quality(client):
    login = client.post('/api/auth/login', json={'username':'admin','password':'ChangeMe123!'})
    assert login.status_code == 200
    return {'Authorization': f"Bearer {login.json()['access_token']}"}


def _seed_quality_context(db):
    c=Customer(code='Q-C',name='Quality Customer');db.add(c);db.flush()
    p=Product(code='Q-P',name='K70 Cylinder block',customer_id=c.id,plant='2020',product_group='CI');db.add(p);db.flush()
    op=Operation(code='DISP_DONE',name='Disp_Done',operation_type=OperationType.DISPATCH);db.add(op);db.flush()
    rv=RouteVersion(product_id=p.id,revision_no=0,effective_from=date(2026,9,1),is_active=True);db.add(rv);db.flush()
    ro=RouteOperation(route_version_id=rv.id,operation_id=op.id,sequence_no=10,is_dispatch=True,standard_yield=Decimal('1'));db.add(ro)
    mis=DailyMIS(mis_date=date(2026,9,30),product_id=p.id,plan_qty=Decimal('2100'),actual_qty=Decimal('2000'),sales_price=Decimal('1'),plan_sales=Decimal('2100'),actual_sales=Decimal('2000'),source=SourceType.MANUAL);db.add(mis)
    ph=upsert_phenomenon(db,'BORE O/S',default_team='Operation')
    db.commit()
    return p,ro,ph


def _make_daily_xlsx(path:Path,reject=15):
    wb=Workbook();ws=wb.active;ws.title='Daily_Rejection_Data'
    ws.append(['Date','Shift','Product','Plant','Customer','Type','Detection Process','Responsible Process','Machine','Phenomenon','Reject Qty','Rework Qty','Scrap Qty','Remark','Raise Action'])
    ws.append([date(2026,9,30),'A','K70 Cylinder block','2020','Quality Customer','CI','Disp_Done','',None,'BORE O/S',reject,5,10,'Daily check','Yes'])
    wb.save(path)


def test_daily_rejection_uses_dispatch_done_denominator(tmp_path):
    path=tmp_path/'daily_quality.xlsx';_make_daily_xlsx(path,15)
    with SessionLocal() as db:
        _seed_quality_context(db)
        stats=import_daily_rejection_workbook(db,path)
        db.commit()
        assert stats['created']==1
        row=db.scalar(select(QualityRejectionDaily))
        assert row.denominator_source=='DISP_DONE'
        assert row.denominator_qty==Decimal('2000')
        assert row.ppm==Decimal('7500.000')
        assert row.action_required is True


def test_daily_rejection_business_key_updates_instead_of_duplicate(tmp_path):
    path=tmp_path/'daily_quality.xlsx';_make_daily_xlsx(path,15)
    with SessionLocal() as db:
        _seed_quality_context(db)
        first=import_daily_rejection_workbook(db,path);db.commit()
        assert first['created']==1
        _make_daily_xlsx(path,18)
        second=import_daily_rejection_workbook(db,path);db.commit()
        assert second['updated']==1
        assert db.query(QualityRejectionDaily).count()==1
        row=db.scalar(select(QualityRejectionDaily))
        assert row.reject_qty==Decimal('18')
        assert row.ppm==Decimal('9000.000')


def test_daily_detail_queries_do_not_grow_per_rejection_row():
    with SessionLocal() as db:
        p, ro, ph = _seed_quality_context(db)
        db.add_all([
            QualityRejectionDaily(record_key=f'bulk-{i}', rejection_date=date(2026, 9, 30),
                                  product_id=p.id, phenomenon_id=ph.id,
                                  detection_route_operation_id=ro.id, reject_qty=Decimal('1'))
            for i in range(12)
        ])
        db.commit()
        rows = db.scalars(select(QualityRejectionDaily)).all()
        statements = []

        def track(_conn, _cursor, statement, _params, _context, _many):
            if statement.lstrip().upper().startswith('SELECT'):
                statements.append(statement)

        event.listen(db.bind, 'before_cursor_execute', track)
        try:
            context = _daily_context(db, rows)
            details = [_serialize_daily(row, context) for row in rows]
        finally:
            event.remove(db.bind, 'before_cursor_execute', track)
        assert len(details) == 12
        assert all(row['product'] == p.name and row['phenomenon'] == ph.name for row in details)
        assert len(statements) <= 7


def test_daily_rejection_template_uses_required_master_dropdowns_and_protected_formulas(tmp_path):
    from fastapi.testclient import TestClient
    from app.main import app

    with SessionLocal() as db:
        _seed_quality_context(db)

    client = TestClient(app)
    response = client.get('/api/quality/template', headers=_headers_for_quality(client))
    assert response.status_code == 200, response.text
    wb = load_workbook(BytesIO(response.content), data_only=False)
    ws = wb['Daily_Rejection_Data']
    masters = wb['Masters']

    assert masters.sheet_state == 'hidden'
    assert ws.protection.sheet is True
    assert ws['C2'].protection.locked is False
    assert ws['D2'].protection.locked is True
    assert ws['D2'].value.startswith('=IFERROR(VLOOKUP(')
    assert ws['C2'].fill.fgColor.rgb.endswith('FFF2CC')
    assert ws['C1'].fill.fgColor.rgb.endswith('C65911')
    assert wb.defined_names['ProductList'].attr_text.startswith("'Masters'!$A$2")
    assert wb.defined_names['PhenomenonList'].attr_text.startswith("'Masters'!$G$2")

    validations = {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}
    assert validations['C2:C501'].formula1 == '=ProductList'
    assert validations['C2:C501'].allow_blank is False
    assert validations['J2:J501'].formula1 == '=PhenomenonList'
    assert validations['J2:J501'].errorStyle == 'stop'
    assert validations['K2:K501'].operator == 'greaterThan'

    # Formula-assisted blank rows must not become 500 false import errors.
    path = tmp_path / 'daily_rejection_template.xlsx'
    path.write_bytes(response.content)
    with SessionLocal() as db:
        stats = import_daily_rejection_workbook(db, path)
        assert stats['rows_read'] == 0
        assert stats['errors'] == []


def test_daily_rejection_preview_rejects_invalid_shift_and_missing_reject_qty(tmp_path):
    path = tmp_path / 'daily_quality_required.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Daily_Rejection_Data'
    ws.append(['Date','Shift','Product','Detection Process','Phenomenon','Reject Qty'])
    ws.append([date(2026,9,30),'Night','K70 Cylinder block','Disp_Done','BORE O/S',5])
    ws.append([date(2026,9,30),'A','K70 Cylinder block','Disp_Done','BORE O/S',None])
    wb.save(path)

    with SessionLocal() as db:
        _seed_quality_context(db)
        stats = import_daily_rejection_workbook(db, path)
        assert stats['rows_read'] == 2
        assert any('Shift is required and must be A, B, C or General' in x for x in stats['errors'])
        assert any('Reject Qty is required and must be greater than zero' in x for x in stats['errors'])


def test_daily_rejection_preview_reports_duplicate_excel_rows(tmp_path):
    path = tmp_path / 'daily_quality_duplicate.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Daily_Rejection_Data'
    ws.append(['Date','Shift','Product','Detection Process','Responsible Process','Phenomenon','Reject Qty'])
    ws.append([date(2026,9,30),'A','K70 Cylinder block','Disp_Done','Disp_Done','BORE O/S',1])
    ws.append([date(2026,9,30),'A','K70 Cylinder block','Disp_Done','Disp_Done','BORE O/S',2])
    wb.save(path)

    with SessionLocal() as db:
        _seed_quality_context(db)
        stats = import_daily_rejection_workbook(db, path)
        db.flush()
        assert stats['created'] == 1
        assert len(stats['errors']) == 1
        assert 'Rows 2 and 3' in stats['errors'][0]
        assert 'Combine the quantities into one row' in stats['errors'][0]
        assert db.query(QualityRejectionDaily).count() == 1


def test_preview_error_message_does_not_render_empty_detail_list():
    class EmptyDetailError(Exception):
        detail = []

    assert preview_error_message(EmptyDetailError('database preview failed')) == 'database preview failed'


def test_phenomenon_upsert_dedupes_same_batch_and_suffixes_code_collisions():
    """Regression: SessionLocal has autoflush=False, so same-workbook masters must be flushed."""
    with SessionLocal() as db:
        p1 = upsert_phenomenon(db, 'DAMAGE', default_team='Operation')
        p2 = upsert_phenomenon(db, 'Damage', default_team='Operation')
        p3 = upsert_phenomenon(db, 'OD -', default_team='Operation')
        p4 = upsert_phenomenon(db, 'OD (-) (+)', default_team='Operation')
        db.commit()

        assert p1.id == p2.id
        assert p1.code == 'DAMAGE'
        assert p3.code == 'OD'
        assert p4.code == 'OD_2'
        rows = db.scalars(select(QualityPhenomenon)).all()
        assert len(rows) == 3


def test_historical_rejection_duplicate_record_keys_fail_preflight(tmp_path):
    from openpyxl import Workbook
    from app.services.quality_import import import_historical_rejection_workbook

    path = tmp_path / 'historical_dup.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Historical_Rejection_Import'
    ws.append([
        'Record_Key','Month','WebApp_Product','Phenomenon_WebApp','Reject_Qty',
        'Source','Source_Sheet','Record_Scope','Include_In_Overall_Aggregate',
        'Detection_Operation','Responsible_Team','Denominator_Source'
    ])
    key = 'HMCL|Total|2026-06|Damage'
    ws.append([key, date(2026,6,1), 'K70 Cylinder block', 'BORE O/S', 10, 'HMCL', 'Total', 'AGGREGATE_TOTAL', 'Yes', 'Disp_Done', 'Operation', 'DISP_DONE'])
    ws.append([key, date(2026,6,1), 'K70 Cylinder block', 'BORE O/S', 15, 'HMCL', 'Total', 'AGGREGATE_TOTAL', 'Yes', 'Disp_Done', 'Operation', 'DISP_DONE'])
    wb.save(path)

    with SessionLocal() as db:
        _seed_quality_context(db)
        try:
            import_historical_rejection_workbook(db, path)
            assert False, 'expected duplicate Record_Key validation error'
        except ValueError as exc:
            text = str(exc)
            assert 'duplicate Record_Key' in text
            assert key in text


def test_historical_rejection_visible_without_dispatch_qty():
    from fastapi.testclient import TestClient
    from app.main import app
    from app.models import QualityRejectionMonthlyHistory

    with SessionLocal() as db:
        p, _, ph = _seed_quality_context(db)
        db.add(QualityRejectionMonthlyHistory(
            record_key='HMCL|Total|2026-06|BORE O/S',
            month=date(2026,6,1), source='HMCL', source_sheet='Total', record_scope='AGGREGATE_TOTAL',
            include_in_aggregate=True, product_id=p.id, plant=None, detection_operation='Disp_Done',
            responsible_team='Operation', phenomenon_id=ph.id, reject_qty=Decimal('25'),
            denominator_source='DISP_DONE', dispatch_qty=None, ppm=None,
        ))
        db.commit()

    client = TestClient(app)
    login = client.post('/api/auth/login', json={'username':'admin','password':'ChangeMe123!'})
    assert login.status_code == 200
    h={'Authorization': f"Bearer {login.json()['access_token']}"}

    r = client.get('/api/quality/dashboard?from_date=2026-06-01&to_date=2026-06-30', headers=h)
    assert r.status_code == 200, r.text
    j = r.json()
    assert j['reject_qty'] == 25.0
    assert j['historical_records'] == 1
    assert j['denominator_qty'] == 0.0
    assert j['ppm'] is None

    pto = client.get('/api/quality/pareto?group_by=phenomenon&from_date=2026-06-01&to_date=2026-06-30', headers=h)
    assert pto.status_code == 200, pto.text
    assert pto.json()[0]['reject_qty'] == 25.0

    hist = client.get('/api/quality/history?from_date=2026-06-01&to_date=2026-06-30', headers=h)
    assert hist.status_code == 200, hist.text
    assert len(hist.json()) == 1
    assert hist.json()[0]['dispatch_qty'] is None
    pack = client.get('/api/quality/report-pack?from_date=2026-06-01&to_date=2026-06-30', headers=h)
    assert pack.status_code == 200, pack.text
    assert pack.json()['summary'] == j
    assert pack.json()['phenomenon_pareto'] == pto.json()
    assert pack.json()['history'] == hist.json()
    filtered = client.get('/api/quality/report-pack?from_date=2026-06-01&to_date=2026-06-30&shift=A', headers=h)
    assert filtered.status_code == 200, filtered.text
    assert filtered.json()['summary']['history_included'] is False
    assert filtered.json()['history'] == hist.json()


def test_zero_rejection_without_denominator_does_not_block_month_ppm():
    from fastapi.testclient import TestClient
    from app.main import app
    from app.models import QualityRejectionMonthlyHistory

    with SessionLocal() as db:
        p, _, ph = _seed_quality_context(db)
        db.add_all([
            QualityRejectionMonthlyHistory(
                record_key='HMCL|Total|2026-06|HAS REJECTION', month=date(2026,6,1), source='HMCL',
                source_sheet='Total', record_scope='AGGREGATE_TOTAL', include_in_aggregate=True,
                product_id=p.id, phenomenon_id=ph.id, reject_qty=Decimal('25'),
                denominator_source='DISP_DONE', dispatch_qty=Decimal('2500'), ppm=Decimal('10000'),
            ),
            QualityRejectionMonthlyHistory(
                record_key='HMCL|Total|2026-06|ZERO', month=date(2026,6,1), source='HMCL',
                source_sheet='Total', record_scope='AGGREGATE_TOTAL', include_in_aggregate=True,
                product_id=p.id, phenomenon_id=ph.id, reject_qty=Decimal('0'),
                denominator_source='DISP_DONE', dispatch_qty=None, ppm=None,
            ),
        ])
        db.commit()

    client = TestClient(app)
    h = _headers_for_quality(client)
    dashboard = client.get('/api/quality/dashboard?from_date=2026-06-01&to_date=2026-06-30', headers=h).json()
    trend = client.get('/api/quality/monthly-trend?from_month=2026-06-01&to_month=2026-06-01', headers=h).json()[0]
    assert dashboard['ppm_pending_rows'] == 0
    assert dashboard['ppm'] == 10000.0
    assert trend['ppm_pending_rows'] == 0
    assert trend['ppm'] == 10000.0


def test_historical_rejection_auto_uses_historical_mis_dispatch_qty():
    from fastapi.testclient import TestClient
    from app.main import app
    from app.models import QualityRejectionMonthlyHistory

    with SessionLocal() as db:
        p, _, ph = _seed_quality_context(db)
        # Matching historical MIS for June; monthly Actual = 1,000 + 1,500 = 2,500.
        db.add(DailyMIS(mis_date=date(2026,6,10), product_id=p.id, plan_qty=Decimal('1100'), actual_qty=Decimal('1000'),
                        sales_price=Decimal('1'), plan_sales=Decimal('1100'), actual_sales=Decimal('1000'), source=SourceType.EXCEL))
        db.add(DailyMIS(mis_date=date(2026,6,20), product_id=p.id, plan_qty=Decimal('1600'), actual_qty=Decimal('1500'),
                        sales_price=Decimal('1'), plan_sales=Decimal('1600'), actual_sales=Decimal('1500'), source=SourceType.EXCEL))
        db.add(QualityRejectionMonthlyHistory(
            record_key='HMCL|Total|2026-06|MIS AUTO',
            month=date(2026,6,1), source='HMCL', source_sheet='Total', record_scope='AGGREGATE_TOTAL',
            include_in_aggregate=True, product_id=p.id, plant=None, detection_operation='Disp_Done',
            responsible_team='Operation', phenomenon_id=ph.id, reject_qty=Decimal('25'),
            denominator_source='DISP_DONE', dispatch_qty=None, ppm=None,
        ))
        db.commit()

    client = TestClient(app)
    login = client.post('/api/auth/login', json={'username':'admin','password':'ChangeMe123!'})
    assert login.status_code == 200
    h={'Authorization': f"Bearer {login.json()['access_token']}"}

    hist = client.get('/api/quality/history?from_date=2026-06-01&to_date=2026-06-30', headers=h)
    assert hist.status_code == 200, hist.text
    row = hist.json()[0]
    assert row['dispatch_qty'] == 2500.0
    assert row['dispatch_qty_source'] == 'MIS_HISTORY'
    assert row['ppm'] == 10000.0

    dash = client.get('/api/quality/dashboard?from_date=2026-06-01&to_date=2026-06-30', headers=h)
    assert dash.status_code == 200, dash.text
    j = dash.json()
    assert j['denominator_qty'] == 2500.0
    assert j['ppm'] == 10000.0


def test_historical_rejection_import_persists_mis_dispatch_fallback(tmp_path):
    from openpyxl import Workbook
    from app.services.quality_import import import_historical_rejection_workbook
    from app.models import QualityRejectionMonthlyHistory

    path = tmp_path / 'historical_mis_denominator.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Historical_Rejection_Import'
    ws.append([
        'Record_Key','Month','WebApp_Product','Phenomenon_WebApp','Reject_Qty',
        'Source','Source_Sheet','Record_Scope','Include_In_Overall_Aggregate',
        'Detection_Operation','Responsible_Team','Denominator_Source'
    ])
    ws.append(['HMCL|Total|2026-06|AUTO', date(2026,6,1), 'K70 Cylinder block', 'BORE O/S', 20,
               'HMCL', 'Total', 'AGGREGATE_TOTAL', 'Yes', 'Disp_Done', 'Operation', 'DISP_DONE'])
    wb.save(path)

    with SessionLocal() as db:
        p, _, _ = _seed_quality_context(db)
        db.add(DailyMIS(mis_date=date(2026,6,15), product_id=p.id, plan_qty=Decimal('2100'), actual_qty=Decimal('2000'),
                        sales_price=Decimal('1'), plan_sales=Decimal('2100'), actual_sales=Decimal('2000'), source=SourceType.EXCEL))
        db.commit()
        stats = import_historical_rejection_workbook(db, path)
        db.commit()
        assert stats['dispatch_from_mis'] == 1
        row = db.scalar(select(QualityRejectionMonthlyHistory).where(QualityRejectionMonthlyHistory.record_key=='HMCL|Total|2026-06|AUTO'))
        assert row.dispatch_qty == Decimal('2000')
        assert row.ppm == Decimal('10000.000')


def test_historical_ppm_plant_filter_uses_plant_scoped_dispatch_qty():
    from fastapi.testclient import TestClient
    from app.main import app
    from app.models import QualityRejectionMonthlyHistory

    with SessionLocal() as db:
        p, _, ph = _seed_quality_context(db)
        db.add_all([
            DailyMIS(mis_date=date(2026,6,10), product_id=p.id, plan_qty=Decimal('5000'), actual_qty=Decimal('5000'),
                     sales_price=Decimal('1'), plan_sales=Decimal('5000'), actual_sales=Decimal('5000'), source=SourceType.EXCEL),
            QualityRejectionMonthlyHistory(
                record_key='PLANT2020|2026-06|BORE O/S', month=date(2026,6,1), source='TEST',
                source_sheet='Plant 2020', record_scope='PRODUCT_TOTAL', include_in_aggregate=True,
                product_id=p.id, plant='2020', phenomenon_id=ph.id, reject_qty=Decimal('20'),
                denominator_source='DISP_DONE', dispatch_qty=Decimal('2000'), ppm=Decimal('10000'),
            ),
            QualityRejectionMonthlyHistory(
                record_key='PLANT2050|2026-06|BORE O/S', month=date(2026,6,1), source='TEST',
                source_sheet='Plant 2050', record_scope='PRODUCT_TOTAL', include_in_aggregate=True,
                product_id=p.id, plant='2050', phenomenon_id=ph.id, reject_qty=Decimal('30'),
                denominator_source='DISP_DONE', dispatch_qty=Decimal('3000'), ppm=Decimal('10000'),
            ),
        ])
        db.commit()

    client = TestClient(app)
    h = _headers_for_quality(client)
    p2020 = client.get('/api/quality/report-pack?from_date=2026-06-01&to_date=2026-06-30&plant=2020', headers=h)
    p2050 = client.get('/api/quality/report-pack?from_date=2026-06-01&to_date=2026-06-30&plant=2050', headers=h)
    assert p2020.status_code == 200, p2020.text
    assert p2050.status_code == 200, p2050.text
    assert p2020.json()['summary']['reject_qty'] == 20.0
    assert p2020.json()['summary']['denominator_qty'] == 2000.0
    assert p2020.json()['summary']['ppm'] == 10000.0
    assert p2050.json()['summary']['reject_qty'] == 30.0
    assert p2050.json()['summary']['denominator_qty'] == 3000.0
    assert p2050.json()['summary']['ppm'] == 10000.0


def test_historical_ppm_plant_filter_does_not_reuse_product_total_mis_for_split_product():
    from fastapi.testclient import TestClient
    from app.main import app
    from app.models import QualityRejectionMonthlyHistory

    with SessionLocal() as db:
        p, _, ph = _seed_quality_context(db)
        db.add(DailyMIS(mis_date=date(2026,6,10), product_id=p.id, plan_qty=Decimal('5000'), actual_qty=Decimal('5000'),
                        sales_price=Decimal('1'), plan_sales=Decimal('5000'), actual_sales=Decimal('5000'), source=SourceType.EXCEL))
        # Simulates rows imported by an older version that persisted the same
        # product-total MIS fallback into both plant rows.
        db.add_all([
            QualityRejectionMonthlyHistory(
                record_key='OLD2020|2026-06|BORE O/S', month=date(2026,6,1), source='TEST',
                source_sheet='Plant 2020', record_scope='PRODUCT_TOTAL', include_in_aggregate=True,
                product_id=p.id, plant='2020', phenomenon_id=ph.id, reject_qty=Decimal('20'),
                denominator_source='DISP_DONE', dispatch_qty=Decimal('5000'), ppm=Decimal('4000'),
            ),
            QualityRejectionMonthlyHistory(
                record_key='OLD2050|2026-06|BORE O/S', month=date(2026,6,1), source='TEST',
                source_sheet='Plant 2050', record_scope='PRODUCT_TOTAL', include_in_aggregate=True,
                product_id=p.id, plant='2050', phenomenon_id=ph.id, reject_qty=Decimal('30'),
                denominator_source='DISP_DONE', dispatch_qty=Decimal('5000'), ppm=Decimal('6000'),
            ),
        ])
        db.commit()

    client = TestClient(app)
    h = _headers_for_quality(client)
    response = client.get('/api/quality/report-pack?from_date=2026-06-01&to_date=2026-06-30&plant=2020', headers=h)
    assert response.status_code == 200, response.text
    summary = response.json()['summary']
    assert summary['reject_qty'] == 20.0
    assert summary['denominator_qty'] == 0.0
    assert summary['ppm'] is None
    assert summary['ppm_pending_rows'] == 1


def test_historical_ppm_never_falls_back_to_company_dispatch_and_rejection_stays_available():
    with SessionLocal() as db:
        p, _, ph = _seed_quality_context(db)
        row = QualityRejectionMonthlyHistory(
            record_key='P0-NONBLOCKING', month=date(2026,9,1), source='TEST',
            source_sheet='Quality', record_scope='AGGREGATE_TOTAL', include_in_aggregate=True,
            product_id=p.id, phenomenon_id=ph.id, reject_qty=Decimal('25'),
            denominator_source='DISP_DONE', dispatch_qty=Decimal('2000'), ppm=Decimal('12500'),
        )
        db.add(row); db.commit(); db.refresh(row)
        resolved = _history_dispatch_resolution(db, [row])
        assert row.reject_qty == Decimal('25')
        assert resolved[row.id] == (None, 'SIDDHARTH_SILVER_PENDING')


def test_historical_ppm_uses_only_siddharth_plus_silver():
    with SessionLocal() as db:
        p, _, ph = _seed_quality_context(db)
        prod = QualityHistoricalPpmProduction(
            month=date(2026,9,1), product_id=p.id,
            siddharth_machining_qty=Decimal('1500'), silver_production_qty=Decimal('500'),
        )
        row = QualityRejectionMonthlyHistory(
            record_key='P0-AUTHORITATIVE', month=date(2026,9,1), source='TEST',
            source_sheet='Quality', record_scope='AGGREGATE_TOTAL', include_in_aggregate=True,
            product_id=p.id, phenomenon_id=ph.id, reject_qty=Decimal('25'),
            denominator_source='DISP_DONE', dispatch_qty=Decimal('9999'), ppm=None,
        )
        db.add_all([prod,row]); db.commit(); db.refresh(row)
        resolved = _history_dispatch_resolution(db, [row])
        assert resolved[row.id] == (2000.0, 'SIDDHARTH_MACHINING_PLUS_SILVER')


def test_historical_ppm_production_requires_both_sources_but_allows_explicit_zero(tmp_path):
    path = tmp_path / 'ppm_production.xlsx'
    wb = Workbook(); ws = wb.active; ws.title = 'Historical_Production_Upload'
    ws.append(['Month','Product','Siddharth_Machining_Production','Silver_Production','Remark'])
    ws.append([date(2026,9,1),'K70 Cylinder block',1500,None,'missing Silver'])
    ws.append([date(2026,10,1),'K70 Cylinder block',1500,0,'explicit zero is valid'])
    wb.save(path)
    with SessionLocal() as db:
        _seed_quality_context(db)
        stats = import_historical_ppm_production_workbook(db, path)
        assert stats['created'] == 1
        assert any('enter both Siddharth Machining Production and Silver Production' in e for e in stats['errors'])
