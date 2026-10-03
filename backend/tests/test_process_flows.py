from datetime import date
from pathlib import Path
from openpyxl import Workbook
from fastapi.testclient import TestClient
from sqlalchemy import select, func
from app.main import app
from app.db import SessionLocal
from app.models import Product,ProcessFlowVersion,ProcessFlowStage,StageScheduleAllocation,DailyRequirement,ProcessDailySummary,DailyMIS,GovernanceAudit,MonthStatus,User,WorkingCalendar
from app.enums import MonthState,UserRole
from app.services.process_flows import import_process_workbook,HEADERS,SHEETS
from test_governance_security import headers
TEMPLATE=Path(__file__).resolve().parents[1]/'templates'/'Production_Process_Upload_v0.4.0.xlsx'
def file(tmp,kind,rows):
 w=Workbook();s=w.active;s.title=SHEETS[kind];s.append(HEADERS[kind])
 for r in rows:s.append([r.get(k) for k in HEADERS[kind]])
 p=tmp/(kind+'.xlsx');w.save(p);return p

def setup():
 with SessionLocal() as db:
  st=import_process_workbook(db,TEMPLATE,'process-design');assert not st['errors'],st['errors'];db.commit()
  return {p.name:p.id for p in db.scalars(select(Product))}

def schedule(**kw):return {'product':'Kubota Head','stage_code':'P04_VMC','month':date(2026,10,1),'effective_from':date(2026,10,1),'allocated_qty':103,'reference':'A1','reason':'Vendor allocation approved','correct_imported_plans':False,**kw}
def actual(code='P04_VMC',qty=10,**kw):return {'product':'Kubota Head','stage_code':code,'date':date(2026,10,2),'actual_qty':qty,'reject_qty':0,'reason':'',**kw}
def commit(db,p,kind):
 st=import_process_workbook(db,p,kind);assert not st['errors'],st;db.commit();return st

def preview(c,h,p,kind):return c.post('/api/import/preview/'+kind,headers=h,files={'file':(p.name,p.read_bytes())})

def test_all_corrected_definitions_preserve_inactive_columns_and_no_fake_plans():
 ids=setup()
 with SessionLocal() as db:
  assert db.scalar(select(func.count()).select_from(ProcessFlowVersion))==23
  assert db.scalar(select(func.count()).select_from(ProcessFlowStage).where(ProcessFlowStage.is_active.is_(True)))==151
  assert not db.scalar(select(StageScheduleAllocation))
  k=db.scalars(select(ProcessFlowStage).join(ProcessFlowVersion).where(ProcessFlowVersion.product_id==ids['Kubota Head'],ProcessFlowStage.is_active.is_(True)).order_by(ProcessFlowStage.sequence_no)).all()
  assert [s.source_column for s in k]==['AN','AO','AP','AQ','AR','AS'];assert k[2].vendor_name==k[3].vendor_name=='Shradha Industries'
  for prod,col in [('Aluminium Head','FT'),('Grab Handle Machined','FJ')]:
   assert db.scalar(select(ProcessFlowStage).join(ProcessFlowVersion).where(ProcessFlowVersion.product_id==ids[prod],ProcessFlowStage.parent_dispatch.is_(True))).source_column==col
  hm=db.scalars(select(ProcessFlowStage).join(ProcessFlowVersion).where(ProcessFlowVersion.product_id==ids['HMCL Cylinder block'],ProcessFlowStage.is_active.is_(True))).all()
  assert {s.variant for s in hm if s.role=='DISPATCH_DETAIL'}=={'ACK Black','AAEE Silver','BS4 KCC PC','BS4 KCC SL','BS4 KCC RT'}
  assert not db.scalar(select(ProcessFlowStage).where(ProcessFlowStage.source_column=='HP')).is_active
  assert not db.get(Product,ids['Platina 99 Cylinder block']).is_active
  assert import_process_workbook(db,TEMPLATE,'process-design')['unchanged']==235

def test_working_day_allocation_exact_total_effective_revision_and_zero(tmp_path):
 setup()
 with SessionLocal() as db:
  db.add(WorkingCalendar(work_date=date(2026,10,2),plant='2020',is_working_day=False));db.commit()
  r=commit(db,file(tmp_path,'stage-schedules',[schedule()]),'stage-schedules');days=r['allocations'][0]['daily']
  assert sum(x['qty'] for x in days)==103;assert '2026-10-02' not in {x['date'] for x in days}
  stage=db.scalar(select(ProcessFlowStage).where(ProcessFlowStage.code=='P04_VMC'))
  before={x.req_date:x.revised_plan_qty for x in db.scalars(select(DailyRequirement).where(DailyRequirement.route_operation_id==stage.route_operation_id,DailyRequirement.req_date<date(2026,10,15)))}
  commit(db,file(tmp_path,'stage-daily',[actual(qty=90)]),'stage-daily')
  commit(db,file(tmp_path,'stage-schedules',[schedule(effective_from=date(2026,10,15),allocated_qty=140,reference='A2')]),'stage-schedules')
  reqs=db.scalars(select(DailyRequirement).where(DailyRequirement.route_operation_id==stage.route_operation_id)).all()
  assert sum(x.revised_plan_qty for x in reqs)==140
  assert {x.req_date:x.revised_plan_qty for x in reqs if x.req_date<date(2026,10,15)}==before
  assert len(db.scalars(select(StageScheduleAllocation)).all())==2
  commit(db,file(tmp_path,'stage-schedules',[schedule(stage_code='P04_CASTING',allocated_qty=0)]),'stage-schedules')
  assert import_process_workbook(db,file(tmp_path,'stage-schedules',[schedule(stage_code='P04_POURING',allocated_qty=None)]),'stage-schedules')['errors'];db.rollback()

def test_parent_mis_corrections_preserve_plan_price_and_audit(tmp_path):
 ids=setup();c=TestClient(app);h=headers(c)
 with SessionLocal() as db:db.add(DailyMIS(product_id=ids['Kubota Head'],mis_date=date(2026,10,2),plan_qty=100,actual_qty=4,sales_price=12,actual_sales=48,plan_sales=1200));db.commit()
 p=file(tmp_path,'stage-daily',[actual(code='P04_DISP_DONE',qty=8)])
 assert not preview(c,h,p,'stage-daily').json()['can_confirm']
 p=file(tmp_path,'stage-daily',[actual(code='P04_DISP_DONE',qty=8,reason='Correct dispatch record')])
 r=preview(c,h,p,'stage-daily');assert r.json()['can_confirm'],r.text
 with SessionLocal() as db:assert not db.scalar(select(ProcessDailySummary))
 r=c.post('/api/import/confirm',headers=h,json={'preview_token':r.json()['preview_token']});assert r.status_code==200,r.text
 with SessionLocal() as db:
  m=db.scalar(select(DailyMIS));assert (m.actual_qty,m.plan_qty,m.sales_price,m.plan_sales,m.actual_sales)==(8,100,12,1200,96)
  assert db.scalar(select(GovernanceAudit).where(GovernanceAudit.entity=='daily_mis')).reason=='Correct dispatch record'

def test_hmcl_variants_reconcile_without_double_counting(tmp_path):
 setup()
 with SessionLocal() as db:
  details=[('P20_SILVER_DISPATCH',1),('P20_ACK_DISPATCH',2),('P20_DISPATCH_KCC_P_C',3),('P20_DISPATCH_KCC_SL',4),('P20_DISPATCH_KCC_RT',5)]
  commit(db,file(tmp_path,'stage-daily',[actual(code=c,qty=q,product='HMCL Cylinder block') for c,q in details]),'stage-daily');assert not db.scalar(select(DailyMIS))
  p=file(tmp_path,'stage-daily',[actual(code='P20_DISPATCH_TOTAL',qty=14,product='HMCL Cylinder block')]);assert 'differs' in import_process_workbook(db,p,'stage-daily')['errors'][0];db.rollback()
  commit(db,file(tmp_path,'stage-daily',[actual(code='P20_DISPATCH_TOTAL',qty=15,product='HMCL Cylinder block')]),'stage-daily');assert db.scalar(select(DailyMIS)).actual_qty==15
  p=file(tmp_path,'stage-daily',[actual(code='P20_SILVER_PRODUCTION',qty=8,product='HMCL Cylinder block')]);assert 'deferred' in import_process_workbook(db,p,'stage-daily')['errors'][0];db.rollback()

def test_row_errors_are_atomic_and_production_only_imports_daily(tmp_path):
 setup();c=TestClient(app);h=headers(c)
 p=file(tmp_path,'stage-daily',[actual(),actual(code='UNKNOWN',qty=3)])
 r=preview(c,h,p,'stage-daily');assert not r.json()['can_confirm'] and 'Row 3' in r.json()['errors'][0],r.text
 with SessionLocal() as db:
  assert not db.scalar(select(ProcessDailySummary));db.scalar(select(User).where(User.username=='admin')).role=UserRole.PRODUCTION;db.commit()
 p=file(tmp_path,'stage-daily',[actual()]);assert preview(c,h,p,'stage-daily').json()['can_confirm']
 assert preview(c,h,TEMPLATE,'process-design').status_code==403
 assert preview(c,h,p,'stage-schedules').status_code==403

def test_month_close_and_legacy_propagation_protect_stage_plans(tmp_path):
 ids=setup();c=TestClient(app);h=headers(c)
 with SessionLocal() as db:
  commit(db,file(tmp_path,'stage-schedules',[schedule()]),'stage-schedules')
  stage=db.scalar(select(ProcessFlowStage).where(ProcessFlowStage.code=='P04_VMC'))
  plans=[(x.req_date,x.revised_plan_qty) for x in db.scalars(select(DailyRequirement).where(DailyRequirement.route_operation_id==stage.route_operation_id))]
  db.add(DailyRequirement(req_date=date(2026,10,2),product_id=ids['Kubota Head'],route_operation_id=None,baseline_plan_qty=1,revised_plan_qty=999));db.commit()
  from app.services.planning import rebuild_process_requirements
  rebuild_process_requirements(db,ids['Kubota Head'],date(2026,10,1));db.commit()
  assert plans==[(x.req_date,x.revised_plan_qty) for x in db.scalars(select(DailyRequirement).where(DailyRequirement.route_operation_id==stage.route_operation_id))]
  db.add(MonthStatus(month=date(2026,10,1),status=MonthState.CLOSED));db.commit()
 p=file(tmp_path,'stage-daily',[actual()]);r=preview(c,h,p,'stage-daily');assert not r.json()['can_confirm'] and 'closed' in r.json()['errors'][0]
 r=preview(c,h,TEMPLATE,'excel');assert 'Explicit process flows' in r.json()['errors'][0]

def test_one_row_reason_cannot_authorize_a_different_row_correction(tmp_path):
 setup()
 with SessionLocal() as db:
  commit(db,file(tmp_path,'stage-daily',[actual(),actual(code='P04_CASTING')]),'stage-daily')
  p=file(tmp_path,'stage-daily',[actual(qty=11,reason='VMC correction'),actual(code='P04_CASTING',qty=11)])
  result=import_process_workbook(db,p,'stage-daily')
  assert result['errors'] and 'Row 3' in result['errors'][0];db.rollback()

def test_new_flow_preserves_legacy_actuals_and_revision_snapshots(tmp_path):
 from openpyxl import load_workbook
 ids=setup()
 with SessionLocal() as db:
  commit(db,file(tmp_path,'stage-daily',[actual()]),'stage-daily')
  old=db.scalar(select(ProcessDailySummary));old_id=old.id;old_operation=old.route_operation_id
  w=load_workbook(TEMPLATE);s=w['Flow_Definition']
  for row in s.iter_rows(min_row=2):row[12].value=date(2026,11,1);row[13].value='November flow revision'
  p=tmp_path/'new-flow.xlsx';w.save(p)
  commit(db,p,'process-design')
  assert db.get(ProcessDailySummary,old_id).route_operation_id==old_operation
  assert db.get(ProcessDailySummary,old_id).actual_qty==10
  assert db.scalar(select(func.count()).select_from(ProcessFlowVersion))==46
  from app.services.process_flows import stage_for
  assert stage_for(db,db.get(Product,ids['Kubota Head']),'P04_VMC',date(2026,10,2)).route_operation_id==old_operation
  assert stage_for(db,db.get(Product,ids['Kubota Head']),'P04_VMC',date(2026,11,2)).route_operation_id!=old_operation

def test_same_workbook_supports_three_separate_confirmed_import_kinds(tmp_path):
 from openpyxl import load_workbook
 w=load_workbook(TEMPLATE)
 for name,row in [('Stage_Schedules',schedule()),('Daily_Actuals',actual())]:
  s=w[name];s.delete_rows(2,s.max_row);s.append([row.get(k) for k in HEADERS['stage-schedules' if name=='Stage_Schedules' else 'stage-daily']])
 p=tmp_path/'combined.xlsx';w.save(p);c=TestClient(app);h=headers(c)
 for kind in ['process-design','stage-schedules','stage-daily']:
  r=preview(c,h,p,kind);assert r.json().get('can_confirm'),r.text
  r=c.post('/api/import/confirm',headers=h,json={'preview_token':r.json()['preview_token']});assert r.status_code==200,r.text
  assert preview(c,h,p,kind).json()['status']=='already_imported'

def test_mid_month_flow_revision_counts_plans_from_prior_stage_version(tmp_path):
 from openpyxl import load_workbook
 ids=setup()
 with SessionLocal() as db:
  commit(db,file(tmp_path,'stage-schedules',[schedule()]),'stage-schedules')
  w=load_workbook(TEMPLATE);s=w['Flow_Definition']
  for row in s.iter_rows(min_row=2):row[12].value=date(2026,10,15);row[13].value='Mid month revision'
  p=tmp_path/'mid-month.xlsx';w.save(p);commit(db,p,'process-design')
  result=commit(db,file(tmp_path,'stage-schedules',[schedule(effective_from=date(2026,10,15),allocated_qty=140,reference='A2')]),'stage-schedules')
  assert result['allocations'][0]['plan_before_effective']>0
  assert result['allocations'][0]['remaining_qty']<140
  assert result['allocations'][0]['remaining_qty']+result['allocations'][0]['plan_before_effective']==140

def test_manual_stage_entry_and_missing_data_dashboard(tmp_path):
 ids=setup();c=TestClient(app);h=headers(c)
 response=c.get(f'/api/flows/monitor?product_id={ids["Kubota Head"]}&monitor_date=2026-10-02',headers=h)
 assert response.status_code==200,response.text
 stage=next(s for s in response.json()['stages'] if s['code']=='P04_VMC')
 assert stage['plan'] is None and stage['actual'] is None
 r=c.post('/api/process/entry',headers=h,json={'summary_date':'2026-10-02','product_id':ids['Kubota Head'],'route_operation_id':stage['route_operation_id'],'actual_qty':8,'good_qty':7,'reject_qty':1,'plan_qty':999,'opening_wip':0,'closing_wip':0,'source':'MANUAL'})
 assert r.status_code==200,r.text
 with SessionLocal() as db:
  row=db.scalar(select(ProcessDailySummary));assert row.source.value=='MANUAL' and row.plan_qty==0
 r=c.get('/api/insights/data-quality?as_of=2026-10-02',headers=h)
 assert r.status_code==200,r.text
 assert r.json()['counts']['stage_schedule_missing']==151
 assert r.json()['counts']['vendor_missing']==6
 assert c.get('/api/flows/template',headers=h).content==TEMPLATE.read_bytes()

def test_closed_month_authorization_survives_preview_and_consumes_at_confirm(tmp_path):
 from app.models import HistoricalCorrectionGrant
 ids=setup();c=TestClient(app);h=headers(c)
 assert c.post('/api/governance/months/close',headers=h,json={'month':'2026-10-01','reason':'October closed for review'}).status_code==200
 with SessionLocal() as db:uid=db.scalar(select(User.id).where(User.username=='admin'))
 r=c.post('/api/governance/corrections',headers=h,json={'month':'2026-10-01','user_id':uid,'reason':'Authorized stage actual correction'})
 gid=r.json()['correction_id'];authorized={**h,'X-Correction-ID':str(gid)}
 p=file(tmp_path,'stage-daily',[actual()]);r=preview(c,authorized,p,'stage-daily');assert r.json()['can_confirm'],r.text
 with SessionLocal() as db:assert db.get(HistoricalCorrectionGrant,gid).used_at is None
 r=c.post('/api/import/confirm',headers=authorized,json={'preview_token':r.json()['preview_token']});assert r.status_code==200,r.text
 with SessionLocal() as db:
  assert db.get(HistoricalCorrectionGrant,gid).used_at is not None
  assert db.scalar(select(ProcessDailySummary)).actual_qty==10


def test_process_monitor_range_returns_daily_plan_actual_working_days_only():
 ids=setup();c=TestClient(app);h=headers(c)
 with SessionLocal() as db:
  product_id=ids['Kubota Head']
  stage=db.scalar(select(ProcessFlowStage).join(ProcessFlowVersion).where(
   ProcessFlowVersion.product_id==product_id,
   ProcessFlowStage.code=='P04_VMC',
   ProcessFlowStage.is_active.is_(True),
  ).order_by(ProcessFlowVersion.effective_from.desc()))
  assert stage and stage.route_operation_id
  db.add_all([
   WorkingCalendar(work_date=date(2026,10,5),plant='2020',is_working_day=False,reason='Planned OFF'),
   WorkingCalendar(work_date=date(2026,10,6),plant='2020',is_working_day=True),
   WorkingCalendar(work_date=date(2026,10,7),plant='2020',is_working_day=True),
   DailyRequirement(req_date=date(2026,10,5),product_id=product_id,route_operation_id=stage.route_operation_id,baseline_plan_qty=999,revised_plan_qty=999),
   DailyRequirement(req_date=date(2026,10,6),product_id=product_id,route_operation_id=stage.route_operation_id,baseline_plan_qty=100,revised_plan_qty=100),
   DailyRequirement(req_date=date(2026,10,7),product_id=product_id,route_operation_id=stage.route_operation_id,baseline_plan_qty=120,revised_plan_qty=120),
   ProcessDailySummary(summary_date=date(2026,10,5),product_id=product_id,route_operation_id=stage.route_operation_id,plan_qty=999,actual_qty=999,good_qty=999,reject_qty=0),
   ProcessDailySummary(summary_date=date(2026,10,6),product_id=product_id,route_operation_id=stage.route_operation_id,plan_qty=100,actual_qty=90,good_qty=90,reject_qty=0),
  ])
  db.commit()

 response=c.get(f'/api/flows/monitor?product_id={ids["Kubota Head"]}&start_date=2026-10-05&end_date=2026-10-07',headers=h)
 assert response.status_code==200,response.text
 body=response.json()
 assert body['period']['working_days']==2
 assert body['period']['off_days']==1
 stage=next(s for s in body['stages'] if s['code']=='P04_VMC')
 assert stage['plan']==110
 assert stage['actual']==90
 assert stage['days_with_plan']==2
 assert stage['days_with_data']==1
 assert stage['daily']==[
  {'date':'2026-10-06','plan':100.0,'actual':90.0,'reject':0.0},
  {'date':'2026-10-07','plan':120.0,'actual':None,'reject':None},
 ]
