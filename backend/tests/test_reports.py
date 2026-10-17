from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from fastapi.testclient import TestClient

from app.db import SessionLocal
from app.main import app
from app.models import (Action, ActionContext, ActionWhyWhy, Customer, DailyMIS, DailyRequirement, Operation,
                        ProcessFlowStage, ProcessFlowVersion, Product, ReviewActionLink,
                        ReviewPoint, ReviewSession, RouteOperation, RouteVersion)


def auth_headers(client: TestClient):
    r = client.post('/api/auth/login', json={'username':'admin','password':'ChangeMe123!'})
    assert r.status_code == 200
    return {'Authorization': f"Bearer {r.json()['access_token']}"}


def test_compliance_reports_use_revised_plan():
    with SessionLocal() as db:
        c=Customer(code='C1',name='Customer 1'); db.add(c); db.flush()
        p=Product(code='P1',name='Part 1',customer_id=c.id); db.add(p); db.flush()
        db.add_all([
            DailyMIS(mis_date=date(2026,9,1),product_id=p.id,plan_qty=Decimal('100'),actual_qty=Decimal('90'),sales_price=Decimal('10'),plan_sales=Decimal('1000'),actual_sales=Decimal('900')),
            DailyMIS(mis_date=date(2026,9,2),product_id=p.id,plan_qty=Decimal('100'),actual_qty=Decimal('120'),sales_price=Decimal('10'),plan_sales=Decimal('1000'),actual_sales=Decimal('1200')),
        ])
        # Revised requirement from the schedule engine must drive report plan.
        db.add(DailyRequirement(req_date=date(2026,9,2),product_id=p.id,route_operation_id=None,baseline_plan_qty=Decimal('100'),revised_plan_qty=Decimal('150')))
        db.commit(); pid=p.id
    client=TestClient(app); h=auth_headers(client)
    r=client.get(f'/api/reports/compliance?as_of=2026-09-02&product_id={pid}&metric=qty',headers=h)
    assert r.status_code == 200, r.text
    j=r.json()
    assert len(j['daily']) == 2
    assert j['summary']['plan'] == 250.0
    assert j['summary']['actual'] == 210.0
    assert round(j['summary']['compliance'],3) == .84
    rs=client.get(f'/api/reports/compliance?as_of=2026-09-02&product_id={pid}&metric=sales',headers=h)
    assert rs.status_code == 200
    assert rs.json()['summary']['plan'] == 2500.0


def test_compliance_filters_and_tonnage():
    with SessionLocal() as db:
        c=Customer(code='C2',name='Customer 2'); db.add(c); db.flush()
        p1=Product(code='P2A',name='Part A',customer_id=c.id,plant='Plant A',product_group='Aluminium',finish_weight_kg=Decimal('2.0')); db.add(p1)
        p2=Product(code='P2B',name='Part B',customer_id=c.id,plant='Plant B',product_group='Cast Iron',finish_weight_kg=Decimal('5.0')); db.add(p2); db.flush()
        db.add_all([
            DailyMIS(mis_date=date(2026,9,1),product_id=p1.id,plan_qty=Decimal('100'),actual_qty=Decimal('90'),sales_price=Decimal('10'),plan_sales=Decimal('1000'),actual_sales=Decimal('900')),
            DailyMIS(mis_date=date(2026,9,1),product_id=p2.id,plan_qty=Decimal('100'),actual_qty=Decimal('100'),sales_price=Decimal('10'),plan_sales=Decimal('1000'),actual_sales=Decimal('1000')),
        ])
        db.commit()
    client=TestClient(app); h=auth_headers(client)
    r=client.get('/api/reports/compliance?as_of=2026-09-01&plant=Plant%20A&metric=tonnage',headers=h)
    assert r.status_code == 200, r.text
    j=r.json()
    assert j['summary']['plan'] == 0.2
    assert j['summary']['actual'] == 0.18
    assert len(j['parts']) == 1
    assert j['parts'][0]['plant'] == 'Plant A'
    assert j['parts'][0]['product_group'] == 'Aluminium'


def test_compliance_weekly_window_is_last_8_weeks():
    with SessionLocal() as db:
        c = Customer(code='C8W', name='Customer 8W'); db.add(c); db.flush()
        p = Product(code='P8W', name='Part 8W', customer_id=c.id); db.add(p); db.flush()
        # Twelve Monday records ending in the selected current week.
        start = date(2026, 7, 13)
        for i in range(12):
            d = start + __import__('datetime').timedelta(weeks=i)
            db.add(DailyMIS(
                mis_date=d, product_id=p.id,
                plan_qty=Decimal('100'), actual_qty=Decimal(str(80+i)),
                sales_price=Decimal('10'), plan_sales=Decimal('1000'), actual_sales=Decimal(str((80+i)*10)),
            ))
        db.commit(); pid = p.id
    client = TestClient(app); h = auth_headers(client)
    r = client.get(f'/api/reports/compliance?as_of=2026-09-28&product_id={pid}&metric=qty', headers=h)
    assert r.status_code == 200, r.text
    weekly = r.json()['weekly']
    assert len(weekly) == 8
    assert weekly[0]['start'] == '2026-08-10'
    assert weekly[-1]['start'] == '2026-09-28'
    assert all(x['label'].startswith('W') and '•' not in x['label'] for x in weekly)


def test_compliance_accepts_multi_select_filters():
    with SessionLocal() as db:
        c=Customer(code='CMS',name='Multi Customer'); db.add(c); db.flush()
        p1=Product(code='MS1',name='Multi A',customer_id=c.id,plant='Plant A',product_group='AL'); db.add(p1)
        p2=Product(code='MS2',name='Multi B',customer_id=c.id,plant='Plant B',product_group='CI'); db.add(p2)
        p3=Product(code='MS3',name='Multi C',customer_id=c.id,plant='Plant C',product_group='CI'); db.add(p3); db.flush()
        for p,actual in [(p1,90),(p2,80),(p3,70)]:
            db.add(DailyMIS(mis_date=date(2026,9,1),product_id=p.id,plan_qty=Decimal('100'),actual_qty=Decimal(str(actual)),sales_price=Decimal('10'),plan_sales=Decimal('1000'),actual_sales=Decimal(str(actual*10))))
        db.commit(); pids=f'{p1.id},{p2.id}'
    client=TestClient(app); h=auth_headers(client)
    r=client.get(f'/api/reports/compliance?as_of=2026-09-01&product_id={pids}&plant=Plant%20A,Plant%20B&metric=qty',headers=h)
    assert r.status_code == 200, r.text
    j=r.json()
    assert j['summary']['plan'] == 200.0
    assert j['summary']['actual'] == 170.0
    assert {x['product'] for x in j['parts']} == {'Multi A','Multi B'}


def test_compliance_daily_action_counts_are_scoped_and_distinct():
    with SessionLocal() as db:
        c=Customer(code='CACT',name='Action Customer'); db.add(c); db.flush()
        p1=Product(code='ACT1',name='Action Part 1',customer_id=c.id,plant='P1'); db.add(p1)
        p2=Product(code='ACT2',name='Action Part 2',customer_id=c.id,plant='P2'); db.add(p2); db.flush()
        for p in (p1,p2):
            db.add(DailyMIS(mis_date=date(2026,9,3),product_id=p.id,plan_qty=Decimal('100'),actual_qty=Decimal('90'),sales_price=Decimal('10'),plan_sales=Decimal('1000'),actual_sales=Decimal('900')))
        a1=Action(action_no='ACT-2026-T001',reference_date=date(2026,9,3),problem_description='Gap 1',action_description='Follow up 1')
        a2=Action(action_no='ACT-2026-T002',reference_date=date(2026,9,3),problem_description='Gap 2',action_description='Follow up 2')
        a3=Action(action_no='ACT-2026-T003',reference_date=date(2026,9,3),problem_description='Other product',action_description='Follow up 3')
        db.add_all([a1,a2,a3]); db.flush()
        # Two contexts for a1 must still count as one distinct action.
        db.add_all([
            ActionContext(action_id=a1.id,context_date=date(2026,9,3),product_id=p1.id),
            ActionContext(action_id=a1.id,context_date=date(2026,9,3),product_id=p1.id),
            ActionContext(action_id=a2.id,context_date=date(2026,9,3),product_id=p1.id),
            ActionContext(action_id=a3.id,context_date=date(2026,9,3),product_id=p2.id),
        ])
        db.commit(); pid=p1.id
    client=TestClient(app); h=auth_headers(client)
    r=client.get(f'/api/reports/compliance?as_of=2026-09-03&product_id={pid}&metric=qty',headers=h)
    assert r.status_code == 200, r.text
    j=r.json()
    assert j['summary']['actions_raised'] == 2
    assert len(j['daily']) == 1
    assert j['daily'][0]['action_count'] == 2



def test_dispatch_schedule_actual_excel_contains_all_parts_and_mom_actions():
    day1 = date(2026, 10, 1)
    day2 = date(2026, 10, 2)
    day3 = date(2026, 10, 3)
    with SessionLocal() as db:
        customer = Customer(code="C_DSA", name="Dispatch Export Customer")
        db.add(customer)
        db.flush()
        p1 = Product(code="DSA-1", name="Dispatch Export Part 1", customer_id=customer.id, plant="Plant A")
        p2 = Product(code="DSA-2", name="Dispatch Export Part 2", customer_id=customer.id, plant="Plant B")
        db.add_all([p1, p2])
        db.flush()

        # Explicit parent-dispatch flow plan is authoritative over the
        # legacy non-stage requirement for the same product/date.
        dispatch_operation = Operation(code="DSA-DISPATCH", name="Dispatch")
        db.add(dispatch_operation)
        db.flush()
        route = RouteVersion(product_id=p1.id, revision_no=1, effective_from=day2)
        db.add(route)
        db.flush()
        route_operation = RouteOperation(
            route_version_id=route.id,
            operation_id=dispatch_operation.id,
            sequence_no=10,
            is_dispatch=True,
        )
        db.add(route_operation)
        db.flush()
        flow = ProcessFlowVersion(
            product_id=p1.id,
            route_version_id=route.id,
            effective_from=day2,
            revision_no=1,
            definition_sha256="a" * 64,
            source_document="dispatch-export-test.xlsx",
            reason="Test explicit parent dispatch plan",
        )
        db.add(flow)
        db.flush()
        parent_stage = ProcessFlowStage(
            flow_id=flow.id,
            code="DISPATCH",
            name="Dispatch",
            route_operation_id=route_operation.id,
            source_column="A",
            role="INTERNAL",
            branch="",
            variant="",
            vendor_name="",
            predecessors_json="[]",
            alias_of="",
            is_active=True,
            parent_dispatch=True,
            sequence_no=10,
        )
        db.add(parent_stage)
        db.flush()

        # Revised daily requirements must override the imported MIS plan.
        db.add_all([
            DailyRequirement(req_date=day1, product_id=p1.id, route_operation_id=None,
                             baseline_plan_qty=Decimal("90"), revised_plan_qty=Decimal("100")),
            DailyRequirement(req_date=day2, product_id=p1.id, route_operation_id=None,
                             baseline_plan_qty=Decimal("100"), revised_plan_qty=Decimal("120")),
            DailyRequirement(req_date=day2, product_id=p1.id, route_operation_id=route_operation.id,
                             baseline_plan_qty=Decimal("100"), revised_plan_qty=Decimal("125")),
            DailyRequirement(req_date=day1, product_id=p2.id, route_operation_id=None,
                             baseline_plan_qty=Decimal("50"), revised_plan_qty=Decimal("50")),
            DailyMIS(mis_date=day1, product_id=p1.id, plan_qty=Decimal("90"), actual_qty=Decimal("90"),
                     sales_price=Decimal("0"), plan_sales=Decimal("0"), actual_sales=Decimal("0")),
            DailyMIS(mis_date=day2, product_id=p1.id, plan_qty=Decimal("100"), actual_qty=Decimal("100"),
                     sales_price=Decimal("0"), plan_sales=Decimal("0"), actual_sales=Decimal("0")),
            DailyMIS(mis_date=day3, product_id=p1.id, plan_qty=Decimal("75"), actual_qty=Decimal("60"),
                     sales_price=Decimal("0"), plan_sales=Decimal("0"), actual_sales=Decimal("0")),
        ])

        meeting = ReviewSession(review_date=day1, participants="Production, Quality")
        db.add(meeting)
        db.flush()
        point_action = Action(
            action_no="ACT-2026-DSA01",
            reference_date=day1,
            problem_category="Delivery",
            problem_description="Dispatch was below target",
            action_description="Recover the shortfall",
        )
        linked_action = Action(
            action_no="ACT-2026-DSA02",
            reference_date=day1,
            problem_category="Delivery",
            problem_description="Part 2 dispatch confirmation",
            action_description="Confirm dispatch with customer",
        )
        db.add_all([point_action, linked_action])
        db.flush()
        db.add_all([
            ActionContext(action_id=point_action.id, context_date=day1, product_id=p1.id),
            ActionContext(action_id=linked_action.id, context_date=day1, product_id=p2.id),
            ActionWhyWhy(
                action_id=point_action.id,
                containment_action="Contain affected dispatch batch",
                why1="Daily dispatch commitment was missed",
                root_cause="No confirmed loading slot",
                corrective_action="Confirm and reserve dispatch slot",
                preventive_action="Review dispatch slot every morning",
                verification_method="Check next five dispatch days",
                verification_result="Verified for five dispatch days",
                lessons_learned="Escalate dispatch constraints before shift start",
            ),
            ReviewPoint(
                review_session_id=meeting.id,
                sequence_no=1,
                category="Delivery",
                discussion_point="[CRITICAL] | Part No: DSA-1 | Dispatch shortfall",
                action_required=True,
                product_id=p1.id,
                action_id=point_action.id,
            ),
            ReviewActionLink(review_session_id=meeting.id, action_id=point_action.id),
            ReviewActionLink(review_session_id=meeting.id, action_id=linked_action.id),
        ])
        db.commit()

    client = TestClient(app)
    headers = auth_headers(client)
    response = client.get(
        "/api/reports/dispatch-schedule-actual.xlsx?from_date=2026-10-01&to_date=2026-10-03",
        headers=headers,
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    workbook = load_workbook(BytesIO(response.content), data_only=True)
    sheet = workbook["Dispatch vs Actual"]
    assert sheet["A1"].value == "Part / Metric"
    assert [sheet.cell(1, col).value.date() for col in (2, 3, 4)] == [day1, day2, day3]

    # First product: one part heading, Plan, Actual, Gap, MOM row, then a blank.
    assert sheet["A2"].value.startswith("PART: DSA-1 - Dispatch Export Part 1")
    assert sheet["A3"].value == "Dispatch Plan"
    assert [sheet.cell(3, col).value for col in (2, 3, 4)] == [100, 125, 75]
    assert sheet["A4"].value == "Dispatch Actual"
    assert [sheet.cell(4, col).value for col in (2, 3, 4)] == [90, 100, 60]
    assert sheet["A5"].value == "Gap (Actual - Plan)"
    assert [sheet.cell(5, col).value for col in (2, 3, 4)] == [-10, -25, -15]
    assert sheet["A6"].value == "MOM Action / Discussion"
    p1_mom = sheet["B6"].value
    assert "MOM-2026-10-01-" in p1_mom
    assert "[CRITICAL]" in p1_mom
    assert "Part No: DSA-1 - Dispatch Export Part 1" in p1_mom
    assert "Review dispatch shortfall" in p1_mom
    assert "Action Plan: Recover the shortfall" in p1_mom
    assert "Containment: Contain affected dispatch batch" in p1_mom
    assert "Root Cause: No confirmed loading slot" in p1_mom
    assert "Corrective Action: Confirm and reserve dispatch slot" in p1_mom
    assert "Preventive Action: Review dispatch slot every morning" in p1_mom
    assert "Verification Method: Check next five dispatch days" in p1_mom
    assert "Verification Result: Verified for five dispatch days" in p1_mom
    assert "Lessons Learned: Escalate dispatch constraints before shift start" in p1_mom
    assert p1_mom.count("ACT-2026-DSA01") == 1
    assert sheet["A7"].value is None

    # Second product is included even though it has no Daily MIS actual row.
    assert sheet["A8"].value.startswith("PART: DSA-2 - Dispatch Export Part 2")
    assert sheet["A9"].value == "Dispatch Plan"
    assert sheet["B9"].value == 50
    assert sheet["A10"].value == "Dispatch Actual"
    assert sheet["B10"].value == 0
    assert sheet["A11"].value == "Gap (Actual - Plan)"
    assert sheet["B11"].value == -50
    assert sheet["A12"].value == "MOM Action / Discussion"
    assert "Part No: DSA-2 - Dispatch Export Part 2" in sheet["B12"].value
    assert "Confirm dispatch with customer" in sheet["B12"].value
    assert "Action Plan: Confirm dispatch with customer" in sheet["B12"].value
